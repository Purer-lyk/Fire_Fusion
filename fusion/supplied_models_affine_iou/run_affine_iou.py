"""Reuse genuine supplied-model predictions; measured affine polygon IoU only."""
import sys, json, csv, hashlib, re, subprocess, zipfile
from pathlib import Path
from datetime import datetime, timezone
from collections import Counter
import numpy as np
import cv2
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parent.parent))
import core as c
ROOT=Path('C:/lyk/fire/qualityDetection')
OUT=Path(__file__).resolve().parent
OLD=ROOT/'fusion/supplied_models_uv_reference'
CPP=Path('C:/lyk/paper/bpa_test')
SOURCE=Path('C:/lyk/paper_data/paper_data_/paper_data')
REPORT=ROOT/'reports/uploadedmodels_AffineIoU_DS_1528.docx'
NAMES={'rgb_baseline':'RGB检测基线','thermal_baseline':'热成像检测基线','affine_dynamic_repeated':'原C++结构：动态IoU两次','affine_spatial_then_classic':'动态RGB/热空间融合＋UV经典DS'}
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(name,obj): (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def affine_polygon(box,a):
    a=np.asarray(a,float)
    if a.shape!=(2,3) or not np.isfinite(a).all() or abs(np.linalg.det(a[:,:2]))<1e-12: raise ValueError('Nondegenerate finite 2x3 affine required')
    return c.box_polygon(box)@a[:,:2].T+a[:,2]
def estimate(src,dst):
    cv2.setRNGSeed(7)
    a,mask=cv2.estimateAffine2D(np.asarray(src,float),np.asarray(dst,float),method=cv2.RANSAC,ransacReprojThreshold=3.,maxIters=2000,confidence=.99,refineIters=10)
    if a is None or a.shape!=(2,3) or not np.isfinite(a).all(): raise ValueError('Affine fitting failed')
    return a,mask

def calibration():
    p=CPP/'main/homo_main.cpp'; text=re.sub(r'//[^\n]*','',p.read_text(encoding='utf-8'))
    points=[]
    for name in ['pts1','pts2']:
        block=re.search(r'vector<Point2f>\s+'+name+r'\s*=\s*\{(.*?)\};',text,re.S).group(1)
        points.append(np.array([(float(x),float(y)) for x,y in re.findall(r'Point2f\(([^,]+),\s*([^\)]+)\)',block)]))
    rgb,thermal=points; assert rgb.shape==thermal.shape==(7,2)
    a,mask=estimate(thermal,rgb)
    errors=np.linalg.norm(thermal@a[:,:2].T+a[:,2]-rgb,axis=1)
    loo=[]; loo_matrices=[]; loo_masks=[]
    for i in range(7):
        keep=np.arange(7)!=i; ai,mi=estimate(thermal[keep],rgb[keep]); loo.append(float(np.linalg.norm(thermal[i]@ai[:,:2].T+ai[:,2]-rgb[i])))
        loo_matrices.append(ai.tolist()); loo_masks.append(mi.ravel().tolist())
    data=dict(status='candidate_unvalidated: no independent scene/full-field checks',direction='thermal pts2 -> RGB pts1 original pixel coordinates',source=str(p),source_sha256=sha(p),thermal_points=thermal.tolist(),rgb_points=rgb.tolist(),affine_matrix_2x3=a.tolist(),estimator='cv2.estimateAffine2D FULL affine RANSAC; seed=7; threshold=3px; maxIters=2000; confidence=0.99; refineIters=10',inlier_mask=mask.ravel().tolist(),inlier_count=int(mask.sum()),fit_errors_px=errors.tolist(),fit_rmse_px=float(np.sqrt(np.mean(errors**2))),inlier_fit_rmse_px=float(np.sqrt(np.mean(errors[mask.ravel().astype(bool)]**2))),leave_one_out_errors_px=loo,leave_one_out_rmse_px=float(np.sqrt(np.mean(np.array(loo)**2))),leave_one_out_max_px=max(loo),leave_one_out_matrices=loo_matrices,leave_one_out_inlier_masks=loo_masks,note='Seven points occupy narrow y band; fitting and LOO do not validate independent scenes, field coverage or parallax. No manual offsets or projective mapping.')
    return a,data

def pairing(rgb,thermal,a):
    rd=rgb['detections']; td=thermal['detections']
    ri=max(range(len(rd)),key=lambda j:(rd[j]['score'],-j)) if rd else None
    # If RGB absent, use thermal highest score for its evidence, not a fabricated match.
    ti=max(range(len(td)),key=lambda j:(td[j]['score'],-j)) if td else None
    s=0.; candidates=[]
    if ri is not None and td:
        for j,d in enumerate(td): candidates.append((c.polygon_iou(c.box_polygon(rd[ri]['bbox_xyxy']),affine_polygon(d['bbox_xyxy'],a)),d['score'],-j,j))
        s,_,_,ti=max(candidates)
        status='positive_overlap' if s>0 else 'both_present_no_overlap'
    else: status='missing_both' if ri is None and ti is None else ('missing_rgb' if ri is None else 'missing_thermal')
    r=rd[ri] if ri is not None else None; t=td[ti] if ti is not None else None
    return dict(s=s,status=status,rgb_detection_index=ri,thermal_detection_index=ti,rgb_box=r['bbox_xyxy'] if r else None,thermal_box=t['bbox_xyxy'] if t else None,affine_thermal_corners=affine_polygon(t['bbox_xyxy'],a).tolist() if t else None,rgb_confidence=r['score'] if r else 0.,thermal_confidence=t['score'] if t else 0.,thermal_frame_max=max((d['score'] for d in td),default=0.),pair_available=ri is not None and ti is not None,candidate_ious=[{'thermal_index':j,'iou':v} for v,_,_,j in candidates],confidence_source='RGB highest retained score; if RGB present thermal highest polygon IoU, ties score then original index; if RGB missing thermal highest score; absent confidence=0')
def metrics(y,p):
    assert len(y)==len(p)>0 and all(v in (0,1) for v in y+p)
    tp=sum(a==b==1 for a,b in zip(y,p)); tn=sum(a==b==0 for a,b in zip(y,p)); fp=sum(a==0 and b==1 for a,b in zip(y,p)); fn=sum(a==1 and b==0 for a,b in zip(y,p))
    return dict(N=len(y),TP=tp,FP=fp,FN=fn,TN=tn,accuracy=(tp+tn)/len(y),precision=tp/(tp+fp) if tp+fp else None,recall=tp/(tp+fn) if tp+fn else None,F1=2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else None)

def run():
    OUT.mkdir(exist_ok=True); (OUT/'examples').mkdir(exist_ok=True)
    protected={str(p):sha(p) for p in list(CPP.rglob('*.cpp'))+list(CPP.rglob('*.h'))+[ROOT/'reports/uploadedmodels_UVreference_1528.docx']}
    samples=[c.parse_sample_line(l) for l in (SOURCE/'paper_sample.txt').read_text(encoding='utf-8-sig').splitlines()]; assert len(samples)==1528
    pred={}; provenance={}
    for mode in ['rgb','thermal']:
        path=OLD/(mode+'_predictions.json'); data=json.loads(path.read_text(encoding='utf-8')); images=data['images']; assert len(images)==1528
        assert len({x['pair_index'] for x in images})==1528
        for i,(rec,sample) in enumerate(zip(images,samples)):
            assert rec['pair_index']==i and rec['source_line']==i+1 and rec['image_name']==getattr(sample,mode)
            assert sha(rec['image_path'])==rec['image_sha256']
            im=cv2.imread(rec['image_path']); assert im.shape[:2]==(rec['height'],rec['width'])
            for d in rec['detections']: assert d['class_id']==0 and d['score']>.5; c.box_polygon(d['bbox_xyxy'])
        assert sha(data['metadata']['checkpoint'])==data['metadata']['checkpoint_sha256']
        pred[mode]=images; provenance[mode]=dict(prediction_path=str(path),prediction_sha256=sha(path),checkpoint_sha256=data['metadata']['checkpoint_sha256'],architecture=data['metadata']['architecture'],rows=1528,retained_boxes=sum(len(r['detections']) for r in images),actual_image_dimensions=dict(Counter(f"{r['width']}x{r['height']}" for r in images)),reused_without_reinference=True,current_source_image_hashes_verified=True)
    a,cal=calibration(); save('candidate_affine_calibration.json',cal)
    rows=[]; pairs=[]; expected=[]; inputs=[]; previous=None; state=c.SensorState(); sequence=0; resets=[]
    for i,sample in enumerate(samples):
        ts=datetime.strptime(sample.rgb.removesuffix('_bgr.jpg'),'%Y_%m_%d_%H_%M_%S_%f'); gap=None if previous is None else (ts-previous).total_seconds(); reset=previous is None or gap<=0 or gap>120
        if reset: state=c.SensorState(); sequence+=1; resets.append(dict(pair_index=i,gap_seconds=gap))
        previous=ts; rgb=pred['rgb'][i]; thermal=pred['thermal'][i]; pair=pairing(rgb,thermal,a); pairs.append(dict(pair_index=i,**pair))
        sr=c.sensor_step(state,sample,pair['rgb_confidence'],pair['thermal_confidence']); state=sr.state
        ab=c.combine(sr.rgb,sr.thermal,'spatial',pair['s'])
        repeated=c.fuse3(sr.rgb,sr.thermal,sr.uv,'spatial',pair['s']); classic=c.fuse_rgb_thermal_uv(sr.rgb,sr.thermal,sr.uv,'spatial',pair['s'])
        row=dict(pair_index=i,source_line=i+1,rgb_name=sample.rgb,thermal_name=sample.thermal,Tmax=sample.temperature,UV_input=sample.uv,UV_reference_label=sample.uv,sequence_id=sequence,state_reset=int(reset),gap_seconds=gap,spatial_polygon_iou=pair['s'],spatial_status=pair['status'],pair_available=int(pair['pair_available']),rgb_detection_index=pair['rgb_detection_index'],thermal_detection_index=pair['thermal_detection_index'],rgb_box_xyxy=json.dumps(pair['rgb_box']),thermal_box_xyxy=json.dumps(pair['thermal_box']),affine_thermal_corners=json.dumps(pair['affine_thermal_corners']),rgb_confidence=pair['rgb_confidence'],thermal_confidence=pair['thermal_confidence'],thermal_frame_max=pair['thermal_frame_max'],rgb_box_count=len(rgb['detections']),thermal_box_count=len(thermal['detections']),rgb_baseline_binary05=int(bool(rgb['detections'])),thermal_baseline_binary05=int(bool(thermal['detections'])),confidence_source=pair['confidence_source'])
        for prefix,m in [('rgb_bpa',sr.rgb),('thermal_bpa',sr.thermal),('uv_bpa',sr.uv),('rgb_thermal_spatial',ab.mass),('affine_dynamic_repeated',repeated.mass),('affine_spatial_then_classic',classic.mass)]:
            for key,v in zip(['fire','nofire','uncertain'],m.values()): row[prefix+'_'+key]=v
        for name,res in [('affine_dynamic_repeated',repeated),('affine_spatial_then_classic',classic)]:
            row[name+'_binary05']=int(res.mass.fire>=.5); row[name+'_decision095']=c.decide(res.mass,.95); row[name+'_degenerate']=int(res.degenerate)
        row['conflict_ab_spatial']=ab.conflict
        rows.append(row); inputs.append(f'{int(reset)} {pair["rgb_confidence"]:.17g} {pair["thermal_confidence"]:.17g} {sample.temperature:.17g} {sample.uv} {pair["s"]:.17g}\n')
        expected.append([*sr.rgb.values(),*sr.thermal.values(),*sr.uv.values(),*ab.mass.values(),*repeated.mass.values(),*classic.mass.values()])
    # Link original C++ sources unchanged; use fuse_sp explicitly with this frame's actual IoU.
    exe=OUT/'cpp_dynamic_reference.exe'; cmd=['C:/msys64/ucrt64/bin/g++.exe','-O2','-static','-std=c++17','-I'+str(CPP/'include'),str(OUT/'cpp_dynamic_reference.cpp')]+[str(CPP/'src'/n) for n in ['bpa.cpp','vision_bpa.cpp','temp_bpa.cpp','uv_bpa.cpp','ds_fusion.cpp']]+['-o',str(exe)]
    build=subprocess.run(cmd,capture_output=True,text=True); (OUT/'cpp_build.log').write_text(build.stdout+build.stderr,encoding='utf-8'); assert build.returncode==0,build.stderr
    proc=subprocess.run([str(exe)],input=''.join(inputs),capture_output=True,text=True,check=True)
    ref=np.array([list(map(float,l.split())) for l in proc.stdout.splitlines()]); exp=np.array(expected); assert ref.shape==exp.shape==(1528,18) and np.isfinite(ref).all()
    err=float(np.abs(ref-exp).max()); assert err<1e-10
    (OUT/'cpp_reference_inputs.txt').write_text(''.join(inputs),encoding='utf-8'); (OUT/'cpp_reference_results.txt').write_text(proc.stdout,encoding='utf-8')
    save('cpp_parity_verification.json',dict(rows=1528,values_per_row=18,maximum_absolute_error=err,build_command=cmd,build_exit=build.returncode,reference_exit=proc.returncode,reference_sha256=sha(exe),actual_dynamic_s=True))
    y=[r['UV_reference_label'] for r in rows]; met={name:metrics(y,[r[name+'_binary05'] for r in rows]) for name in NAMES}; coverage={}
    for name in ['affine_dynamic_repeated','affine_spatial_then_classic']:
        counts=Counter(r[name+'_decision095'] for r in rows); selected=[r for r in rows if r[name+'_decision095']!='uncertain']
        coverage[name]=dict(fire=counts['fire'],nofire=counts['nofire'],uncertain=counts['uncertain'],decision_coverage=len(selected)/1528,uncertainty_rate=counts['uncertain']/1528,degenerate_count=sum(r[name+'_degenerate'] for r in rows),covered_only_metrics=metrics([r['UV_reference_label'] for r in selected],[int(r[name+'_decision095']=='fire') for r in selected]) if selected else None)
    with (OUT/'fusion_results.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    save('fusion_results.json',dict(primary='affine_dynamic_repeated',scientifically_nonspatial_UV_variant='affine_spatial_then_classic',rows=rows)); save('geometry_pairing.json',pairs)
    metric_data=dict(reference='User column4 UV0=nofire UV1=fire; UV is also fusion input, circular agreement not independent truth',methods=met,UV_label_counts=dict(Counter(y)),abstention=coverage,binary_threshold=.5,abstention_threshold=.95); save('metrics.json',metric_data)
    # Five real samples: highest overlap, low positive, multiple, missing, nonoverlap where available.
    selected=[]; sample_labels=[]
    selectors=[('最大正重叠',lambda r:r['spatial_status']=='positive_overlap',True),('较小正重叠',lambda r:r['spatial_status']=='positive_overlap',False),('多目标配对',lambda r:r['thermal_box_count']>1 or r['rgb_box_count']>1,True),('缺失模态',lambda r:r['spatial_status'].startswith('missing'),False),('双模态无重叠',lambda r:r['spatial_status']=='both_present_no_overlap',False)]
    for label,condition,descending in selectors:
        choices=sorted([r for r in rows if condition(r) and r['pair_index'] not in selected],key=lambda r:r['spatial_polygon_iou'],reverse=descending)
        if not choices: choices=[r for r in rows if r['pair_index'] not in selected]; label+='（该类不存在，替代真实帧）'
        selected.append(choices[0]['pair_index']); sample_labels.append(label)
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',18); examples=[]
    for index,label in zip(selected,sample_labels):
        rec=pred['rgb'][index]; im=Image.open(rec['image_path']).convert('RGB'); draw=ImageDraw.Draw(im); pair=pairs[index]
        for j,d in enumerate(rec['detections']):
            draw.rectangle(d['bbox_xyxy'],outline=(0,220,0),width=2); draw.text((d['bbox_xyxy'][0],max(0,d['bbox_xyxy'][1]-20)),f'R{j} {d["score"]:.3f}',font=font,fill=(0,220,0))
        for j,d in enumerate(pred['thermal'][index]['detections']):
            poly=affine_polygon(d['bbox_xyxy'],a); xy=[tuple(x) for x in poly]; color=(255,0,255) if j==pair['thermal_detection_index'] else (200,140,200)
            draw.line(xy+[xy[0]],fill=color,width=3 if j==pair['thermal_detection_index'] else 1); draw.text(xy[0],f'T{j} {d["score"]:.3f}',font=font,fill=color)
        im=im.resize((960,round(im.height*960/im.width))); canvas=Image.new('RGB',(960,im.height+105),(20,20,20)); canvas.paste(im,(0,105)); d=ImageDraw.Draw(canvas)
        d.text((10,5),f'{label} | pair {index} | UV={y[index]} | s={pair["s"]:.6f}',font=font,fill='white')
        d.text((10,35),f'RGB锚R{pair["rgb_detection_index"]} + 仿射热T{pair["thermal_detection_index"]} | {pair["status"]}',font=font,fill='white')
        d.text((10,65),'绿色RGB检测框；紫色选中热框四角仿射映射；IoU按多边形面积计算（非外包矩形）',font=font,fill='white')
        dest=OUT/'examples'/f'pair_{index}.png'; canvas.save(dest); examples.append(dict(pair_index=index,path=str(dest),s=pair['s'],status=pair['status'],label=label,sha256=sha(dest)))
    summary=dict(created_utc=datetime.now(timezone.utc).isoformat(),rows=1528,provenance=provenance,calibration=cal,spatial_status_counts=dict(Counter(r['spatial_status'] for r in rows)),spatial_nonzero_rows=sum(r['spatial_polygon_iou']>0 for r in rows),selected_thermal_not_frame_max=sum(r['thermal_confidence']!=r['thermal_frame_max'] for r in rows),metrics=met,abstention=coverage,examples=examples,resets=resets,sequence_count=sequence,cpp_max_abs_error=err,source_sample_sha256=sha(SOURCE/'paper_sample.txt'))
    save('execution_summary.json',summary); report(summary)
    with zipfile.ZipFile(REPORT) as z: assert z.testzip() is None; embedded=len([n for n in z.namelist() if n.startswith('word/media/')])
    doc=Document(REPORT); readback='\n'.join([p.text for p in doc.paragraphs]+[cell.text for t in doc.tables for rr in t.rows for cell in rr.cells]); assert all(t in readback for t in ['1528','动态IoU两次','UV经典DS','2×3','循环','未验证'])
    (OUT/'report_readback.txt').write_text(readback,encoding='utf-8'); loaded=list(csv.DictReader((OUT/'fusion_results.csv').open(encoding='utf-8-sig'))); assert len(loaded)==1528 and [int(r['pair_index']) for r in loaded]==list(range(1528))
    for name in NAMES: assert metrics([int(r['UV_reference_label']) for r in loaded],[int(r[name+'_binary05']) for r in loaded])==met[name]
    assert embedded==5 and all(sha(p)==v for p,v in protected.items())
    verify=dict(csv_rows=1528,unique_ids=1528,metrics_recomputed_from_csv=True,docx_zip_crc_ok=True,docx_readback_ok=True,embedded_images=embedded,source_and_previous_report_unchanged=True,prediction_reinference=False,docx_path=str(REPORT),docx_sha256=sha(REPORT),csv_sha256=sha(OUT/'fusion_results.csv'),metrics_sha256=sha(OUT/'metrics.json'),cpp_parity_max_abs_error=err)
    save('deliverable_verification.json',verify); print(json.dumps(dict(metrics=met,abstention=coverage,calibration=cal,status_counts=summary['spatial_status_counts'],selected_thermal_not_frame_max=summary['selected_thermal_not_frame_max'],examples=examples,verification=verify),ensure_ascii=False,indent=2))

def report(s):
    d=Document(); style=d.styles['Normal']; style.font.name='Microsoft YaHei'; style.font.size=Pt(10); style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    def p(text): d.add_paragraph(text)
    def h(text): d.add_heading(text,1)
    def table(headers,records):
        t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
        for cell,v in zip(t.rows[0].cells,headers): cell.text=str(v)
        for record in records:
            for cell,v in zip(t.add_row().cells,record): cell.text=str(v)
    def f(v): return '未定义' if v is None else f'{v:.6f}'
    d.add_heading('上传模型：仿射IoU动态空间系数D-S融合',0)
    p('完整1528对真实预测重用，不重新推理。best.pt为RGB、best-2.pt为热成像；实际架构YOLO26n而非PPYOLOv2。只使用class0 fire、score>0.5原始保留框。第三列Tmax与第四列UV保持原输入语义。用户UV0/1作为参考标准，同时UV参与融合，属于循环评价的UV符合度，不是独立场景真值或定位AP。未验证模型训练/测试独立性，不承诺融合性能提升。')
    h('1. FULL affine候选配准：2×3，不是投影单应性')
    p('从当前main/homo_main.cpp七个未注释thermal pts2→RGB pts1点估计cv2.estimateAffine2D FULL affine，非estimateAffinePartial2D。RANSAC阈值3px，seed7，maxIters2000，confidence0.99，refineIters10。仅2×3线性仿射加平移，无透视除法，无人工矩形偏移，无固定空间系数。')
    p('矩阵：'+json.dumps(s['calibration']['affine_matrix_2x3'])+'；inliers='+str(s['calibration']['inlier_mask']))
    p('全部点fit RMSE='+f(s['calibration']['fit_rmse_px'])+'px；inlier fit RMSE='+f(s['calibration']['inlier_fit_rmse_px'])+'px；留一RMSE='+f(s['calibration']['leave_one_out_rmse_px'])+'px，最大='+f(s['calibration']['leave_one_out_max_px'])+'px。七点集中狭窄y区域，无独立场景/全视场检查，因此candidate_unvalidated（未验证），误差指标不能证明跨场景配准有效。')
    table(['热点','RGB点','fit误差px','留一误差px'],[[a,b,f(e),f(l)] for a,b,e,l in zip(s['calibration']['thermal_points'],s['calibration']['rgb_points'],s['calibration']['fit_errors_px'],s['calibration']['leave_one_out_errors_px'])])
    h('2. 实际polygon IoU与置信度配对政策')
    p('热矩形四角q=A[:,0:2]p+A[:,2]映射到RGB原始像素坐标，形成凸四边形。s=area(RGB框∩映射热四边形)/[area(RGB框)+area(映射热四边形)-intersection]；cv2.intersectConvexConvex计算交面积。使用完整原始多边形，不裁剪到图像边界，不以轴对齐外包矩形替代。')
    p('RGB最高分fire框为锚；在全部保留热框中选择实际polygon IoU最高者，IoU相等优先置信度，再原始检测索引较小。RGB BPA用锚分数、热BPA用该选中框分数，几何与BPA来源一致，非各模态独立最大分。双模态均存在但无重叠s=0，仍按同一政策选最高分热框；RGB缺失时无有效配对，热BPA使用最高保留分数；热缺失时热置信度0；双缺失时两者0。所有缺失状态s=0且明确标记。s=0不证明无火。检测基线仍是本模态任一保留fire框存在。')
    p('状态计数：'+json.dumps(s['spatial_status_counts'],ensure_ascii=False)+'；实际非零IoU帧='+str(s['spatial_nonzero_rows'])+'；选中热分数不同于帧最高分的帧='+str(s['selected_thermal_not_frame_max'])+'。CSV存原框坐标/索引/分数、仿射四角、s、状态与BPA；geometry_pairing.json另存候选全部IoU。')
    h('3. 原C++结构与非空间UV变体：分别报告')
    p('主要兼容结果“原C++结构：动态IoU两次”：ab=fuse_sp(RGB,thermal,s)，abc=fuse_sp(ab,UV,s)。仅把原fuse3两个空间调用的固定值替换为实际s；第二次重复同一RGB–热IoU是保留原结构的算法规则，不是假称UV具有空间框。科学上非空间UV变体“动态RGB/热空间融合＋UV经典DS”：ab=fuse_sp(RGB,thermal,s)，abc=fuse(ab,UV)。两种定义冻结后计算，不能按测试结果择优包装。')
    p('原公式：Ksp=(1-s)(aF*bN+aN*bF)，F=s(aF*bF+aF*bU+aU*bF)/(1-Ksp)，N=(aN*bN+aN*bU+aU*bN)/(1-Ksp)，U=aU*bU/(1-Ksp)，随后归一化。共同分母归一化后抵消，实际比例(sFsupport,Nsupport,Usupport)，不能单独宣称动态冲突系数提升性能。经典DS不乘s。总/近总冲突或零支持显式回退(0,0,1)并标degenerate；真实数据退化数见表。')
    p('BPA忠实重用core.py：窗口10，稳定性exp(-std)，λ0.7平滑；RGB可靠性0.3；热PC=1-exp(-10c)、PT=clamp(1-exp(-0.1(Tmax-50)),0,1)；UV q初始0.001，UV1增0.01q(1-q)、UV0减0.1q(1-q)，原UV BPA(q,0.8(1-q),0.2(1-q))。这些0.1是温度/UV公式参数，不是空间系数。UV不是直接注入硬标签BPA。原startup与平滑均保留。')
    p('时序沿原文件顺序；开始/时间不递增/间隔>120s重置，属既有启发式而非独立验证场景界线。sequence数='+str(s['sequence_count'])+'；resets='+json.dumps(s['resets']))
    p('重新编译链接原始C++四类源文件，只读未改；参考入口显式传每帧实际s调用原fuse_sp及fuse。全部1528行×18个BPA/融合值与Python最大绝对误差='+f'{s["cpp_max_abs_error"]:.3e}'+'，exit0。')
    h('4. 全量UV参考帧级指标：阈值0.5')
    p('融合mF≥0.5为fire，否则nofire；基线任一严格score>0.5框为fire。Accuracy=(TP+TN)/N，Precision=TP/(TP+FP)，Recall=TP/(TP+FN)，F1=2TP/(2TP+FP+FN)，零分母未定义。所有1528帧均纳入，无丢弃困难帧。')
    table(['方法','TP','FP','FN','TN','Accuracy','P','R','F1'],[[NAMES[k],v['TP'],v['FP'],v['FN'],v['TN'],f(v['accuracy']),f(v['precision']),f(v['recall']),f(v['F1'])] for k,v in s['metrics'].items()])
    h('5. 0.95弃权与覆盖子集')
    p('mF≥0.95 fire；否则mN≥0.95 nofire；其余uncertain。覆盖率=(fire+nofire)/1528；弃权不视为正确。覆盖子集指标与全量0.5分开，不能混比。')
    table(['方法','fire','nofire','uncertain','覆盖率','退化'],[[NAMES[k],v['fire'],v['nofire'],v['uncertain'],f(v['decision_coverage']),v['degenerate_count']] for k,v in s['abstention'].items()])
    table(['方法（仅覆盖）','N','TP','FP','FN','TN','Acc','P','R','F1'],[[NAMES[k]]+([v['covered_only_metrics'][x] for x in ['N','TP','FP','FN','TN']]+[f(v['covered_only_metrics'][x]) for x in ['accuracy','precision','recall','F1']] if v['covered_only_metrics'] else ['无覆盖']*9) for k,v in s['abstention'].items()])
    h('6. 五个真实RGB叠加样例')
    p('原始RGB检测框和实际仿射热框四角叠加；检测索引与配对s明确标注，中文由TrueType字体渲染。非标注真值；样例按情况展示，不是随机性能代表。')
    for ex in s['examples']: p(ex['label']+'；pair '+str(ex['pair_index'])+'；s='+f(ex['s'])); d.add_picture(ex['path'],width=Inches(6.2))
    h('7. 来源与交付核验')
    for mode,pr in s['provenance'].items(): p(mode+'：'+json.dumps(pr,ensure_ascii=False))
    p('结果目录fusion/supplied_models_affine_iou；旧uploadedmodels_UVreference_1528.docx及原算法、图像、预测保持不变。附CSV/JSON、指标、七点2×3矩阵及LOO、C++参考输入输出、五个PNG及验证清单。当前paper_sample SHA256='+s['source_sample_sha256']+'。配准候选未独立验证；UV循环评价不能支撑真实场景泛化或独立性能提升结论。')
    REPORT.parent.mkdir(exist_ok=True); d.save(REPORT)
if __name__=='__main__': run()
