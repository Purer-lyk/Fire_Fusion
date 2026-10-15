"""Execute fixed nominal spatial reliability D-S on genuine 1528 cached pairs."""
from pathlib import Path
import sys,json,csv,hashlib,math,zipfile,subprocess
from datetime import datetime,timezone
from collections import Counter
import numpy as np
import cv2
import spatial_core as s
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
OUT=Path(__file__).resolve().parent
ROOT=OUT.parent.parent
OLD=OUT.parent/'supplied_models_uv_reference'
AFF=OUT.parent/'supplied_models_affine_iou'
SOURCE=Path('C:/lyk/paper_data/paper_data_/paper_data')
REPORT=ROOT/'reports/spatial_reliability_DS_1528_最终对比报告.docx'
sys.path.insert(0,str(OUT.parent))
import core as legacy
NAMES={'rgb':'RGB检测基线','thermal':'热成像检测基线','ordinary_ds':'普通DS（无空间、无UV）','spatial_ds':'IoU可靠性DS（无UV）','ordinary_uv':'普通DS＋瞬时UV','spatial_uv':'IoU可靠性DS＋瞬时UV（新方案）','ordinary_legacy_uv':'普通DS＋原时序UV','spatial_legacy_uv':'IoU可靠性DS＋原时序UV','legacy_thermal_spatial':'原温度热BPA＋新RGB／IoU（无UV）','legacy_thermal_uv':'原温度热BPA＋新RGB／IoU＋瞬时UV','original_dynamic':'原C++动态IoU重复两次','original_fixed':'原C++固定0.1重复两次','uv_alone':'UV单独（定义恒等）'}

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def load(p): return json.loads(Path(p).read_text(encoding='utf-8'))
def save(name,data): (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def load_config(): return load(OUT/'frozen_config.json')
def temperature_reliability(t,cfg):
    if cfg['temperature_unit'] is None: return 1.
    center,scale=cfg['temperature_center'],cfg['temperature_scale']
    if center is None or scale is None or scale<=0: raise ValueError('calibrated temperature center/scale required')
    z=float(np.clip((t-center)/scale,-60,60))
    return .5+.5/(1+math.exp(-z))

def metrics(y,p):
    if len(y)!=len(p) or not y or any(v not in (0,1) for v in list(y)+list(p)): raise ValueError('binary arrays required')
    tp=sum(a==b==1 for a,b in zip(y,p)); tn=sum(a==b==0 for a,b in zip(y,p))
    fp=sum(a==0 and b==1 for a,b in zip(y,p)); fn=sum(a==1 and b==0 for a,b in zip(y,p))
    return dict(N=len(y),TP=tp,FP=fp,FN=fn,TN=tn,accuracy=(tp+tn)/len(y),precision=tp/(tp+fp) if tp+fp else None,recall=tp/(tp+fn) if tp+fn else None,F1=2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else None)

def frame(rgb,thermal,a,t,uv,legacy_uv,cfg):
    candidates=s.match(rgb,thermal,a,cfg['matching_min_iou'])
    if not candidates: candidates=[dict(rgb_index=None,thermal_index=None,iou=0.,status='missing_both')]
    details=[]
    for cand in candidates:
        ri,ti=cand['rgb_index'],cand['thermal_index']
        rc=rgb[ri]['score'] if ri is not None else None; tc=thermal[ti]['score'] if ti is not None else None
        rm=s.base(rc,cfg['base_ignorance']); tm=s.discount(s.base(tc,cfg['base_ignorance']),temperature_reliability(t,cfg))
        ordinary,k,d=s.ds(rm,tm)
        r=cfg['matched_reliability_floor']+cfg['matched_reliability_slope']*cand['iou'] if cand['status']=='matched' else cfg['unmatched_present_reliability']
        dr=s.discount(rm,r); dt=s.discount(tm,r); spatial,ks,ds=s.ds(dr,dt)
        details.append(dict(**cand,rgb_confidence=rc,thermal_confidence=tc,reliability=r,rgb_base=rm.tolist(),thermal_base=tm.tolist(),rgb_discounted=dr.tolist(),thermal_discounted=dt.tolist(),ordinary_mass=ordinary.tolist(),spatial_mass=spatial.tolist(),ordinary_conflict=k,spatial_conflict=ks,degenerate=d or ds))
    # Same pre-UV candidate for every ablation: no label-conditioned selection.
    index=max(range(len(details)),key=lambda j:(details[j]['ordinary_mass'][0],min([x for x in [details[j]['rgb_confidence'],details[j]['thermal_confidence']] if x is not None],default=0),-j))
    chosen=details[index]; ordinary=np.array(chosen['ordinary_mass']); spatial=np.array(chosen['spatial_mass'])
    um=np.array([cfg['uv_instant_reliability']*uv,cfg['uv_instant_reliability']*(1-uv),1-cfg['uv_instant_reliability']])
    masses=dict(ordinary_ds=ordinary,spatial_ds=spatial,ordinary_uv=s.ds(ordinary,um)[0],spatial_uv=s.ds(spatial,um)[0],ordinary_legacy_uv=s.ds(ordinary,legacy_uv)[0],spatial_legacy_uv=s.ds(spatial,legacy_uv)[0],uv_alone=um)
    return masses,dict(selected_candidate=index,selected=chosen,candidates=details,instant_uv_mass=um.tolist(),legacy_uv_mass=list(legacy_uv),temperature_reliability=temperature_reliability(t,cfg))

def run():
    cfg=load_config(); receipt=load(OUT/'config_freeze_receipt.json'); assert sha(OUT/'frozen_config.json')==receipt['sha256']
    protected_paths=list(Path('C:/lyk/paper/bpa_test').rglob('*.cpp'))+list(Path('C:/lyk/paper/bpa_test').rglob('*.h'))+list((ROOT/'reports').glob('*.docx'))
    protected_paths += [OLD/'rgb_predictions.json',OLD/'thermal_predictions.json',AFF/'candidate_affine_calibration.json',OLD/'fusion_results.json',AFF/'fusion_results.json',SOURCE/'paper_sample.txt']
    protected={str(p):sha(p) for p in protected_paths if p!=REPORT}
    cal=load(AFF/'candidate_affine_calibration.json'); assert sha(cal['source'])==cal['source_sha256']; a=cal['affine_matrix_2x3']
    samples=[legacy.parse_sample_line(l) for l in (SOURCE/'paper_sample.txt').read_text(encoding='utf-8-sig').splitlines()]
    assert len(samples)==1528 and len({x.rgb for x in samples})==len({x.thermal for x in samples})==1528
    pred={}; provenance={}
    for mode in ['rgb','thermal']:
        path=OLD/(mode+'_predictions.json'); data=load(path); records=data['images']; assert len(records)==1528
        assert len({r['pair_index'] for r in records})==1528
        for i,(r,sm) in enumerate(zip(records,samples)):
            assert r['pair_index']==i and r['source_line']==i+1 and r['image_name']==getattr(sm,mode)
            assert sha(r['image_path'])==r['image_sha256']; im=cv2.imread(r['image_path']); assert im is not None and im.shape[:2]==(r['height'],r['width'])
            for d in r['detections']: assert d['score']>.5 and d['class_id']==0; s.polygon(d['bbox_xyxy'])
        assert sha(data['metadata']['checkpoint'])==data['metadata']['checkpoint_sha256']
        provenance[mode]=dict(path=str(path),sha256=sha(path),metadata=data['metadata'],rows=1528,boxes=sum(len(r['detections']) for r in records),missing_frames=sum(not r['detections'] for r in records),dimensions=dict(Counter(f"{r['width']}x{r['height']}" for r in records)))
        pred[mode]=records
    old=load(OLD/'fusion_results.json')['rows']; aff=load(AFF/'fusion_results.json')['rows']
    # Re-exercise exact original C++ executable/inputs and verify stored genuine output.
    cpp={}
    for tag,folder,exe in [('fixed',OLD,'cpp_reference_harness.exe'),('dynamic',AFF,'cpp_dynamic_reference.exe')]:
        proc=subprocess.run([str(folder/exe)],input=(folder/'cpp_reference_inputs.txt').read_text(encoding='utf-8'),capture_output=True,text=True,check=True)
        output=np.array([list(map(float,l.split())) for l in proc.stdout.splitlines()]); cached=np.loadtxt(folder/'cpp_reference_results.txt')
        assert output.shape==cached.shape and output.shape[0]==1528 and np.max(np.abs(output-cached))<1e-10
        cpp[tag]=dict(rows=1528,exit=proc.returncode,max_abs_error=float(np.max(np.abs(output-cached))),exe_sha256=sha(folder/exe),inputs_sha256=sha(folder/'cpp_reference_inputs.txt'))
    rows=[]; detail=[]; state=legacy.SensorState(); previous=None; resets=[]
    for i,sm in enumerate(samples):
        assert old[i]['pair_index']==aff[i]['pair_index']==i and old[i]['UV_reference_label']==aff[i]['UV_reference_label']==sm.uv
        lu=[old[i]['uv_bpa_'+x] for x in ['fire','nofire','uncertain']]
        masses,d=frame(pred['rgb'][i]['detections'],pred['thermal'][i]['detections'],a,sm.temperature,sm.uv,lu,cfg)
        selected=d['selected']; ts=datetime.strptime(sm.rgb.removesuffix('_bgr.jpg'),'%Y_%m_%d_%H_%M_%S_%f')
        gap=None if previous is None else (ts-previous).total_seconds(); reset=previous is None or gap<=0 or gap>cfg['legacy_reset_gap_seconds']; previous=ts
        if reset: state=legacy.SensorState(); resets.append(i)
        sr=legacy.sensor_step(state,sm,selected['rgb_confidence'] or 0.,selected['thermal_confidence'] or 0.); state=sr.state
        # Same candidate, legacy thermal temperature/history only; absent thermal stays vacuous.
        thermal_mass=np.array(sr.thermal.values()) if selected['thermal_index'] is not None else s.base(None)
        legacy_ab=s.ds(np.array(selected['rgb_discounted']),s.discount(thermal_mass,selected['reliability']))[0]
        masses['legacy_thermal_spatial']=legacy_ab; masses['legacy_thermal_uv']=s.ds(legacy_ab,np.array(d['instant_uv_mass']))[0]
        masses['original_dynamic']=np.array([aff[i]['affine_dynamic_repeated_'+k] for k in ['fire','nofire','uncertain']])
        masses['original_fixed']=np.array([old[i]['legacy_fixed_01_'+k] for k in ['fire','nofire','uncertain']])
        row=dict(pair_index=i,source_line=i+1,rgb_name=sm.rgb,thermal_name=sm.thermal,Tmax=sm.temperature,UV_input=sm.uv,UV_reference_label=sm.uv,rgb_count=len(pred['rgb'][i]['detections']),thermal_count=len(pred['thermal'][i]['detections']),selected_candidate=d['selected_candidate'],rgb_index=selected['rgb_index'],thermal_index=selected['thermal_index'],polygon_iou=selected['iou'],status=selected['status'],reliability=selected['reliability'],rgb_confidence=selected['rgb_confidence'],thermal_confidence=selected['thermal_confidence'],temperature_reliability=d['temperature_reliability'],rgb_binary=int(bool(pred['rgb'][i]['detections'])),thermal_binary=int(bool(pred['thermal'][i]['detections'])))
        for name,m in masses.items():
            assert np.isfinite(m).all() and min(m)>=-1e-12 and np.isclose(sum(m),1)
            for key,val in zip(['fire','nofire','uncertain'],m): row[name+'_'+key]=float(val)
            row[name+'_betp']=float(m[0]+.5*m[2]); row[name+'_binary']=s.binary(m)
            row[name+'_mf_binary']=int(m[0]>=.5)
            row[name+'_decision08']=s.decision(m,.8); row[name+'_decision095']=s.decision(m,.95)
            row[name+'_vacuous']=int(np.allclose(m,[0,0,1],atol=1e-12)); row[name+'_tie']=int(abs(m[0]+.5*m[2]-.5)<=1e-12)
        # Old published thresholds preserved explicitly; common BetP supplied separately.
        row['original_dynamic_binary']=aff[i]['affine_dynamic_repeated_binary05']; row['original_fixed_binary']=old[i]['legacy_fixed_01_binary05']
        rows.append(row); detail.append(dict(pair_index=i,**d,legacy_thermal_mass=thermal_mass.tolist(),legacy_reset=int(reset)))
    y=[sm.uv for sm in samples]; met={name:metrics(y,[r[name+'_binary'] for r in rows]) for name in NAMES}
    common={name:metrics(y,[int(r[name+'_betp']>.5+1e-12) for r in rows]) for name in NAMES if name not in ['rgb','thermal']}
    mf={name:metrics(y,[r[name+'_mf_binary'] for r in rows]) for name in NAMES if name not in ['rgb','thermal']}
    coverage={}
    for name in common:
        coverage[name]={}
        for code in ['08','095']:
            counts=Counter(r[name+'_decision'+code] for r in rows); selected=[r for r in rows if r[name+'_decision'+code]!='uncertain']
            coverage[name][code]=dict(**{k:counts[k] for k in ['fire','nofire','uncertain']},coverage=len(selected)/1528,covered_only_metrics=metrics([r['UV_reference_label'] for r in selected],[int(r[name+'_decision'+code]=='fire') for r in selected]) if selected else None)
        coverage[name]['vacuous_count']=sum(r[name+'_vacuous'] for r in rows); coverage[name]['tie_count']=sum(r[name+'_tie'] for r in rows)
    with (OUT/'fusion_results.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    save('fusion_results.json',dict(rows=rows)); save('geometry_evidence.json',detail)
    allmet=dict(reference='UV-reference agreement; UV also input: circular NOT independent scene validity',methods=met,common_betp_metrics=common,common_mf05_metrics=mf,coverage=coverage,label_counts=dict(Counter(y)),parameters_sha256=receipt['sha256'])
    save('metrics.json',allmet)
    summary=dict(created_utc=datetime.now(timezone.utc).isoformat(),total_pairs=1528,source_sample_sha256=sha(SOURCE/'paper_sample.txt'),config=cfg,freeze=receipt,calibration=cal,provenance=provenance,cpp_reexecuted=cpp,selected_status_counts=dict(Counter(r['status'] for r in rows)),candidate_status_counts=dict(Counter(c['status'] for d in detail for c in d['candidates'])),both_missing_frames=sum(r['rgb_count']==r['thermal_count']==0 for r in rows),both_present_no_match_frames=sum(r['rgb_count']>0 and r['thermal_count']>0 and all(c['status']!='matched' for c in detail[r['pair_index']]['candidates']) for r in rows),legacy_resets=resets,metrics=allmet)
    examples=overlays(rows,detail,pred,a); summary['examples']=examples
    save('execution_summary.json',summary); make_report(summary)
    loaded=list(csv.DictReader((OUT/'fusion_results.csv').open(encoding='utf-8-sig'))); assert len(loaded)==1528 and [int(r['pair_index']) for r in loaded]==list(range(1528))
    for name in NAMES: assert metrics([int(r['UV_reference_label']) for r in loaded],[int(r[name+'_binary']) for r in loaded])==met[name]
    with zipfile.ZipFile(REPORT) as z: assert z.testzip() is None; embedded=sum(n.startswith('word/media/') for n in z.namelist())
    doc=Document(REPORT); text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for rr in t.rows for c in rr.cells]); assert all(x in text for x in ['1528','循环','未验证','Hungarian','温度','UV单独'])
    (OUT/'report_readback.txt').write_text(text,encoding='utf-8'); assert embedded==len(examples)
    assert all(sha(p)==h for p,h in protected.items()) and sha(OUT/'frozen_config.json')==receipt['sha256']
    verify=dict(csv_rows=1528,unique_pair_ids=1528,current_image_hashes_verified=3056,metrics_recomputed_from_csv=True,docx_zip_crc_ok=True,docx_readback=True,docx_tables=len(doc.tables),embedded_images=embedded,original_sources_and_old_reports_unchanged=True,frozen_config_unchanged=True,report_path=str(REPORT),report_sha256=sha(REPORT),csv_sha256=sha(OUT/'fusion_results.csv'),metrics_sha256=sha(OUT/'metrics.json'),protected_file_hashes=protected)
    save('deliverable_verification.json',verify)
    print(json.dumps(dict(metrics=met,coverage=coverage,selected_status=summary['selected_status_counts'],both_missing=summary['both_missing_frames'],report=str(REPORT),verified=True),ensure_ascii=False,indent=2))

def overlays(rows,detail,pred,a):
    folder=OUT/'examples'; folder.mkdir(exist_ok=True)
    choices=[]
    for condition in [lambda r:r['status']=='matched',lambda r:r['rgb_count']>1 or r['thermal_count']>1,lambda r:r['rgb_count']==r['thermal_count']==0,lambda r:r['status']=='unmatched_rgb',lambda r:r['status']=='unmatched_thermal']:
        available=[r for r in rows if condition(r) and r['pair_index'] not in choices]
        if available: choices.append(max(available,key=lambda r:r['polygon_iou'])['pair_index'])
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',16); result=[]
    for i in choices:
        image=Image.open(pred['rgb'][i]['image_path']).convert('RGB'); draw=ImageDraw.Draw(image)
        for j,d in enumerate(pred['rgb'][i]['detections']): draw.rectangle(d['bbox_xyxy'],outline='lime',width=2); draw.text((d['bbox_xyxy'][0],d['bbox_xyxy'][1]),f'R{j}:{d["score"]:.3f}',font=font,fill='lime')
        for j,d in enumerate(pred['thermal'][i]['detections']):
            corners=[tuple(p) for p in s.affine(d['bbox_xyxy'],a)]; draw.line(corners+[corners[0]],fill='magenta',width=2); draw.text(corners[0],f'T{j}:{d["score"]:.3f}',font=font,fill='magenta')
        image=image.resize((960,round(960*image.height/image.width))); canvas=Image.new('RGB',(960,image.height+100),(20,20,20)); canvas.paste(image,(0,100)); draw=ImageDraw.Draw(canvas)
        r=rows[i]; pairs='; '.join(f'R{x["rgb_index"]}/T{x["thermal_index"]}:IoU={x["iou"]:.3f}' for x in detail[i]['candidates'])
        draw.text((8,4),f'pair {i} UV={r["UV_input"]} selected={r["selected_candidate"]} r={r["reliability"]:.3f} {r["status"]}',font=font,fill='white')
        draw.text((8,32),pairs,font=font,fill='white'); draw.text((8,62),f'绿色RGB／紫色仿射热框；新方案 BetP={r["spatial_uv_betp"]:.4f}；非独立真值',font=font,fill='white')
        path=folder/f'pair_{i}.png'; canvas.save(path); result.append(dict(pair_index=i,path=str(path),sha256=sha(path)))
    return result

def make_report(summary):
    d=Document(); st=d.styles['Normal']; st.font.name='Microsoft YaHei'; st.font.size=Pt(9); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    section=d.sections[0]; section.left_margin=section.right_margin=Inches(.65)
    def p(text): d.add_paragraph(str(text))
    def h(text): d.add_heading(text,1)
    def f(v): return '未定义' if v is None else f'{100*v:.4f}%'
    def table(headers,records):
        t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
        for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
        for record in records:
            for c,v in zip(t.add_row().cells,record): c.text=str(v)
    met=summary['metrics']['methods']; cfg=summary['config']
    d.add_heading('空间可靠性折扣D-S融合：1528对最终真实对比',0)
    p('本次实际执行全部1528对，复用用户提供best.pt（RGB）和best-2.pt（热成像）的真实YOLO26n预测；没有重训、重新推理、测试集参数搜索或人工修改检测。原C++、旧预测及旧报告保持不变。生成UTC：'+summary['created_utc'])
    p('重要：用户第四列UV0无火／UV1有火是本报告参考标准，同时也是融合输入。以下结果是UV符合度、循环评价，不是独立场景真实准确率，不是定位AP。UV单独100%按定义成立；必须用不含UV消融区分空间模块与UV注入，不能将全部提升归因于IoU。')
    h('1. 全量真实结果与直接结论')
    table(['方法','TP','FP','FN','TN','Acc','P','R','F1'],[[NAMES[k],m['TP'],m['FP'],m['FN'],m['TN'],f(m['accuracy']),f(m['precision']),f(m['recall']),f(m['F1'])] for k,m in met.items()])
    new,plain,spatial=met['spatial_uv'],met['ordinary_ds'],met['spatial_ds']
    p(f'新方案与普通DS（无UV）Accuracy差={100*(new["accuracy"]-plain["accuracy"]):+.4f}个百分点；仅空间可靠性（无UV）相对普通DS差={100*(spatial["accuracy"]-plain["accuracy"]):+.4f}个百分点；新方案相对无空间＋瞬时UV差={100*(new["accuracy"]-met["ordinary_uv"]["accuracy"]):+.4f}个百分点。这是实际观测，不是独立性能改善证明。')
    p('实际结论：无UV时普通DS与IoU可靠性DS的BetP二分类结果完全相同；加入瞬时UV后，有无空间模块均为100% UV符合度，UV单独也是100%。因此本次没有观察到IoU模块带来的二分类增益，100%主要反映直接引入同一个UV参考信号，不能写成独立融合检测准确率100%。保留旧时序UV时，IoU折扣结果反而低于无空间对照，详见表。')
    p('可靠性折扣降低明确证据、增加未知是设计目的，不等于提高所有决策指标。统一mF≥0.5无UV对照会明显降低召回；0.8阈值的新无UV方案全1528帧弃权；0.95下新含UV方案仅少量帧覆盖。BetP全量结果不能替代高确信度覆盖率。报告同时呈现这些负面结果，不调参掩盖。')
    p('标签分布：'+json.dumps(summary['metrics']['label_counts'])+'；正例比例高，Accuracy不可单独解释。P零分母为未定义。所有1528帧纳入；没有把困难帧删除或弃权视为正确。')
    h('2. 配准、目标匹配与帧聚合')
    cal=summary['calibration']; p('热框四角按2×3 FULL affine直接变换到RGB：q=A[:,0:2]p+A[:,2]。矩阵='+json.dumps(cal['affine_matrix_2x3'])+'。不使用外包矩形IoU、不裁剪多边形、不施加人工偏移。IoU由cv2.intersectConvexConvex真实交面积／并面积计算。')
    p(f'配准来源7对点，RANSAC内点{cal["inlier_count"]}/7，fit RMSE={cal["fit_rmse_px"]:.4f}px，留一RMSE={cal["leave_one_out_rmse_px"]:.4f}px，最大={cal["leave_one_out_max_px"]:.4f}px。候选未验证，点覆盖狭窄、视差／全视场／跨场景有效性未知。')
    p('全部保留目标构成IoU矩阵。Hungarian最大总IoU一对一匹配，dummy允许未匹配，预定最小门限0.01；无重复RGB或thermal索引。低于门限不强配，对两侧剩余目标各建单模态候选。门限只决定匹配，不把低IoU判无火。')
    p('每个候选按实际对应检测分数建立BPA。用普通DS的mF最大候选作为全帧代表，平手优先较高最小在场分数、再候选顺序；所有新方案消融固定使用同一候选，不根据UV选择目标。多框之间不连续DS累加，避免相关框数导致置信度膨胀。未匹配目标仍参与候选，保守可靠性0.5；缺失模态vacuous(0,0,1)。两模态全缺失保持全未知，而非无火证据。')
    p('选中状态：'+json.dumps(summary['selected_status_counts'])+'；所有候选状态：'+json.dumps(summary['candidate_status_counts'])+'；双缺失帧='+str(summary['both_missing_frames'])+'；双在场但无门限内匹配帧='+str(summary['both_present_no_match_frames'])+'。完整索引、候选、分数、质量与BPA见geometry_evidence.json。')
    h('3. 冻结公式、温度单位限制与透明算法变更')
    p('视觉基础BPA：(mF,mN,mU)=(0.8c,0.8(1-c),0.2)。此处c是检测fire置信度，不是经校准后验；互补无火项只适用于实际在场检测，缺框不是c=0无火证据。新公式与原RGB可靠性0.3／窗口平滑／温度饱和公式不同，不声称保持原C++不变。')
    p('匹配可靠性r=0.5+0.5s，s是实际polygon IoU；未匹配在场r=0.5。折扣mF′=r*mF，mN′=r*mN，mU′=1-r+r*mU；低IoU把明确证据移入未知，不自动增加无火比例。经典DS：K=aF*bN+aN*bF；Fsupport=aF*bF+aF*bU+aU*bF；Nsupport=aN*bN+aN*bU+aU*bN；Usupport=aU*bU；除以支持总和。总冲突回退vacuous并标记。先RGB／thermal经典DS，再UV经典DS，UV无空间框、绝不重复IoU。')
    p('瞬时UV新BPA：UV1=(0.9,0,0.1)，UV0=(0,0.9,0.1)，可靠性0.9预先固定。这是显式UV算法变更，不是原q=0.001递推。原时序UV消融保留原α0.01／β0.1、λ0.7、原startup、文件顺序与120s重置，反映原UV缓慢启动问题，不能用新UV收益冒充IoU收益。')
    p('第三列Tmax原样保存。温度单位/传感器标定未确认，主方案温度可靠性中性1：不会私自把50当摄氏阈值。已实现有单位、经外部给定center/scale才能启用的有界软门控rT=0.5+0.5*sigmoid((T-center)/scale)，当前未启用。因此主方案是空间RGB／热置信度＋UV，并未证明物理温度门控收益。另列“原温度热BPA”受控消融：重用原PC=1-exp(-10c)、PT=clamp(1-exp(-0.1(T-50)),0,1)、窗口10稳定性exp(-std)、λ0.7时序热BPA，用相同选中热分数与候选，热缺失仍未知；仅为历史数值规则参考，非单位已验证。')
    p('配置冻结于全量指标运行前，无训练/参数搜索。冻结UTC='+summary['freeze']['utc']+'；SHA256='+summary['freeze']['sha256']+'。参数全文：'+json.dumps(cfg,ensure_ascii=False))
    h('4. 二分类口径、同阈值审查与弃权')
    p('新方法二分类采用pignistic P(F)=mF+0.5*mU>0.5；精确0.5强制nofire供全量表，同时单独报告tie／未知。原C++行保留旧发布mF≥0.5，以免把历史指标偷偷改掉。下面两表对所有融合方法分别使用统一BetP与统一mF≥0.5；这样可区分算法与决策口径改变。无框全未知BetP=0.5，不可声称有证据支持正确无火。')
    for key,title in [('common_betp_metrics','全部融合统一BetP>0.5'),('common_mf05_metrics','全部融合统一mF≥0.5')]:
        p(title); table(['方法','Acc','P','R','F1'],[[NAMES[k],f(v['accuracy']),f(v['precision']),f(v['recall']),f(v['F1'])] for k,v in summary['metrics'][key].items()])
    p('弃权：mF≥0.8 fire；否则mN≥0.8 nofire；其他uncertain。覆盖率=(fire+nofire)/1528，弃权不是正确；另报0.95覆盖。覆盖子集完整TP/FP/FN/TN及P/R/F1见metrics.json。')
    table(['方法','F/.8','N/.8','弃权/.8','覆盖/.8','覆盖/.95','全未知','tie'],[[NAMES[k],v['08']['fire'],v['08']['nofire'],v['08']['uncertain'],f(v['08']['coverage']),f(v['095']['coverage']),v['vacuous_count'],v['tie_count']] for k,v in summary['metrics']['coverage'].items()])
    h('5. 真实模型、哈希与原C++控制')
    for mode,pr in summary['provenance'].items():
        md=pr['metadata']; p(mode+'：'+md['architecture']+'；checkpoint '+md['checkpoint']+'；SHA256 '+md['checkpoint_sha256']+'；预测SHA256 '+pr['sha256'])
        p('图像数1528，框数'+str(pr['boxes'])+'，无框帧'+str(pr['missing_frames'])+'，尺寸分布'+json.dumps(pr['dimensions'])+'；class0严格score>0.5；原预测imgsz640／iou0.7／max_det300／halfFalse／rectFalse；训练配置imgsz320。非PPYOLOv2，模型训练／测试重叠未审计。')
    p('重新核验3056源图像SHA256、文件名、行序、唯一pair_id、尺寸与权重SHA256。只读原paper_sample SHA256='+summary['source_sample_sha256']+'。')
    p('原C++动态与固定参考exe均再次运行全部1528行既有真实输入，exit0、与缓存最大误差<1e-10：'+json.dumps(summary['cpp_reexecuted'])+'。旧动态参考使用旧最高RGB锚／最高IoU热框及旧BPA；新方案一对一／选择策略／BPA／UV／决策均已披露，历史对比不是单因素因果实验。')
    h('6. 真实叠加样例')
    p('按匹配、多目标、全缺失及单模态选择真实样例，绿色RGB框、紫色仿射thermal四边形。显示全部候选与最终选中索引；不是人工场景真值，不能代表随机性能样本。')
    for ex in summary['examples']: p('pair '+str(ex['pair_index'])); d.add_picture(ex['path'],width=Inches(6.5))
    h('7. 可交付结果与未完成的独立验证')
    p('新目录fusion/spatial_reliability_ds：可重跑Python、冻结配置/哈希回执、TDD测试/RED日志、全1528 CSV/JSON、几何及证据、指标/执行摘要、样例、DOCX回读与核验清单。未覆盖：独立盲评场景标签、全视场跨场景配准验证、温度单位/标定、模型训练重叠审计。无独立验证时，不宣称真实火灾泛化性能改善。')
    d.save(REPORT)

if __name__=='__main__': run()
