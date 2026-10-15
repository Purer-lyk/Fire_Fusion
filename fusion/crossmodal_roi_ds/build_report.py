import json,hashlib,csv,zipfile,collections
from pathlib import Path
import numpy as np,cv2
from scipy.stats import beta
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import core
OUT=Path(__file__).parent; ROOT=OUT.parents[1]
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def load(n):return json.loads((OUT/n).read_text(encoding='utf-8'))
def build():
 rows=load('full_results.json'); metrics=load('full_metrics.json'); execution=load('full_execution.json'); config=load('frozen_config.json'); receipt=load('full_receipt.json'); a=np.array(json.loads((ROOT/'fusion/supplied_models_affine_iou/candidate_affine_calibration.json').read_text())['affine_matrix_2x3'])
 assert len(rows)==1528 and {r['pair_index'] for r in rows}==set(range(1528))
 truth=[r['uv_reference'] for r in rows]
 for key,field in [('primary_vision_only','visual_fire'),('optional_uv_gate','uv_fire'),('before_roi_vision','pre_roi_visual_fire')]:assert core.metrics(truth,[r[field] for r in rows])==metrics['methods'][key]
 assert sum(t['status']=='inferred' for r in rows for t in r['roi_attempts'])==execution['roi_actual_forward_images']
 assert all(sha(p)==h for p,h in receipt['protected_hashes'].items())
 audit=load('full_image_audit.json'); assert len(audit)==3056
 for record in audit: assert sha(record['path'])==record['sha256']
 selected=[]
 predicates=[lambda r:r['best'] and any(r[m+'_detections'][r['best'][m+'_index']]['provenance']=='roi_actual_detection' for m in ['rgb','thermal']),lambda r:r['visual_fire'] and r['matched_pairs'],lambda r:r['reason']=='roi_no_response' and r['roi_attempts']]
 for pred in predicates:
  found=next(r for r in rows if pred(r) and r['pair_index'] not in {s['pair_index'] for s in selected});selected.append(found)
 example_paths=[]; (OUT/'examples').mkdir(exist_ok=True)
 for r in selected:
  rgb=cv2.imread(r['rgb_path']); thermal=cv2.imread(r['thermal_path']); h=max(rgb.shape[0],thermal.shape[0]); canvas=np.zeros((h+80,rgb.shape[1]+thermal.shape[1],3),np.uint8); canvas[80:80+rgb.shape[0],:rgb.shape[1]]=rgb; canvas[80:80+thermal.shape[0],rgb.shape[1]:]=thermal; off=rgb.shape[1]
  for m,dx in [('rgb',0),('thermal',off)]:
   for j,b in enumerate(r[m+'_detections']):
    x1,y1,x2,y2=map(int,b['bbox']); color=(0,255,0) if b['provenance']=='roi_actual_detection' else (0,200,255);cv2.rectangle(canvas,(dx+x1,80+y1),(dx+x2,80+y2),color,1);cv2.putText(canvas,f'{m[0]}{j} {b["score"]:.2f}',(dx+x1,80+y1-2),cv2.FONT_HERSHEY_SIMPLEX,.3,color,1)
  for j,b in enumerate(r['thermal_detections']):
   p=core.transform(core.corners(b['bbox']),a); p[:,1]+=80;cv2.polylines(canvas,[p.astype(np.int32)],True,(255,0,255),1)
  for attempt in r['roi_attempts']:
   x1,y1,x2,y2=attempt['crop_xyxy_original'];dx=0 if attempt['review_modality']=='rgb' else off;cv2.rectangle(canvas,(dx+x1,80+y1),(dx+x2,80+y2),(255,255,0),1)
  cv2.putText(canvas,f'pair {r["pair_index"]} UV={r["uv_reference"]} mF={r["mf"]:.4f} visual={r["visual_fire"]}',(8,24),cv2.FONT_HERSHEY_SIMPLEX,.48,(255,255,255),1)
  cv2.putText(canvas,'left RGB / right thermal; magenta=thermal->RGB polygon',(8,45),cv2.FONT_HERSHEY_SIMPLEX,.4,(255,255,255),1)
  cv2.putText(canvas,'green=real ROI detection, cyan=crop, yellow=full detection',(8,66),cv2.FONT_HERSHEY_SIMPLEX,.4,(255,255,255),1)
  path=OUT/'examples'/f'pair_{r["pair_index"]}.png';assert cv2.imwrite(str(path),canvas);example_paths.append(path)
 doc=Document();style=doc.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(10);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 doc.add_heading('跨模态 ROI 复核与 D-S 融合：1528 对真实验证报告',0)
 doc.add_paragraph('结论：固定名义参数下，主视觉算法未达到全部目标。FPR 为 0%，但 FNR 为 88.69814%，高于热成像基线 47.63948%。ROI 产生真实复核框，但在同样的去重和匹配协议下未增加最终真阳性，无法抵消严格空间一致性和固定 mF 决策的漏报。不得据此宣称已达到目标或保证部署效果。')
 doc.add_heading('1. 数据、参考标准与执行证据',1)
 doc.add_paragraph('数据仅来自 C:/lyk/paper_data/paper_data_/paper_data/paper_sample.txt。1528 个唯一配对、3056 张输入图片逐一验证原始 SHA256 与尺寸。第四列 UV=1 有 1398 对、UV=0 有 130 对；这些统计是 UV 参考一致性，不是独立人工标注的场景真实性评估。第三列温度仅保留原值，单位未知，未启用物理温度门控。')
 doc.add_paragraph(f'用户原始权重只读使用 best.pt（RGB）及 best-2.pt（热成像）；实测架构 YOLO26n，不替换为 PPYOLOv2、不重新训练。Torch {execution["torch"]}，Ultralytics {execution["ultralytics"]}；实际设备 {execution["actual_devices"]}，{execution["gpu_name"]}，CPU 回退 {execution["fallbacks"]}。完整执行耗时 {execution["elapsed_seconds"]:.3f} 秒。')
 doc.add_paragraph(f'真实前向：全图 {execution["full_frame_inferences"]} 次，ROI {execution["roi_actual_forward_images"]} 次，合计 {sum(execution["model_input_counts"].values())} 个模型输入。ROI 尝试 {execution["roi_attempts"]} 次，产生 {execution["roi_real_detections"]} 个真实火焰检测输出（含跨裁剪重复，非独立目标数）。GPU 与单次预处理/推理/后处理耗时均写入 full_execution.json。原权重、历史预测、仿射标定和数据清单哈希前后不变。')
 doc.add_heading('2. 预冻结参数与实现范围',1)
 doc.add_paragraph(f'全量测试前冻结配置 SHA256：{execution["config_sha256"]}；启动时刻 UTC：{receipt["started_utc"]}。不使用测试 UV 标签优化任何阈值或挑选参数。')
 table=doc.add_table(rows=1,cols=2);table.style='Table Grid';table.rows[0].cells[0].text='参数';table.rows[0].cells[1].text='实际值'
 for k,v in config.items():cells=table.add_row().cells;cells[0].text=k;cells[1].text=str(v)
 doc.add_paragraph('置信分数校准仅提供可调用接口，本次为 identity：没有拟合独立校准器，因此不能称为“概率已校准”。q_A=1 仅为名义图像质量，无独立质量估计。仿射矩阵为既有 FULL thermal→RGB 候选标定，LOO RMSE 5.7718px，独立场景与全视场未验证。时间平滑关闭；未启用温度、重训练或其他未实际执行的特征。')
 doc.add_heading('3. 算法和严格坐标绑定',1)
 steps=['逐一按 paper_sample 的明确 pair_index/文件名读取 OpenCV BGR 数组；对列表预测逐项 zip 原数组与结果，检查 orig_shape，绝不依赖 Ultralytics 数组列表的虚拟文件名。两模型独立以 conf=0.1、imgsz=640、max_det=300 推理。','热成像真实框四角通过 2×3 FULL 仿射映射到 RGB；RGB 框保持原像素坐标。凸多边形相交面积计算 IoU，含 dummy 节点的 Hungarian 最大总 IoU 一对一匹配，硬门限 s≥0.3。','初始未匹配且置信分数≥0.25 的候选，按各模态最高分最多 3 个触发对方原始图像裁剪；RGB→thermal 用逆仿射，thermal→RGB 用正仿射。投影四角外包络每边扩展尺寸的 25% 加 10px 注册裕量，向外取整并裁至原图。每帧最多 6 个 ROI，不递归触发。','ROI 仅用于定义待复核区域，不作为检测证据。对裁剪原图实际再次运行对方 YOLO；真实输出框加 crop 左上角偏移回原图，再按同一仿射进行空间匹配。每模态原坐标 IoU>0.7 的重复框按置信分数保留一个。未匹配或无真实输出均不得报警。','每个真实匹配对：c 为未校准 detector confidence，u=0.1；基础 BPA=(0.9c,0.9(1−c),0.1)。r=0.5+0.5s，对称可靠性折扣 m=(r·mF,r·mN,1−r+r·mU)。','经典 D-S：K=mRF·mTN+mRN·mTF；mF=(mRF·mTF+mRF·mTU+mRU·mTF)/(1−K)，mN 对称，mU=mRU·mTU/(1−K)。记录冲突 K；奇异冲突退为全未知。不额外声称经过独立验证的冲突自适应机制。','每帧只选择最高视觉 mF 的实际匹配对，不累计同源多框，不以 UV 挑选框。主判定必须真实双模态匹配且 mF>0.5。无匹配视为 NoFire；这包含可能漏报而非安全证明。','可选 UV 只做乘性门控 mF_UV=mF·[(1−λ)+λUV]，λ=0.5。UV=0 将分数减半，UV=1 保持，不能复活缺失证据或凭 UV 合成检测。该支路把 UV 同时用作输入和参考，存在循环评价。']
 for s in steps:doc.add_paragraph(s,style='List Number')
 doc.add_paragraph('实际热成像→RGB 仿射矩阵：\n'+np.array2string(a,precision=10))
 doc.add_heading('4. 实际混淆矩阵与目标判定',1)
 table=doc.add_table(rows=1,cols=10);table.style='Table Grid';headers=['方法','TP','FP','FN','TN','FPR%','FNR%','Precision','Recall','F1']
 for c,h in zip(table.rows[0].cells,headers):c.text=h
 labels={'rgb_historical_baseline':'历史RGB conf0.5','thermal_historical_baseline':'历史热成像 conf0.5','before_roi_vision':'ROI前视觉DS','primary_vision_only':'主视觉ROI+DS','optional_uv_gate':'可选UV乘性门控'}
 for k,label in labels.items():
  m=metrics['methods'][k];vals=[label,*[str(m[x]) for x in ['TP','FP','FN','TN']],f'{100*m["FPR"]:.5f}',f'{100*m["FNR"]:.5f}',f'{m["precision"]:.6f}',f'{m["recall"]:.6f}',f'{m["F1"]:.6f}'];cells=table.add_row().cells
  for c,v in zip(cells,vals):c.text=v
 primary=metrics['methods']['primary_vision_only'];upper=float(beta.ppf(.975,1,130));doc.add_paragraph(f'主算法 meets_all_target={primary["meets_all_target"]}；可选 UV 门控结果与主算法相同，不依赖 UV 才取得 0 FP。相对 RGB 基线观察 FPR 从 96.92308% 降至 0%，但主算法召回率仅 11.30186%，不能以低 FPR 掩盖漏报。0/130 FP 的双侧 95% Clopper–Pearson FPR 区间约 [0, {100*upper:.4f}%]（独立二项假设；时间相关样本会削弱解释），不是总体 FPR 必为零的保证。')
 doc.add_heading('5. 漏报与 ROI 收益诊断',1)
 counts=collections.Counter({k:0 for k in ['both_no_candidates','candidate_weak_or_missing','roi_no_response','iou_failed','ds_low_score']});counts.update(metrics['miss_diagnostics'])
 for k,v in counts.items():doc.add_paragraph(f'{k}: {v} 个 UV 阳性漏报。')
 eligible=sum(bool(r['matched_pairs']) and bool(r['uv_reference']) for r in rows);rescued=[r['pair_index'] for r in rows if r['uv_reference'] and r['visual_fire'] and not r['pre_roi_visual_fire']];lost=[r['pair_index'] for r in rows if r['uv_reference'] and not r['visual_fire'] and r['pre_roi_visual_fire']]
 doc.add_paragraph(f'最终 UV 阳性空间匹配覆盖 {eligible}/1398；即使只放宽 DS 分数而保持此冻结匹配结果，召回也受该覆盖限制。ROI 后新检出 {len(rescued)} 个，ROI 前阳性变阴性 {len(lost)} 个；净变化 {len(rescued)-len(lost)}。rescued IDs={rescued}，lost IDs={lost}。低候选门限改善检索覆盖不意味着低原始置信分数的双模态 BPA 能超过 mF=0.5；本次主要障碍为 DS 低分。注册候选未经独立验证，IoU 失败不能都归为检测器错误。')
 doc.add_heading('6. 三个真实复核/匹配案例',1)
 for r,p in zip(selected,example_paths):
  doc.add_heading(f'Pair {r["pair_index"]} / {r["reason"]}',2);doc.add_paragraph(f'{r["rgb_name"]}\n{r["thermal_name"]}\nUV参考={r["uv_reference"]}，mF={r["mf"]:.6f}，ROI前={r["pre_roi_visual_fire"]}，ROI后={r["visual_fire"]}。');doc.add_picture(str(p),width=Inches(6))
  if r['best']:
   b=r['best']; rb=r['rgb_detections'][b['rgb_index']];tb=r['thermal_detections'][b['thermal_index']];doc.add_paragraph(f'实际选择：RGB score={rb["score"]:.6f} {rb["provenance"]} bbox={rb["bbox"]}；thermal score={tb["score"]:.6f} {tb["provenance"]} bbox={tb["bbox"]}；polygonIoU={b["iou"]:.6f}，K={b["conflict"]:.6f}。')
  for t in r['roi_attempts']:doc.add_paragraph(f'触发 {t["trigger_modality"]}[{t["trigger_index"]}] score={t["trigger_score"]:.4f} → {t["review_modality"]}，原图crop={t["crop_xyxy_original"]}，状态={t["status"]}，真实响应数={len(t["detections"])}。')
 doc.add_heading('7. 复现、验证与交付',1)
 for name in ['test_core.py','test_smoke.py','run_experiment.py','test_report.py']:doc.add_paragraph('C:/lyk/fire/qualityDetection/.venv/Scripts/python.exe '+str(OUT/name))
 doc.add_paragraph('运行 test_smoke 和 run_experiment 将再次进行真实推理并覆盖本实验自己的输出；不修改历史代码/结果。测试记录 red_geometry.log、red_fusion.log、red_smoke.log、red_report.log 与 green_tests.log、smoke.log、report_test.log。测试覆盖仿射逆映射、四角、ROI偏移与裁剪、真实数组尺寸绑定、一对一匹配、缺失证据硬门、UV不能复活证据、DS质量和已知混淆矩阵。')
 for name in ['frozen_config.json','core.py','run_experiment.py','full_results.csv','full_results.json','full_metrics.json','full_execution.json','full_image_audit.json','full_receipt.json','full_model_metadata.json','full_inference.log']:doc.add_paragraph(f'{OUT/name}\nSHA256={sha(OUT/name)}')
 doc.add_paragraph('后续应独立场景验证配准，并使用与本测试集隔离的标注验证集拟合概率校准及阈值。不能根据本次1528 UV标签调参后仍称其为独立测试，也不应为达标取消用户要求的真实双模态证据门。')
 path=ROOT/'reports/跨模态ROI复核_DS_1528验证报告.docx';doc.save(path); reread=Document(path);text='\n'.join(p.text for p in reread.paragraphs)+'\n'+'\n'.join(c.text for table in reread.tables for row in table.rows for c in row.cells)
 assert '88.69814' in text and '1240' in text and '未达到全部目标' in text
 (OUT/'report_readback.txt').write_text(text,encoding='utf-8')
 with zipfile.ZipFile(path) as z:
  assert z.testzip() is None; images=[n for n in z.namelist() if n.startswith('word/media/')];assert len(images)>=3
 verification={'unique_pairs':len(rows),'verified_image_hashes':len(audit),'embedded_images':len(images),'report_path':str(path),'report_sha256':sha(path),'docx_zip_ok':True,'report_readback_ok':True,'metric_recompute_ok':True,'protected_readback_unchanged':True,'meets_all_target':primary['meets_all_target'],'example_pair_ids':[r['pair_index'] for r in selected],'roi_rescued_ids':rescued,'roi_lost_ids':lost,'positive_spatial_match_coverage':eligible}
 (OUT/'verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(verification,ensure_ascii=False));return verification
if __name__=='__main__':build()
