from pathlib import Path
import json,csv,hashlib,zipfile
import numpy as np
import cv2
from scipy.optimize import linear_sum_assignment
from docx import Document
OUT=Path(__file__).resolve().parent;ROOT=OUT.parent.parent
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
load=lambda n:json.loads((OUT/n).read_text(encoding='utf-8'))
protocol=load('frozen_protocol.json');receipt=load('freeze_receipt.json');assert sha(OUT/'frozen_protocol.json')==receipt['protocol_sha256']
assert receipt['created_utc']<load('smoke_3pairs.json')['created_utc']<load('execution.json')['completed_utc']
for p,h in protocol['old_source_sha256'].items():assert sha(p)==h
assert sha(protocol['geometry']['source'])==protocol['geometry']['source_sha256']
for m,x in load('model_metadata.json').items():assert sha(x['source'])==x['sha256']
TEST=Path('C:/lyk/paper_data/paper_data_/paper_data');lines=[s.split() for s in (TEST/'paper_sample.txt').read_text().splitlines() if s.strip()]
y=[int(s[3]) for s in lines];assert len(y)==1528 and sum(y)==1398
pred={m:load(m+'_predictions.json') for m in ['rgb','thermal']};allhash=[]
for m,folder,col in [('rgb','bgr',0),('thermal','thermal',1)]:
 assert len(pred[m])==1528 and {r['pair_index'] for r in pred[m]}==set(range(1528))
 for i,r in enumerate(pred[m]):
  p=TEST/folder/lines[i][col];assert r['pair_index']==i and r['image_name']==p.name and Path(r['image_path'])==p;assert sha(p)==r['image_sha256'];allhash.append(sha(p))
  assert cv2.imread(str(p)).shape[:2]==tuple(r['orig_shape']);assert all(s[-2:]==[320,320] for s in r['model_input_shapes']);assert r['device']=='cuda:0'
assert len(set(allhash))==3056
# Independent implementation of geometry, Hungarian, selected BPA and counts.
def poly(b):
 x,y,X,Y=b;return np.array([[x,y],[X,y],[X,Y],[x,Y]],np.float32)
A=np.array(protocol['geometry']['affine_matrix_2x3'])
rows=list(csv.DictReader((OUT/'frame_predictions.csv').open(encoding='utf-8-sig')));stored=load('matched_evidence.json');assert len(rows)==1528
recomputed={};methods=['rgb_baseline','thermal_baseline','rgb_operating','thermal_operating','main_fire','optional_uv_fire'];out={m:[] for m in methods}
for i,row in enumerate(rows):
 assert int(row['pair_index'])==i and row['rgb_name']==lines[i][0] and row['thermal_name']==lines[i][1] and int(row['UV_reference'])==y[i]
 rb=pred['rgb'][i]['detections'];tb=pred['thermal'][i]['detections'];ri=[j for j,b in enumerate(rb) if b['score']>=.2];ti=[j for j,b in enumerate(tb) if b['score']>=.1]
 selected=[]
 if ri and ti:
  s=np.zeros((len(ri),len(ti)))
  for a,j in enumerate(ri):
   p=poly(rb[j]['bbox'])
   for b,k in enumerate(ti):
    q=(np.c_[poly(tb[k]['bbox']),np.ones(4)]@A.T).astype(np.float32);intersection,_=cv2.intersectConvexConvex(p,q);den=abs(cv2.contourArea(p))+abs(cv2.contourArea(q))-intersection;s[a,b]=intersection/den if den>0 else 0
  n,m=s.shape;w=np.zeros((n+m,n+m));w[:n,:m]=np.where(s>=.01,s,-1);rr,cc=linear_sum_assignment(-w)
  for a,b in zip(rr,cc):
   if a<n and b<m and s[a,b]>=.01:selected.append((ri[a],ti[b],s[a,b],1-(1-rb[ri[a]]['score'])*(1-tb[ti[b]]['score'])))
 assert len(selected)==int(row['matched_count'])==len(stored[i]['matches'])
 assert len({s[0] for s in selected})==len(selected)==len({s[1] for s in selected})
 best=max(selected,key=lambda s:s[3],default=None);mf=best[3] if best else 0
 assert abs(mf-float(row['mass_fire']))<1e-12
 if best:assert int(row['selected_rgb_index'])==best[0] and int(row['selected_thermal_index'])==best[1] and abs(float(row['selected_iou'])-best[2])<1e-12
 assert abs(sum(float(row[k]) for k in ['mass_fire','mass_nofire','mass_uncertain'])-1)<1e-12 and float(row['conflict'])==0
 decisions=dict(main_fire=int(mf>.3),optional_uv_fire=int(mf*(.5+.5*y[i])>.3),rgb_baseline=int(any(b['score']>.5 for b in rb)),thermal_baseline=int(any(b['score']>.5 for b in tb)),rgb_operating=int(any(b['score']>=.2 for b in rb)),thermal_operating=int(any(b['score']>=.1 for b in tb)))
 for m,p in decisions.items():assert p==int(row[m]);out[m].append(p)
for m,p in out.items():
 counts=dict(TP=sum(a==b==1 for a,b in zip(y,p)),FP=sum(a==0 and b==1 for a,b in zip(y,p)),FN=sum(a==1 and b==0 for a,b in zip(y,p)),TN=sum(a==b==0 for a,b in zip(y,p)))
 c=counts;v=dict(N=1528,**c,accuracy=(c['TP']+c['TN'])/1528,precision=c['TP']/(c['TP']+c['FP']),recall=c['TP']/1398,F1=2*c['TP']/(2*c['TP']+c['FP']+c['FN']),FPR=c['FP']/130,FNR=c['FN']/1398)
 for k,value in v.items():assert abs(value-load('metrics.json')['results'][m][k])<1e-12
 recomputed[m]=v
report=ROOT/'reports/320输入_空间一致性DS_论文式实验报告.docx'
with zipfile.ZipFile(report) as z:assert z.testzip() is None
read=Document(report);assert len(read.tables)==2 and len(read.inline_shapes)==3
assert len(read.tables[0].rows)==7 and len(read.tables[1].rows)==7
for i,m in enumerate(methods,1):
 assert [c.text for c in read.tables[0].rows[i].cells][1:]==[str(recomputed[m][k]) for k in ['TP','FP','FN','TN']]
text='\n'.join(p.text for p in read.paragraphs);assert '6 结论' in text and '不是独立测试' in text and '5.771805' in text
# Validate the prior320 operating-rule predictions genuinely reproduce, not merely quote old table.
OLD=ROOT/'newdirfusion/dataset_grounded_ds';prior={m:json.loads((OLD/name).read_text()) for m,name in [('rgb','rgb5_native320_predictions.json'),('thermal','thermal5_native320_predictions.json')]}
prior_operating_box_differences={}
for m,threshold in [('rgb',.2),('thermal',.1)]:
 differences=[]
 for i,(a,b) in enumerate(zip(pred[m],prior[m])):
  assert b['pair_index']==i and b['image_sha256']==a['image_sha256']
  aa=[x for x in a['detections'] if x['score']>=threshold];bb=[x for x in b['detections'] if x['score']>=threshold]
  if aa!=bb:differences.append(i)
 prior_operating_box_differences[m]=differences
oldrows=[r for r in csv.DictReader((OLD/'frame_predictions.csv').open(encoding='utf-8-sig')) if r['variant']=='primary_native320']
assert len(oldrows)==1528
prior_decision_differences=[i for i,(a,b) in enumerate(zip(rows,oldrows)) if a['main_fire']!=b['primary_fire']]
receipt=dict(success=True,independently_recomputed_metrics=recomputed,unique_pair_ids_each=1528,missing=0,unique_image_sha256=3056,all_geometry_and_candidate_selection_recomputed=True,actual_input_320_verified=True,old_source_hashes_unchanged=True,prior320_operating_box_difference_pair_ids=prior_operating_box_differences,prior320_main_decision_difference_pair_ids=prior_decision_differences,report_crc_ok=True,report_tables=2,report_images=3,report_sha256=sha(report))
(OUT/'validation.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
artifacts={str(p.relative_to(OUT)):sha(p) for p in OUT.rglob('*') if p.is_file() and p.suffix not in ['.pyc'] and p.name!='artifact_sha256.json'};artifacts[str(report)]=sha(report)
(OUT/'artifact_sha256.json').write_text(json.dumps(artifacts,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,indent=2))
