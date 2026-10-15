from pathlib import Path
import json,csv,hashlib,zipfile
import cv2,numpy as np
from docx import Document
from docx.shared import Pt,Inches,Cm
from docx.oxml.ns import qn
from fuse320 import corners,transform
OUT=Path(__file__).resolve().parent;ROOT=OUT.parent.parent
load=lambda n:json.loads((OUT/n).read_text(encoding='utf-8'))
results=load('metrics.json')['results'];diag=load('coverage_diagnostics.json');protocol=load('frozen_protocol.json');execution=load('execution.json');meta=load('model_metadata.json');audit=load('source_overlap_audit.json')['stats']
rows=list(csv.DictReader((OUT/'frame_predictions.csv').open(encoding='utf-8-sig')));pred={m:load(m+'_predictions.json') for m in ['rgb','thermal']};A=np.array(protocol['geometry']['affine_matrix_2x3'])
# One success, one false negative, one UV-reference negative. Never fabricate box ground truth.
ids=[next(i for i,r in enumerate(rows) if int(r['main_fire'])==1),next(i for i,r in enumerate(rows) if int(r['UV_reference'])==1 and int(r['main_fire'])==0),next(i for i,r in enumerate(rows) if int(r['UV_reference'])==0)]
(OUT/'examples').mkdir(exist_ok=True);example_records=[]
for i in ids:
 r=rows[i];im=cv2.imread(pred['rgb'][i]['image_path']);thermal=cv2.imread(pred['thermal'][i]['image_path']);original=im.copy();transformed=im.copy()
 for b in pred['rgb'][i]['detections']:
  if b['score']<.2:continue
  x,y,X,Y=map(int,b['bbox']);cv2.rectangle(original,(x,y),(X,Y),(0,230,0),1);cv2.putText(original,f"R {b['score']:.3f}",(x,max(10,y-2)),cv2.FONT_HERSHEY_SIMPLEX,.3,(0,230,0),1)
 for b in pred['thermal'][i]['detections']:
  if b['score']<.1:continue
  p=transform(corners(b['bbox']),A);cv2.polylines(transformed,[p.astype(np.int32)],True,(255,150,0),1);cv2.putText(transformed,f"T {b['score']:.3f}",tuple(p[0].astype(int)),cv2.FONT_HERSHEY_SIMPLEX,.3,(255,150,0),1)
 rr=int(r['selected_rgb_index']);tt=int(r['selected_thermal_index'])
 if rr>=0:
  cv2.polylines(transformed,[corners(pred['rgb'][i]['detections'][rr]['bbox']).astype(np.int32)],True,(0,230,0),2)
 canvas=np.full((385,960,3),245,np.uint8);canvas[45:365,:320]=original;canvas[45:365,320:640]=transformed
 # Thermal panel is padded, not stretched; retain real original image coordinates.
 h,w=thermal.shape[:2];canvas[45:45+h,640:640+w]=thermal
 for j,label in enumerate(['RGB real detections','RGB + affine thermal polygons','Original thermal image']):cv2.putText(canvas,label,(j*320+5,20),cv2.FONT_HERSHEY_SIMPLEX,.43,(0,0,0),1)
 label=f"pair={i} UV-reference={r['UV_reference']} main={r['main_fire']} selected IoU={float(r['selected_iou']):.4f} mF={float(r['mass_fire']):.4f} (NO box GT)"
 cv2.putText(canvas,label,(5,380),cv2.FONT_HERSHEY_SIMPLEX,.43,(0,0,0),1)
 p=OUT/'examples'/f'pair_{i}.png';assert not p.exists();assert cv2.imwrite(str(p),canvas)
 example_records.append(dict(pair_index=i,path=str(p),selected_rgb_index=rr,selected_thermal_index=tt,selected_iou=float(r['selected_iou']),mF=float(r['mass_fire']),reference=int(r['UV_reference'])))
(OUT/'example_receipts.json').write_text(json.dumps(example_records,indent=2),encoding='utf-8')
# No fabricated significance tests or independent labels.
main=results['main_fire'];rgb=results['rgb_baseline'];th=results['thermal_baseline'];pct=lambda x:f'{100*x:.4f}%'
labels={'rgb_baseline':'RGB单模态，score>0.5','thermal_baseline':'热成像单模态，score>0.5','rgb_operating':'RGB辅助阈值，score≥0.2','thermal_operating':'热成像辅助阈值，score≥0.1','main_fire':'空间IoU门控+D-S（主方法，无UV）','optional_uv_fire':'D-S后UV乘性门控（循环消融）'}
headers=['方法','TP','FP','FN','TN','准确率','精确率','召回率','F1','FPR','FNR']
table=[]
for m in labels:
 v=results[m];table.append([labels[m]]+[str(v[k]) for k in ['TP','FP','FN','TN']]+[pct(v[k]) for k in ['accuracy','precision','recall','F1','FPR','FNR']])
sections=[]
def sec(title,text):sections.append((title,text))
sec('摘要',f'为复核指定检查点在320输入下的帧级火情判别，本实验对1528组RGB—热成像配对样本重新执行完整GPU前向，使用既有七点全仿射变换及多边形IoU完成空间配对，在冻结阈值下组合正向火情简单支持证据。主方法不输入UV，第四列UV仅用作0/1参考标签。得到TP={main["TP"]}、FP={main["FP"]}、FN={main["FN"]}、TN={main["TN"]}，FPR={pct(main["FPR"])}、FNR={pct(main["FNR"])}、F1={pct(main["F1"])}。同次320前向的RGB与热成像score>0.5基线分别为FPR={pct(rgb["FPR"])}、{pct(th["FPR"])}，FNR={pct(rgb["FNR"])}、{pct(th["FNR"])}。结果为已见数据上的UV参考一致性复核，不能表述为独立泛化验证或动态可靠性折扣增益。')
sec('关键词','火情检测；RGB—热成像；全仿射配准；多边形IoU；Dempster–Shafer证据；误报率；漏报率')
sec('1 引言','不同模态在火情候选检出与干扰响应上存在差异。本文复核一个可重复的空间一致性门控及证据组合流程，重点给出误报率FPR与漏报率FNR，避免将假阳性率与1−精确率混淆。本文不重新训练、不更换权重、不启用ROI，也不以历史640结果替代同输入尺寸对照。由于阈值曾在已见配对数据上选取，本工作定位为受控复核而非新独立测试。')
sec('2 数据与模型',f'数据目录为C:/lyk/paper_data/paper_data_/paper_data，paper_sample.txt每行记录RGB文件、热成像文件、最高温度和UV参考值。共1528行、1398个UV=1与130个UV=0；两模态各1528个唯一源文件，重新读取并计算3056个唯一图像SHA256，无缺失或重复配对索引。原始RGB为320×320，热成像为256×192；原尺寸检测框用于几何映射，不在320字面尺寸下强行改坐标。最高温度不参与本次决策；UV不进入主方法。参考标签不是独立人工火情真值，更无目标框GT，因此只报告帧级指标，不报告AP/mAP。')
sec('2.1 指定权重与训练元数据','RGB直接加载C:/Users/Administrator/AppData/Local/hermes/attachments/best.pt；thermal直接加载同目录best-2.pt。二者均为Ultralytics YOLO26n，单类别class 0=fire/Fire，end2end=True，非PPYOLOv2。\nRGB SHA256='+meta['rgb']['sha256']+'\nthermal SHA256='+meta['thermal']['sha256']+'\n检查点SHA与此前RGB5/thermal5原始320实验一致，不使用RGB6、thermal3或后续安全训练模型。保存的RGB训练参数name=fireOrigin_yolo26n-5、data=cfgs\\bgr.yaml、epochs=100、imgsz=320；thermal为thermal_yolo26n-5、cfgs\\thermal.yaml、epochs=10、imgsz=320。本次训练次数为零。完整检查点yaml和train_args存于model_metadata.json。')
sec('2.2 数据重叠审计','对当前TorchDetection源划分重新计算图像哈希：'+json.dumps(audit,ensure_ascii=False)+'。RGB训练/验证与测试分别重叠298/74张，thermal分别重叠74/15张。此为当前源数据内容与测试的确切重复，不足以独立证明历史检查点实际使用了每一张图，但结合保存data配置构成明确风险。不得声称train/val/test独立或无泄漏，也不得把后续去重训练结果混入本表。')
sec('3 方法：空间约束与正向证据组合','本次沿用calibrated_spatial_ds/frozen_selected.json中的冻结参数：RGB≥0.2、thermal≥0.1、IoU≥0.01、rmin=1、最终mF>0.3。先按候选置信度筛选，再匹配；缺失或不配对时mF=0、mN=0、mU=1。该二分类工作点将未满足火情条件的帧归为0，不保留弃权类别，不能把mN=0称为模型已证明无火。')
sec('3.1 七点全仿射与真实多边形IoU','采用thermal→RGB的FULL affine矩阵A='+json.dumps(A.tolist())+'。热框四个角均以pR=A[pT.x,pT.y,1]^T映射，生成四边形；IoU=交集面积/(RGB框面积+映射多边形面积−交集面积)，交集通过cv2.intersectConvexConvex实际计算，不取AABB、不加人工像素偏移、不使用投影框替代检测。七点RANSAC内点6/7；全点拟合RMSE=2.479927 px，内点RMSE=0.996947 px，留一法RMSE=5.771805 px、最大误差9.142144 px。留一法是点集诊断，不是跨场景配准验证；该变换尚无独立全视场/多场景验证。')
sec('3.2 一对一配对与候选选择','对已筛选真实检测构造IoU矩阵，以dummy节点扩充为(n+m)×(n+m)矩阵，真实边IoU≥0.01时权重为IoU，否则为−1，dummy边为0；Hungarian最大总IoU完成一对一配对。遍历合格匹配对，以最高融合火情质量mF作为唯一帧级候选，选择过程独立于UV。不跨多个相关框累积组合证据。0.01为宽松硬门限，不能被描述为高精度空间一致性的证明。')
sec('3.3 D-S公式与决策','识别框架Θ={F,N}，每个检测的名义BPA为m(F)=c、m(N)=0、m(Θ)=1−c。冻结rmin=1意味着r=1，空间IoU仅决定硬门控、不执行动态可靠性折扣。两个正向简单支持由Dempster规则组合：K=0；mF=1−(1−cR)(1−cT)，mN=0，mU=(1−cR)(1−cT)。两质量之和为1；该运算是OR型正向支持增强，不是冲突建模或校准后验。只有真实双模态匹配且mF>0.3时输出火情。使用mF而非BetP；阈值严格大于，候选与IoU门限含等号。')
sec('3.4 UV仅作为可选循环消融','单独记录mF_uv=mF×[0.5+0.5×UV]，仍以mF_uv>0.3判断。这是D-S之后的乘性门控，不是额外空间证据，UV没有图像位置。由于UV同时充当参考标签，此消融具有循环评价风险，不作为主结果或独立提升证据。本次消融与主方法全部1528帧判决相同，不能宣称UV改善性能。')
sec('4 实验协议与评价定义',f'在任何smoke/full前向之前保存frozen_protocol.json及SHA回执。先真实3组样本smoke（6图），随后两模态各1528图完整前向，共3056图；设备cuda:0，{execution["gpu"]}，Torch {protocol["torch"]}、Ultralytics {protocol["ultralytics"]}。输入OpenCV BGR数组，batch=32、rect=False、固定imgsz=320；forward hook验证实际张量H=W=320，逐结果验证orig_img与源数组逐像素一致，并显式绑定源路径、pair_index与哈希，规避Ultralytics批路径丢失。half=False、augment=False、classes=[0]、max_det=300、iou参数0.7。YOLO26n端到端Nx6路径跳过常规NMS，因此0.7不作为实际NMS抑制阈值的作用依据。')
sec('4.1 与上次320协议的边界差异','上次原始320完整前向使用conf=0.01；本次按本轮要求使用conf=0.05，Ultralytics实际保留边界为score>0.05。后续操作门限仍为RGB≥0.2与thermal≥0.1，均高于新旧候选地板。这不是把旧0.01原始框缓存伪装为新前向，也不声称低分原始框列表逐位相同。本次实际新前向复现主方法1359/0/39/130计数。未执行ROI；协议、代码与来源哈希保存在独立新目录，旧结果未覆盖。')
sec('4.2 指标定义','以UV=1为参考阳性，TP为参考阳性且预测火情；FP为参考阴性但预测火情；FN为参考阳性但预测0；TN为参考阴性且预测0。Accuracy=(TP+TN)/1528，Precision=TP/(TP+FP)，Recall=TP/(TP+FN)，F1=2TP/(2TP+FP+FN)，FPR=FP/(FP+TN)=FP/130，FNR=FN/(FN+TP)=FN/1398。FPR不是FP/(TP+FP)。所有方法使用同一1528行同一次320预测。主单模态对照以任一真实框score>0.5判断；操作候选阈值RGB≥0.2与thermal≥0.1作为辅助对照完整披露，不挑选有利阈值。')
sec('5 结果与讨论','表1给出同次320输入的完整帧级结果。百分数保留四位小数，原始未舍入数值见metrics.json；TP/FP/FN/TN均为整数实测计数。')
sec('5.1 相对同尺寸单模态基线',f'主方法相对RGB>0.5基线，FPR降低{100*(rgb["FPR"]-main["FPR"]):.4f}个百分点，FNR降低{100*(rgb["FNR"]-main["FNR"]):.4f}个百分点；相对thermal>0.5基线，FPR降低{100*(th["FPR"]-main["FPR"]):.4f}个百分点，FNR降低{100*(th["FNR"]-main["FNR"]):.4f}个百分点。本次thermal>0.5并非零误报/零漏报，实际FP=11、FN=513；不存在用旧640替代320基线的问题。该对照共同改变了单模态置信度工作点、真实双模态约束与组合规则，不能将全部提升单独归因于D-S或空间配准。')
sec('5.2 辅助操作阈值与覆盖约束','RGB≥0.2辅助基线TP=1362、FP=70、FN=36；主方法漏报39，比该RGB工作点增加3个，换取减少70个假阳性。thermal≥0.1辅助基线TP=1392、FP=130、FN=6，提示低候选地板检出更多火情但也响应全部130个参考阴性。主方法不是对所有工作点都同时更优，不应作无条件优势论断。\n候选与几何覆盖：'+json.dumps(diag,ensure_ascii=False)+'。合格匹配帧1359，无合格匹配169，其中UV阳性39、阴性130；主方法39个漏报全部来自双模态候选/几何门控失败，所有合格匹配帧均通过最终质量阈值。selected IoU是每帧选定最高mF匹配的IoU，不是全体框的最大IoU。')
sec('5.3 零误报的不确定性与限制',f'主方法在130个UV参考阴性上FP=0，观察FPR=0%。二项独立假设下Clopper–Pearson双侧95%区间为[0,{100*main["FPR_exact95_upper"]:.4f}%]，不保证部署真实误报为零。数据带时间关联且类别不均衡，该区间只是条件性数值示例，不解决独立性问题。\n模型置信度未经校准；冻结配置源自此前已见1528配对数据的搜索；源训练/验证存在精确内容重叠；UV参考未获独立人工火情复核；配准缺少跨场景验证；没有目标级GT、跨场景外部集或单独空间/D-S因果消融。本结果仅支持这些模型与既有参数在该参考集上的受控复核。')
sec('5.4 真实检测可视化','图1—图3依次展示正确检出、参考阳性漏报和参考阴性示例。绿色为RGB模型真实框，蓝色为thermal真实检测经全仿射变换后的四边形，右侧为原始热图；图中列出真实选定IoU与mF，不画伪造GT。漏报/未匹配时selected IoU=0与mF=0是未选定状态，不代表已绘制全部候选间IoU都为零。')
sec('5.5 执行与可复现记录',f'本次始于{execution["started_utc"]}、完成于{execution["completed_utc"]}；总执行时间{execution["elapsed_seconds"]:.3f} s，RGB完整前向与读取记录{execution["timing"]["rgb"]["seconds"]:.3f} s，thermal为{execution["timing"]["thermal"]["seconds"]:.3f} s。计时含批预测、图像读取、哈希和记录处理，不是纯GPU kernel延时，不直接推导部署FPS。保存真实逐批时钟、模型输入shape与逐图speed。验证器重新读源标签、图像哈希、候选几何与融合，逐行重算混淆矩阵，Word包CRC和文档结构另行校验。')
sec('6 结论',f'新完整320前向确认：指定RGB/thermal权重与冻结空间门控+D-S规则在1528组UV参考样本上得到FPR={pct(main["FPR"])}、FNR={pct(main["FNR"])}、准确率={pct(main["accuracy"])}、精确率={pct(main["precision"])}、召回率={pct(main["recall"])}、F1={pct(main["F1"])}。相对同次score>0.5单模态基线达到更低观察FPR和FNR。该结论限定于已见数据上的复核；不是独立测试、不是动态折扣算法验证，也不构成零误报部署保证。后续需另行建立时序/场景隔离的外部标签集、验证几何并在独立校准数据上固定阈值。')
sec('附录A 可复现工件','目录：'+str(OUT)+'。主要文件：infer320.py、fuse320.py、test_320.py、frozen_protocol.json、freeze_receipt.json、model_metadata.json、smoke_3pairs.json、forward.log、rgb_predictions.json、thermal_predictions.json、frame_predictions.csv、matched_evidence.json、metrics.json、coverage_diagnostics.json、source_overlap_audit.json、execution.json、examples/、validation.json及artifact_sha256.json。无需从旧结果抄写任何本次指标。')
md=['# 320输入下基于空间一致性门控与D-S正向证据组合的火情判别复核报告','']
for title,text in sections:
 md.extend(['## '+title,'',text,''])
 if title=='5 结果与讨论':
  md.append('| '+' | '.join(headers)+' |');md.append('| '+' | '.join(['---']*len(headers))+' |')
  md.extend('| '+' | '.join(r)+' |' for r in table);md.append('')
 if title=='5.4 真实检测可视化':md.extend(f'![真实示例pair {e["pair_index"]}](examples/pair_{e["pair_index"]}.png)' for e in example_records)
(OUT/'paper.md').write_text('\n'.join(md),encoding='utf-8')
doc=Document();section=doc.sections[0];section.top_margin=Cm(2);section.bottom_margin=Cm(2);section.left_margin=Cm(1.6);section.right_margin=Cm(1.6)
style=doc.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(10.5);style.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
style.paragraph_format.space_after=Pt(6)
doc.add_heading('320输入下基于空间一致性门控与D-S正向证据组合的火情判别复核报告',0)
for title,text in sections:
 doc.add_heading(title,1)
 for para in text.split('\n'):doc.add_paragraph(para)
 if title=='5 结果与讨论':
  doc.add_paragraph('表1 同次完整320前向：帧级UV参考一致性结果')
  # Split count and rate tables to keep Chinese paper readable on portrait pages.
  for cols in [[0,1,2,3,4],[0,5,6,7,8,9,10]]:
   t=doc.add_table(rows=1,cols=len(cols));t.style='Table Grid'
   for j,k in enumerate(cols):t.rows[0].cells[j].text=headers[k]
   for row in table:
    cells=t.add_row().cells
    for j,k in enumerate(cols):cells[j].text=row[k]
   for row in t.rows:
    for cell in row.cells:
     for p in cell.paragraphs:
      for run in p.runs:run.font.size=Pt(8)
 if title=='5.4 真实检测可视化':
  for j,e in enumerate(example_records,1):doc.add_picture(e['path'],width=Inches(6.6));doc.add_paragraph(f'图{j} 真实配对索引{e["pair_index"]}；UV参考={e["reference"]}，IoU={e["selected_iou"]:.6f}，mF={e["mF"]:.6f}。')
report=ROOT/'reports/320输入_空间一致性DS_论文式实验报告.docx';report.parent.mkdir(exist_ok=True);assert not report.exists();doc.save(report)
read=Document(report);assert len(read.inline_shapes)==3 and len(read.tables)==2
with zipfile.ZipFile(report) as z:assert z.testzip() is None
(OUT/'report_readback.txt').write_text('\n'.join(p.text for p in read.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in read.tables for r in t.rows),encoding='utf-8')
print(report)
