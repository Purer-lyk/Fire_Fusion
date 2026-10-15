"""Controlled cache-only fusion rerun. prepare must precede evaluate; source artifacts read-only."""
from pathlib import Path
import sys,json,csv,hashlib,datetime,collections,zipfile,inspect
import numpy as np
import cv2
from scipy.optimize import linear_sum_assignment
OUT=Path(__file__).resolve().parent
OLD=OUT.parent/'adaptive_spatial_ds'
ROOT=OUT.parent.parent
sys.dont_write_bytecode=True
sys.path.insert(0,str(OLD))
import run_adaptive as old
import continuous_core as c
load=old.load
sha=old.sha
now=old.now
POLICIES={'fixed_source_gate001':(.001,False),'continuous_primary_source_gate001':(.001,True),'fixed_requested_gate01':(.01,False),'continuous_requested_gate01':(.01,True),'continuous_strict_confirmation_gate3':(.3,True)}

def save(name,obj):
 (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def csvsave(name,rows):
 with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def prepare():
 assert not (OUT/'protocol.json').exists(),'Protocol already frozen; never silently overwrite'
 protocol=dict(created_utc=now(),N=1528,reference_source=str(old.prior.old.SOURCE),source_sha256='7fa9997f14d65213063651f8448dc56344f797911185a2d6170147f1960bd483',rR=1.,rTmax=.1,sL=.1,sH=.3,threshold=.5,strict_decision='mF > .5; not BetP',primary_hard_gate=.001,requested_hard_gate=.01,secondary_strict_gate=.3,gate_discrepancy='Actual anchor_protocol.json AND adaptive_core.anchor_fuse both use .001; .01 in narrative is not the executed source gate. Preserve .001 as exact single-change primary; also execute requested .01 controlled pair and strict .3 separately.',g='clip((s-.1)/(.3-.1),0,1)',defaults_provenance='Predeclared historical transition defaults, not values supplied by user or optimized on current labels; all data previously exposed',matching='Dummy-enabled Hungarian maximum-sum nominal polygon IoU; one-to-one, eligible gate before assignment; matching before fusion; maximum final mF among all matched candidates; stable first tie',BPA='mR=(cR,0,1-cR);mT=(.1*g*cT,0,1-.1*g*cT);K=0;mF=1-(1-cR)*(1-.1*g*cT)',UV_runtime=False,inputs={'rgb':'rgb_all_full320.json','thermal':'thermal_all_predictions.json'},cache_source=str(OLD),fresh_forward=False,retraining=False,geometry=old.A.tolist(),models=old.P['models'],policies=POLICIES,prior_label_exposure=True,independent_test=False,search_trials=0,physical_temperature_input=False)
 protected={str(p):sha(p) for p in OLD.rglob('*') if p.is_file() and '__pycache__' not in p.parts}
 for p in (ROOT/'reports').glob('*.docx'):protected[str(p)]=sha(p)
 for v in old.P['models'].values():assert sha(v['source'])==v['sha256'];protected[v['source']]=sha(v['source'])
 geometry_source=old.P['geometry']['source'];protected[geometry_source]=sha(geometry_source)
 assert sha(old.prior.old.SOURCE)==protocol['source_sha256']
 protected[str(old.prior.old.SOURCE)]=sha(old.prior.old.SOURCE)
 save('protected_sources.json',protected);save('protocol.json',protocol)
 save('freeze_receipt.json',dict(utc=now(),protocol_sha256=sha(OUT/'protocol.json'),protected_sha256=sha(OUT/'protected_sources.json'),before_metrics=True,prior_exposed_data=True))
 print('FROZEN',sha(OUT/'protocol.json'),'actual source gate=.001; requested=.01 separately')

def match(rb,tb,gate):
 n,m=len(rb),len(tb)
 if not n or not m:return []
 s=np.array([[old.c.piou(old.c.corners(r['bbox']),old.c.project(old.c.corners(t['bbox']),old.A)) for t in tb] for r in rb])
 w=np.zeros((n+m,n+m));w[:n,:m]=np.where(s>=gate,s,-1.)
 rr,tt=linear_sum_assignment(-w)
 return [dict(rgb_index=int(i),thermal_index=int(j),cR=rb[i]['score'],cT=tb[j]['score'],s=float(s[i,j])) for i,j in zip(rr,tt) if i<n and j<m and s[i,j]>=gate]

def metric(y,p):
 tp=sum(a==1 and b==1 for a,b in zip(y,p));fp=sum(a==0 and b==1 for a,b in zip(y,p));fn=sum(a==1 and b==0 for a,b in zip(y,p));tn=sum(a==0 and b==0 for a,b in zip(y,p))
 return dict(N=len(y),TP=tp,FP=fp,FN=fn,TN=tn,FPR=fp/(fp+tn),FNR=fn/(tp+fn),precision=tp/(tp+fp) if tp+fp else 0.,recall=tp/(tp+fn),F1=2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else 0.,accuracy=(tp+tn)/len(y))

def evaluate():
 p=load(OUT/'protocol.json');assert sha(OUT/'protocol.json')==load(OUT/'freeze_receipt.json')['protocol_sha256']
 rows=old.prior.old.parse();R=load(OLD/'rgb_all_full320.json');T=load(OLD/'thermal_all_predictions.json');RM=load(OLD/'rgb_all_predictions.json');y=[int(r['UV_reference']) for r in rows];assert len(R)==len(T)==len(RM)==len(rows)==1528 and sum(y)==1288
 # Verify all actual image/weight metadata before reusing cached forwards.
 image_receipts=[]
 for mod,data in [('rgb',R),('thermal',T)]:
  for i,r in enumerate(data):
   assert r['pair_index']==i and Path(r['image_path']).name==rows[i]['rgb_name' if mod=='rgb' else 'thermal_name']
   h=sha(r['image_path']);assert h==r['image_sha256'];im=cv2.imread(r['image_path']);assert list(im.shape[:2])==r['orig_shape']
   for b in r['detections']:assert b['actual_detector'] and b['class_id']==0 and b['score']>.005
   if mod=='rgb':assert all(b['view']=='full320' and b['input_size']==320 for b in r['detections'])
   image_receipts.append(dict(modality=mod,pair_index=i,image_path=r['image_path'],sha256=h,shape=r['orig_shape']))
 for v in old.P['models'].values():assert sha(v['source'])==v['sha256']
 save('image_hash_verification.json',image_receipts)
 baseline={name:metric(y,[int(any(b['score']>.5 for b in r['detections'])) for r in data]) for name,data in [('rgb_original320',R),('thermal_multiview',T),('rgb_multiview_auxiliary',RM)]}
 historical=load(OLD/'anchor_evidence.json')
 evidence=[];pool=[];flat=[]
 for i,(r,t) in enumerate(zip(R,T)):
  gates={gate:match(r['detections'],t['detections'],gate) for gate in [.001,.01,.3]}
  o={key:c.select(gates[gate],continuous=cont,gate=gate) for key,(gate,cont) in POLICIES.items()}
  # Exact source function reproduction checked every row, not only aggregate counts.
  actual=old.c.anchor_fuse(r['detections'],t['detections'],old.A,rT=.1,threshold=.5,soft=False)
  fixed=o['fixed_source_gate001'];assert actual['fire']==fixed['fire']==historical[i]['fixed']['fire'];np.testing.assert_allclose(actual['masses'],[fixed['mF'],fixed['mN'],fixed['mU']],atol=1e-12)
  if actual['selected']:
   assert (actual['selected']['rgb_index'],actual['selected']['thermal_index'])==(fixed['rgb_index'],fixed['thermal_index'])
  for gate,pairs in gates.items():
   for pair in pairs:
    q=dict(pair_index=i,hard_gate=gate,**pair);q['g']=c.gain(pair['s']);q['actual_rT']=.1*q['g'];q['mF_fixed']=c.mass(pair['cR'],pair['cT'],pair['s'],False)[0];q['mF_continuous']=c.mass(pair['cR'],pair['cT'],pair['s'])[0];q['mN']=0.;q['mU_continuous']=1-q['mF_continuous'];pool.append(q)
  e=dict(pair_index=i,rgb_name=rows[i]['rgb_name'],thermal_name=rows[i]['thermal_name'],UV_reference=y[i],rgb_detections=r['detections'],thermal_detections=t['detections'],matches_by_gate={str(k):v for k,v in gates.items()},outputs=o);evidence.append(e)
  f=dict(pair_index=i,rgb_name=e['rgb_name'],thermal_name=e['thermal_name'],UV_reference=y[i],rgb_candidate_count=len(r['detections']),thermal_candidate_count=len(t['detections']),matched001=len(gates[.001]),matched01=len(gates[.01]),matched3=len(gates[.3]))
  for key,v in o.items():f.update({key+'_'+k:val for k,val in v.items()})
  flat.append(f)
 metrics={key:metric(y,[e['outputs'][key]['fire'] for e in evidence]) for key in POLICIES}
 assert {k:metrics['fixed_source_gate001'][k] for k in ['TP','FP','FN','TN']}==dict(TP=1265,FP=0,FN=23,TN=240)
 assert {k:baseline['rgb_original320'][k] for k in ['TP','FP','FN','TN']}==dict(TP=1262,FP=0,FN=26,TN=240)
 assert {k:baseline['thermal_multiview'][k] for k in ['TP','FP','FN','TN']}==dict(TP=1199,FP=180,FN=89,TN=60)
 transitions={}
 for key in POLICIES:
  if key=='fixed_source_gate001':continue
  a=[e['outputs']['fixed_source_gate001'] for e in evidence];b=[e['outputs'][key] for e in evidence]
  ids=lambda cond:[i for i,(u,v) in enumerate(zip(a,b)) if cond(i,u,v)]
  transitions[key]=dict(oldTP_newFN=ids(lambda i,u,v:y[i]==1 and u['fire']==1 and v['fire']==0),oldFP_newTN=ids(lambda i,u,v:y[i]==0 and u['fire']==1 and v['fire']==0),oldFN_newTP=ids(lambda i,u,v:y[i]==1 and u['fire']==0 and v['fire']==1),oldTN_newFP=ids(lambda i,u,v:y[i]==0 and u['fire']==0 and v['fire']==1),selected_candidate_changed=ids(lambda i,u,v:(u['rgb_index'],u['thermal_index'])!=(v['rgb_index'],v['thermal_index'])),mass_changed=ids(lambda i,u,v:abs(u['mF']-v['mF'])>1e-12))
 pc='continuous_primary_source_gate001';pairs001=[q for q in pool if q['hard_gate']==.001];selected=[e['outputs'][pc] for e in evidence if e['outputs'][pc]['rgb_index'] is not None]
 stats=dict(all_matched_candidates=len(pairs001),candidates_g_less_1=sum(q['g']<1 for q in pairs001),candidates_g_zero=sum(q['g']==0 for q in pairs001),selected_frames=len(selected),selected_g_less_1=sum(q['g']<1 for q in selected),selected_g_zero=sum(q['g']==0 for q in selected),firing_selected_g_less_1=sum(q['g']<1 and q['fire']==1 for q in selected),firing_selected_g_zero=sum(q['g']==0 and q['fire']==1 for q in selected))
 gains=[dict(pair_index=i,reference=y[i],outputs=evidence[i]['outputs']) for i in [727,747,879]]
 illustrations=[dict(s=s,g=c.gain(s),rT=.1*c.gain(s),cR=.49,cT=.8,mF_fixed=c.mass(.49,.8,s,False)[0],mF_continuous=c.mass(.49,.8,s)[0]) for s in [0.,.001,.01,.05,.1,.100001,.15,.2,.25,.299999,.3,.449,1.]]
 result=dict(created_utc=now(),N=len(y),positive=sum(y),negative=len(y)-sum(y),source_sha256=sha(old.prior.old.SOURCE),UV_runtime=False,fresh_forward=False,independent_test=False,baselines=baseline,results=metrics,transitions=transitions,affected=stats,known_fixed_gains=gains,illustrative_only=illustrations,targets={key:dict(FP0_FN_le25=m['FP']==0 and m['FN']<=25) for key,m in metrics.items()})
 save('frame_evidence.json',evidence);csvsave('frame_predictions.csv',flat);csvsave('all_matched_candidates.csv',pool);save('metrics.json',result);save('delta_cases.json',transitions);save('known_gains.json',gains);save('continuity_illustration.json',illustrations)
 print(json.dumps(dict(results=metrics,baselines=baseline,affected=stats,transitions=transitions),ensure_ascii=False))

def report():
 from docx import Document
 from docx.shared import Inches,Pt,Cm
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 m=load(OUT/'metrics.json');p=load(OUT/'protocol.json');d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(2.2);sec.left_margin=sec.right_margin=Cm(2.4)
 normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10.5);normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');normal.paragraph_format.line_spacing=1.3
 for h in ['Title','Heading 1','Heading 2']:d.styles[h].element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
 def para(s):d.add_paragraph(s)
 def table(headers,rows):
  t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
  for cell,txt in zip(t.rows[0].cells,headers):cell.text=str(txt)
  for row in rows:
   for cell,txt in zip(t.add_row().cells,row):cell.text=str(txt)
  return t
 d.add_heading('连续IoU可靠性 D-S受控复验报告',0)
 para('固定真实检测缓存｜全量1,528对｜仅修改热证据可靠性｜探索性配对比较')
 d.add_heading('摘要',1)
 primary=m['results']['continuous_primary_source_gate001'];fixed=m['results']['fixed_source_gate001'];strict=m['results']['continuous_strict_confirmation_gate3']
 para(f"本研究复用原始全图320 RGB与真实热图多视角检测输出，对1,528对样本重新进行名义仿射多边形IoU、全候选一对一匹配和D-S融合，不重新训练、不重新网络前向、不输入当前UV。历史固定法完整复现TP={fixed['TP']}、FP={fixed['FP']}、FN={fixed['FN']}、TN={fixed['TN']}。连续可靠性主对照得到TP={primary['TP']}、FP={primary['FP']}、FN={primary['FN']}、TN={primary['TN']}。严格确认附加对照得到TP={strict['TP']}、FP={strict['FP']}、FN={strict['FN']}、TN={strict['TN']}。所有结果均为已暴露数据上的描述性复验，不构成独立泛化认证。")
 d.add_heading('1 数据、缓存与受控范围',1)
 para('参考标签为当前paper_sample.txt第四列UV：阳性1,288、阴性240。UV只用于结果评估，不进入匹配、可靠性或最终预测。温度、文件名、pair_index和标签均不作为预测特征；pair_index仅用于审计关联。所有3,056张图像与两份检查点逐项SHA256核验后复用真实缓存。RGB采用full320原图输出，热图采用原已执行的full320、full640及四个tile320多视角输出及同一跨视角去重；不是仅320双模态新前向。')
 para('源代码发现：anchor_protocol.json与adaptive_core.anchor_fuse实际硬门槛均为0.001，非先前文字描述的0.01。本报告保留0.001作为源码同策略、单变量主复验；另列0.01固定/连续配对对照，以及0.3严格确认对照。三个硬门槛不混称同一主实验，也没有将sL/sH偷换成硬门槛。')
 para('全部参数在本次指标计算前写入protocol.json并生成冻结SHA256回执；sL=0.1、sH=0.3为历史转折区间的预先声明默认值，不是用户专属给定值，也未利用本次标签调优。该预声明不消除历史全量数据暴露事实。')
 table(['参数/策略','冻结值'],[['rR / rTmax','1 / 0.1'],['sL / sH','0.1 / 0.3'],['主对照硬门槛','IoU ≥ 0.001（源码原值）'],['文字要求门槛附加对照','IoU ≥ 0.01（相对源码属于额外门槛变动）'],['严格确认附加对照','IoU ≥ 0.3'],['最终决策','mF > 0.5（严格大于）'],['候选与匹配','原floor>0.005；原去重不变；dummy Hungarian最大总IoU'],['帧内汇总','所有合格一对一匹配中取最终mF最大；同分取原顺序首项'],['网络/训练/阈值搜索','无新前向／无训练／0次搜索']])
 d.add_heading('2 方法与连续性',1)
 para('将热图框四角以原2×3仿射矩阵映射到RGB坐标，使用凸多边形交并比s，不使用外接矩形IoU、透视除法或Smax。保持候选过滤、几何矩阵及最大总IoU的一对一Hungarian规则。先匹配、后融合；修改可靠性后对所有实际匹配重新计算并重新选取最终mF最大者，而非仅修改历史选中框。')
 para('g(s) = clip((s − 0.1)/(0.3 − 0.1), 0, 1)；rT(s) = 0.1 g(s)。')
 para('mR = (cR, 0, 1 − cR)；mT = (0.1 g(s)cT, 0, 1 − 0.1 g(s)cT)。')
 para('K = 0；mF = 1 − (1 − cR)(1 − 0.1 g(s)cT)；mN = 0；mU = (1 − cR)(1 − 0.1 g(s)cT)。')
 para('此为正向简单支持BPA的Dempster组合；没有负证据或冲突抑制，置信度不宣称为校准类别后验。不得以BetP阈值0.5替代mF>0.5，否则含义改变。缺任一真实证据或无合格匹配时输出空泛质量(0,0,1)且不触发。')
 para('在0.001/0.01硬门槛下，s≤0.1时g=0仅使热证据增量为零；如匹配RGB本身cR>0.5，仍可触发。这与“弱空间对应完全不能报火”不同，后者由单列的s≥0.3严格确认实验评估。s≥0.3时g=1，连续法与固定法该候选完全相等，例s=0.449也不改变。')
 table(['s（示例）','g(s)','rT(s)','固定mF','连续mF'],[[q['s'],f"{q['g']:.6f}",f"{q['rT']:.6f}",f"{q['mF_fixed']:.6f}",f"{q['mF_continuous']:.6f}"] for q in m['illustrative_only']])
 para('上表仅为cR=0.49、cT=0.8的公式演示，不是生成或替代实测检测数据。')
 d.add_heading('3 全量受控结果',1)
 labels={'rgb_original320':'RGB原图320','thermal_multiview':'热图多视角','rgb_multiview_auxiliary':'RGB多视角（辅助公平对照）','fixed_source_gate001':'原固定法，门槛0.001','continuous_primary_source_gate001':'连续法主对照，门槛0.001','fixed_requested_gate01':'固定法附加对照，门槛0.01','continuous_requested_gate01':'连续法附加对照，门槛0.01','continuous_strict_confirmation_gate3':'连续法严格确认，门槛0.3'}
 allmetrics={**m['baselines'],**m['results']}
 table(['方法','TP','FP','FN','TN','FPR %','FNR %'],[[labels[k],v['TP'],v['FP'],v['FN'],v['TN'],f"{100*v['FPR']:.4f}",f"{100*v['FNR']:.4f}"] for k,v in allmetrics.items()])
 table(['方法','Precision %','Recall %','F1 %','Accuracy %'],[[labels[k],*[f"{100*v[t]:.4f}" for t in ['precision','recall','F1','accuracy']]] for k,v in allmetrics.items()])
 para('FPR=FP/240；FNR=FN/1288；Precision=TP/(TP+FP)；Recall=TP/1288；F1=2TP/(2TP+FP+FN)；Accuracy=(TP+TN)/1528。上述指标均为帧级UV参考一致性，不是目标级AP，也不是独立人工火焰真值性能。RGB多视角辅助对照不替换主实验原始320 RGB输入。')
 d.add_heading('4 实际受影响候选与逐帧差异',1)
 para(f"主对照共有{m['affected']['all_matched_candidates']}个实际一对一匹配，其中{m['affected']['candidates_g_less_1']}个候选g<1；{m['affected']['selected_g_less_1']}帧的最终质量发生下降，但全部1,528帧的二值决策不变，未新增TP、未减少FP，也未新增FN。连续化改变证据强度，不应宣称本次提高了分类指标。0.01固定/连续配对对照的二值决策同样完全一致；其相对历史少4个TP来自额外门槛变化，不是连续折扣。")
 table(['审计项','实际数量'],[[k,v] for k,v in m['affected'].items()])
 for key,v in m['transitions'].items():
  para(labels[key]+'：'+ '；'.join(f'{name}={len(ids)}，IDs={ids}' for name,ids in v.items()))
 para('所有ID均为0起始pair_index；delta_cases.json列出完整变化列表，frame_predictions.csv包含全1,528行选中证据cR/cT/s/g/rT、质量与决策；all_matched_candidates.csv保存各门槛的全部实际一对一匹配及固定/连续质量。frame_evidence.json同时保存每帧全部原始候选框，以便核验非历史选中候选的重选。')
 table(['已知新增TP ID','cR','cT','IoU','g','固定mF','连续mF','前/后决策'],[[q['pair_index'],f"{q['outputs']['fixed_source_gate001']['cR']:.9f}",f"{q['outputs']['fixed_source_gate001']['cT']:.9f}",f"{q['outputs']['fixed_source_gate001']['s']:.9f}",f"{q['outputs']['continuous_primary_source_gate001']['g']:.6f}",f"{q['outputs']['fixed_source_gate001']['mF']:.9f}",f"{q['outputs']['continuous_primary_source_gate001']['mF']:.9f}",f"{q['outputs']['fixed_source_gate001']['fire']}/{q['outputs']['continuous_primary_source_gate001']['fire']}"] for q in m['known_fixed_gains']])
 para('727、747、879为历史固定法相对RGB原图320的三个新增TP；其真实IoU及前后质量如上，不以推测替代实算。')
 d.add_heading('5 有限目标与局限',1)
 for key,v in m['targets'].items():para(labels[key]+'：FP=0且FN≤25：'+('达到' if v['FP0_FN_le25'] else '未达到')+'（仅本次已暴露数据描述性目标）。')
 para('本实验目标是检验连续IoU折扣在固定输入和决策规则下的影响，不进行新的阈值搜索、校准优化、OOF选择或GPU重训。历史已重复查看全部数据，并存在训练/测试内容重叠风险；不能称为新的独立测试或认证。原名义仿射仅依赖七点标定，历史留一误差约5.77像素且视场覆盖有限，不足以证明所有场景的空间对应有效。热图以及部分RGB中的烧录检测框/Fire文字可能污染检测，真实网络输出可追溯不等于经人工确认的火焰对应。')
 para('正证据DS只会强化置信度。连续折扣可降低弱重叠的增量，但不能创造真实缺失证据，不能自动修复配准；零误报样本计数也不证明真实部署误报率为零。严格确认分支同时改变硬门槛，仅用于附加解释，不把其变化归因于单独的连续可靠性修改。')
 d.add_heading('6 可复现产物与核验',1)
 para('脚本：run_continuous.py（prepare/evaluate/report）、continuous_core.py。运行环境使用项目.venv/Scripts/python.exe；默认python缺少numpy，未安装或改动全局环境。协议回执、3,056图像哈希核验、独立验证及文档CRC回执均保存在本实验根目录。原缓存、源代码、权重、标签文件与旧DOCX均只读保护。')
 para('本次协议SHA256：'+sha(OUT/'protocol.json'))
 para('标签源SHA256：'+m['source_sha256'])
 path=ROOT/'reports/连续IoU可靠性_DS复验报告.docx';path.parent.mkdir(parents=True,exist_ok=True);d.save(path)
 print(str(path))

if __name__=='__main__':globals()[sys.argv[1]]()
