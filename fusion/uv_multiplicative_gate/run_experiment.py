"""Frozen literal UV multiplication over 1528 real cached pairs; no inference."""
from pathlib import Path
import sys,json,csv,hashlib,zipfile,math
from datetime import datetime,timezone
from collections import Counter
import numpy as np
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml.ns import qn
from multiplicative_core import apply
OUT=Path(__file__).resolve().parent
ROOT=OUT.parent.parent
REPORT=ROOT/'reports/UV乘性门控_DS_1528统计报告.docx'
SOURCE=Path('C:/lyk/paper_data/paper_data_/paper_data/paper_sample.txt')
REL=OUT.parent/'spatial_reliability_ds'
GATE=OUT.parent/'spatial_gate_ds'
PRED=OUT.parent/'supplied_models_uv_reference'
sys.path.insert(0,str(REL))
import spatial_core as s
LABELS={'A':'普通视觉DS（无空间硬门）','B':'原IoU可靠性DS（无硬门）','C':'严格空间门控DS（主方案）'}
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def load(p): return json.loads(Path(p).read_text(encoding='utf-8'))
def save(name,data): (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def writecsv(name,rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
def metric(y,p):
    tp=sum(a==b==1 for a,b in zip(y,p)); tn=sum(a==b==0 for a,b in zip(y,p))
    fp=sum(a==0 and b==1 for a,b in zip(y,p)); fn=sum(a==1 and b==0 for a,b in zip(y,p))
    recall=tp/(tp+fn); spec=tn/(tn+fp)
    return dict(N=len(y),TP=tp,FP=fp,FN=fn,TN=tn,accuracy=(tp+tn)/len(y),precision=tp/(tp+fp) if tp+fp else None,recall=recall,F1=2*tp/(2*tp+fp+fn),balanced_accuracy=(recall+spec)/2,specificity=spec,FPR=1-spec,FNR=1-recall,binary_decision_coverage=1.,predicted_fire_count=tp+fp)
def invariants():
    cases=0
    for lam in [0,.25,.5,.75,1]:
        for uv in [0,1]:
            for mass in [(0.,0.,1.),(.5,.2,.3),(.8,.1,.1),(1.,0.,0.)]:
                for matched,iou in [(False,0.),(True,.1),(True,.299999),(True,.3),(True,1.)]:
                    for hard in [False,True]:
                        r=apply(mass,uv,lam,hard_gate=hard,matched=matched,iou=iou)
                        assert r['mass'][0]<=mass[0] and abs(sum(r['mass'])-1)<1e-12
                        assert r['mass'][1]==mass[1] and r['fire']<=r['base_fire']
                        if uv==1 or lam==1: assert r['mass']==mass
                        if uv==0 and lam==0: assert r['mass'][0]==0
                        if mass[0]==.5: assert r['fire']==0
                        if hard and (not matched or iou<.3): assert r['fire']==0
                        if mass==(0.,0.,1.): assert r['fire']==0
                        cases+=1
    assert apply((.8,.1,.1),1,hard_gate=True,matched=True,iou=.3)['fire']==1
    return dict(cases=cases,lambda_endpoints=True,uv_endpoints=True,mass_conservation=True,no_score_increase=True,no_resurrection=True,strict_mF_boundary=True,inclusive_spatial_boundary=True,UV_cannot_override_spatial_or_both_missing=True)
def run():
    OUT.mkdir(exist_ok=True); REPORT.parent.mkdir(exist_ok=True)
    assert not REPORT.exists() or (OUT/'deliverable_verification.json').exists(), 'refuse overwriting unowned report'
    protected_paths=[]
    for folder in [REL,GATE,PRED,OUT.parent/'supplied_models_affine_iou',Path('C:/lyk/paper/bpa_test')]:
        protected_paths += [p for p in folder.rglob('*') if p.is_file() and '__pycache__' not in str(p)]
    protected_paths += [p for p in (ROOT/'reports').glob('*.docx') if p!=REPORT]+[SOURCE]
    protected={str(p):sha(p) for p in protected_paths}
    config=dict(lambda_nominal=.5,tau_F=.5,decision_operator='>',tau_L=.1,tau_H=.3,spatial_operator='>=',UV='paper_sample.txt fourth column binary',formula='mprimeF=mF*(lambda+(1-lambda)*guv)',BPA_conservation='mprimeN=mN; mprimeU=mU+(1-q)*mF',UV_is_third_DS_source=False,renormalize=False,threshold_source='pre-fixed configurable defaults, not user-specified numerics; no test tuning',lambda_not_historical_smoothing_07=True,sensitivity_lambda=[0,.25,.5,.75,1],candidate_selection={'A':'original spatial_reliability_ds ordinary mF maximum candidate','B':'same original candidate as A','C':'latest spatial_gate_ds maximum IoU, tie visual BetP/min confidence/index'},visual_mass_sources={'A':'spatial_reliability_ds ordinary_ds_*','B':'spatial_reliability_ds spatial_ds_*','C':'spatial_gate_ds raw_spatial_ds_*'},source_sha256={str(p):sha(p) for p in [SOURCE,REL/'fusion_results.csv',REL/'fusion_results.json',GATE/'fusion_results.csv',GATE/'fusion_results.json']})
    if (OUT/'frozen_config.json').exists(): assert load(OUT/'frozen_config.json')==config
    else: save('frozen_config.json',config)
    config_hash=sha(OUT/'frozen_config.json')
    if not (OUT/'config_freeze_receipt.json').exists(): save('config_freeze_receipt.json',dict(sha256=config_hash,frozen_utc=datetime.now(timezone.utc).isoformat(),before_metrics=True))
    assert load(OUT/'config_freeze_receipt.json')['sha256']==config_hash
    tests=invariants(); save('numerical_invariant_tests.json',tests)
    samples=[l.split() for l in SOURCE.read_text(encoding='utf-8-sig').splitlines()]
    assert len(samples)==1528 and all(len(x)==4 and x[3] in ('0','1') for x in samples)
    assert len({x[0] for x in samples})==len({x[1] for x in samples})==1528
    rel=load(REL/'fusion_results.json')['rows']; gate=load(GATE/'fusion_results.json')['rows']
    for folder,jrows in [(REL,rel),(GATE,gate)]:
        crows=list(csv.DictReader((folder/'fusion_results.csv').open(encoding='utf-8-sig')))
        assert len(crows)==len(jrows)==1528
        for cr,jr in zip(crows,jrows):
            for key,val in jr.items():
                if val is None: assert cr[key]==''
                elif isinstance(val,(float,int)): assert float(cr[key])==val
                else: assert cr[key]==str(val)
    provenance={}; verified_images=0
    for mode,pos in [('rgb',0),('thermal',1)]:
        path=PRED/(mode+'_predictions.json'); data=load(path)
        assert len(data['images'])==1528 and sha(data['metadata']['checkpoint'])==data['metadata']['checkpoint_sha256']
        for i,r in enumerate(data['images']):
            assert r['pair_index']==i and r['source_line']==i+1 and r['image_name']==samples[i][pos]
            assert sha(r['image_path'])==r['image_sha256']; verified_images+=1
            assert all(d['class_id']==0 and d['score']>.5 for d in r['detections'])
            assert len(r['detections'])==rel[i][mode+'_count']==gate[i][mode+'_count']
        provenance[mode]=dict(path=str(path),sha256=sha(path),architecture=data['metadata']['architecture'],checkpoint=data['metadata']['checkpoint'],checkpoint_sha256=data['metadata']['checkpoint_sha256'],rows=1528)
    rows=[]; y=[]
    for i,(sm,rr,gr) in enumerate(zip(samples,rel,gate)):
        uv=int(sm[3]); y.append(uv)
        for old in [rr,gr]:
            assert old['pair_index']==i and old['source_line']==i+1 and old['rgb_name']==sm[0] and old['thermal_name']==sm[1]
            assert old['UV_input']==old['UV_reference_label']==uv and old['Tmax']==float(sm[2])
        row=dict(pair_index=i,source_line=i+1,rgb_name=sm[0],thermal_name=sm[1],Tmax=float(sm[2]),UV_input=uv,UV_reference_label=uv,q=.5+.5*uv,rgb_count=rr['rgb_count'],thermal_count=rr['thermal_count'],rgb_binary=rr['rgb_binary'],thermal_binary=rr['thermal_binary'],prior_ordinary_BetP_binary=rr['ordinary_ds_binary'],prior_reliability_BetP_binary=rr['spatial_ds_binary'],prior_strict_BetP_visual_binary=gr['gate_visual_binary'],prior_strict_BetP_classicUV_binary=gr['gate_visual_uv_binary'])
        for branch,old,prefix in [('A',rr,'ordinary_ds'),('B',rr,'spatial_ds'),('C',gr,'raw_spatial_ds')]:
            mass=tuple(old[prefix+'_'+k] for k in ['fire','nofire','uncertain'])
            status=old.get('matching_status',old.get('status')); matched=status=='matched'; iou=old['polygon_iou']
            # Re-derive VISUAL-ONLY masses from selected real detector scores, never old UV masses.
            rm,tm=s.base(old['rgb_confidence'],.2),s.base(old['thermal_confidence'],.2)
            if branch!='A': rm,tm=s.discount(rm,old['reliability']),s.discount(tm,old['reliability'])
            expected=s.ds(rm,tm)[0]; assert np.allclose(expected,mass,atol=1e-12,rtol=0)
            r=apply(mass,uv,.5,.5,branch=='C',matched,iou,.3)
            assert r['fire']<=r['base_fire']
            for k in ['selected_candidate','rgb_index','thermal_index','rgb_confidence','thermal_confidence','reliability']:
                row[branch+'_'+k]=old[k]
            row[branch+'_matching_status']=status; row[branch+'_polygon_iou']=iou
            state='NoFire' if not matched or iou<=.1 else 'Transition' if iou<.3 else 'FireCandidate'
            row[branch+'_geometry_state']=state; row[branch+'_eligible']=r['eligible']
            for k,v,vp in zip(['F','N','U'],mass,r['mass']): row[branch+'_m'+k]=v; row[branch+'_mprime'+k]=vp
            row[branch+'_preUV_BetP']=mass[0]+.5*mass[2]
            row[branch+'_noUV_mF_binary']=r['base_fire']; row[branch+'_UV_mprimeF_binary']=r['fire']; row[branch+'_reason']=r['reason']
        rows.append(row)
    assert Counter(y)=={1:1398,0:130}
    methods={k:metric(y,[r[col] for r in rows]) for k,col in [('RGB','rgb_binary'),('Thermal','thermal_binary'),('prior_ordinary_BetP','prior_ordinary_BetP_binary'),('prior_reliability_BetP','prior_reliability_BetP_binary'),('prior_strict_BetP_visual','prior_strict_BetP_visual_binary'),('prior_strict_BetP_classicUV','prior_strict_BetP_classicUV_binary')]}
    transitions={}; coverage={}
    for b in LABELS:
        for tag in ['noUV_mF','UV_mprimeF']: methods[b+'_'+tag]=metric(y,[r[b+'_'+tag+'_binary'] for r in rows])
        transitions[b]=dict(kept_fire=sum(r[b+'_noUV_mF_binary']==r[b+'_UV_mprimeF_binary']==1 for r in rows),suppressed_fire=sum(r[b+'_noUV_mF_binary']==1 and r[b+'_UV_mprimeF_binary']==0 for r in rows),resurrected_fire=sum(r[b+'_noUV_mF_binary']==0 and r[b+'_UV_mprimeF_binary']==1 for r in rows),kept_nofire=sum(r[b+'_noUV_mF_binary']==r[b+'_UV_mprimeF_binary']==0 for r in rows),reason_counts=dict(Counter(r[b+'_reason'] for r in rows)))
        coverage[b]=dict(geometry_states=dict(Counter(r[b+'_geometry_state'] for r in rows)),matching_status=dict(Counter(r[b+'_matching_status'] for r in rows)),eligible_count=sum(r[b+'_eligible'] for r in rows),hard_gate_eligibility_fraction=sum(r[b+'_eligible'] for r in rows)/1528,binary_decision_coverage=1.,abstentions=0,visual_mF_range=[min(r[b+'_mF'] for r in rows),max(r[b+'_mF'] for r in rows)],UV_mprimeF_range=[min(r[b+'_mprimeF'] for r in rows),max(r[b+'_mprimeF'] for r in rows)],vacuous_count=sum(r[b+'_mF']==r[b+'_mN']==0 for r in rows))
    sensitivity=[]
    for lam in config['sensitivity_lambda']:
        for b in LABELS:
            ps=[int(r[b+'_eligible'] and r[b+'_mF']*(lam+(1-lam)*r['UV_input'])>.5) for r in rows]
            sensitivity.append(dict(branch=b,lambda_value=lam,tau_F=.5,diagnostic_only=True,**metric(y,ps)))
    allmet=dict(reference='UV fourth column; circular evaluation not independent scene truth',label_counts=dict(Counter(y)),config_sha256=config_hash,methods=methods,transitions=transitions,coverage=coverage)
    writecsv('fusion_results.csv',rows); save('fusion_results.json',dict(rows=rows)); save('metrics.json',allmet)
    writecsv('metrics.csv',[dict(method=k,**m) for k,m in methods.items()]); writecsv('sensitivity_metrics.csv',sensitivity)
    save('provenance.json',provenance); save('protected_before_hashes.json',protected)
    calibration=load(OUT.parent/'supplied_models_affine_iou/candidate_affine_calibration.json')
    make_report(allmet,calibration,sensitivity,config_hash)
    # Independent readback of every prediction plus recomputation from output CSV.
    reread=list(csv.DictReader((OUT/'fusion_results.csv').open(encoding='utf-8-sig')))
    assert len(reread)==1528 and [int(r['pair_index']) for r in reread]==list(range(1528))
    for b in LABELS:
        for tag in ['noUV_mF','UV_mprimeF']:
            assert metric([int(r['UV_reference_label']) for r in reread],[int(r[b+'_'+tag+'_binary']) for r in reread])==methods[b+'_'+tag]
        for r in reread:
            f=float(r[b+'_mF']); q=float(r['q']); mf=float(r[b+'_mprimeF']); eligible=int(r[b+'_eligible'])
            assert mf==f*q and int(r[b+'_UV_mprimeF_binary'])==int(eligible and mf>.5)
            assert abs(sum(float(r[b+'_mprime'+k]) for k in ['F','N','U'])-1)<1e-12
    with zipfile.ZipFile(REPORT) as z: assert z.testzip() is None
    d=Document(REPORT); text='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for rr in t.rows for c in rr.cells])
    assert all(x in text for x in ['1528','循环评价','m′(F)','5.7718','BetP','严格大于','不增加','0.7'])
    (OUT/'report_readback.txt').write_text(text,encoding='utf-8')
    assert sha(OUT/'frozen_config.json')==config_hash and all(sha(p)==h for p,h in protected.items())
    verification=dict(csv_rows=1528,unique_pair_ids=1528,unique_rgb_names=1528,unique_thermal_names=1528,current_image_hashes_verified=verified_images,input_csv_json_parity_verified=True,visual_only_masses_rederived=True,metrics_recomputed_from_csv=True,all_output_formulas_rechecked=True,docx_zip_crc_ok=True,docx_readback=True,docx_tables=len(d.tables),original_sources_and_old_reports_unchanged=True,frozen_config_unchanged=True,config_sha256=config_hash,report_path=str(REPORT),report_sha256=sha(REPORT),numerical_invariants=tests,verified_utc=datetime.now(timezone.utc).isoformat())
    save('deliverable_verification.json',verification)
    save('artifact_manifest.json',{str(p):sha(p) for p in OUT.iterdir() if p.is_file() and p.name not in ['artifact_manifest.json','tests_green.log']})
    print(json.dumps(dict(main=methods['C_UV_mprimeF'],ordinary=methods['A_UV_mprimeF'],reliability=methods['B_UV_mprimeF'],transitions=transitions,coverage=coverage,report=str(REPORT),verified=True),ensure_ascii=True,indent=2))
def make_report(m,cal,sens,cfg_hash):
    d=Document(); st=d.styles['Normal']; st.font.name='Microsoft YaHei'; st.font.size=Pt(9); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    sec=d.sections[0]; sec.left_margin=sec.right_margin=Inches(.6)
    def p(t): d.add_paragraph(str(t))
    def h(t): d.add_heading(t,1)
    def f(v): return '未定义（无预测正例）' if v is None else f'{100*v:.4f}%'
    def tab(head,records):
        t=d.add_table(rows=1,cols=len(head)); t.style='Table Grid'
        for c,v in zip(t.rows[0].cells,head): c.text=str(v)
        for r in records:
            for c,v in zip(t.add_row().cells,r): c.text=str(v)
    d.add_heading('UV乘性门控D-S融合：1528对真实重算统计',0)
    p('实际全量执行1528对、逐对唯一索引，UV标签1398有火／130无火。复用用户best.pt可见光、best-2.pt热成像的Ultralytics YOLO26n真实预测（不是PPYOLOv2），检测分数>0.5；无训练、无推理、无新标定。全部3056图像SHA256重新核验。旧文件和报告只读保留。')
    h('1. 结论与全量混淆矩阵')
    p('主方案保留最新双模态空间一致性必要条件，UV改为乘性门控而不是第三个D-S证据。以下m(F)严格是仅可见光和热成像所得的质量，不使用旧UV融合质量，不静默改成BetP。结果真实报告，不根据测试集调参弥补负面结果。')
    keys=['RGB','Thermal','A_noUV_mF','A_UV_mprimeF','B_noUV_mF','B_UV_mprimeF','C_noUV_mF','C_UV_mprimeF']
    names={'RGB':'RGB检测','Thermal':'热成像检测'}
    for b in LABELS:
        names[b+'_noUV_mF']=LABELS[b]+'：无UV mF>0.5'; names[b+'_UV_mprimeF']=LABELS[b]+'：UV乘性 m′F>0.5'
    tab(['方法','TP','FP','FN','TN','Acc','P','R','F1'],[[names[k],*(m['methods'][k][x] for x in ['TP','FP','FN','TN']),*(f(m['methods'][k][x]) for x in ['accuracy','precision','recall','F1'])] for k in keys])
    tab(['方法','Balanced Acc','Specificity','FPR','FNR'],[[names[k],*(f(m['methods'][k][x]) for x in ['balanced_accuracy','specificity','FPR','FNR'])] for k in keys])
    for b in LABELS:
        base,new=m['methods'][b+'_noUV_mF'],m['methods'][b+'_UV_mprimeF']
        p(f'{LABELS[b]}：UV乘性相对完全相同视觉mF门限的Accuracy变化={100*(new["accuracy"]-base["accuracy"]):+.4f}个百分点；Recall变化={100*(new["recall"]-base["recall"]):+.4f}个百分点。视觉mF范围={m["coverage"][b]["visual_mF_range"]}；转移={json.dumps(m["transitions"][b],ensure_ascii=False)}。')
    h('2. 用户字面公式、冻结默认和质量守恒')
    p('用户公式：m′(F)=m(F)(λ+(1−λ)g_uv)。g_uv为paper_sample.txt第四列原始二值0/1；q=λ+(1−λ)g_uv。预先固定λ=0.5、τF=0.5，判火用m′(F)>τF（严格大于，不含等号）。这两个数值是可配置工程默认，不是用户指定数值，也不是验证最优值；λ不是历史时序平滑λ=0.7，本次不使用时序UV。')
    p('UV1：q=1，保持视觉mF；UV0：q=λ，只抑制视觉mF。q∈[λ,1]，所以m′F≤mF，绝不增加有火质量；同一固定阈值下不可能把视觉无火恢复成有火，也不能提高相对同一视觉基线的召回率。UV仍可减少UV0对应误报，但不保证总体提升。')
    p('为保持BPA：m′N=mN，m′U=mU+(1−q)mF，总和为1。抑制的F质量只转移到未知Θ，不转移到无火。不重新归一化、不放大用户字面m′F。不把UV作为第三D-S源，也不使用旧经典UV质量作为本次基础。m(F)是D-S证据质量，并非已校准的概率或统计置信度。BetP(F)=mF+0.5mU只是旧判决量，此次不用。')
    p('冻结配置SHA256='+cfg_hash+'；配置与回执在指标计算前保存，固定敏感性λ={0,0.25,0.5,0.75,1}，τF=0.5不变，只诊断不挑选最优。')
    h('3. 三分支及空间/视觉证据链')
    p('基础在场检测BPA=(0.8c,0.8(1−c),0.2)；缺失为(0,0,1)，不是c=0无火证据。两视觉质量D-S：K=aF bN+aN bF；mF=(aF bF+aF bU+aU bF)/(1−K)，mN对应互换F/N，mU=aU bU/(1−K)。实现沿用旧支持总和归一化与全冲突退回未知规则。仅在视觉阶段使用DS；UV乘法不再进行DS。')
    p('A：普通视觉DS无空间硬门，沿用原spatial_reliability_ds候选（最大普通mF，平手在场最小分数，再顺序）；回应前表普通分支。B：同A候选，但用原可靠性r=0.5+0.5s（matched）、未匹配r=0.5，无硬门；折扣(r mF,r mN,1−r+r mU)。A/B匹配最小IoU=0.01。')
    p('C主方案：沿用最新spatial_gate_ds候选和几何；Hungarian一对一多边形IoU匹配，正IoU匹配底限1e−12，优先最大IoU，平手视觉BetP、最小分数、候选索引。BetP仅用于原有候选选取，不用于最终判火。g=clip((s−0.1)/(0.3−0.1),0,1)，matched才折扣两视觉BPA，未匹配g=0。τL=0.1、τH=0.3保持原冻结默认。最终Fire=matched且s≥0.3且m′F>0.5；不一致、Transition、缺任一框或全缺框均二分类无火，UV1也不得绕过。')
    p('几何NoFire：未匹配或s≤0.1；Transition：0.1<s<0.3；FireCandidate：matched且s≥0.3。硬无火是用户指定工程决策，不代表D-S本身证实物理无火。A/B不是严格门控的替代；候选政策不同，跨A/C差异不能全部归因于空间模块。每一分支内部无UV与乘性UV用完全相同候选和视觉质量，才能隔离UV作用。')
    p('每帧单一代表候选，不把相关多框D-S累加；来源selected_candidate、两侧索引、检测分数、IoU、可靠性、几何状态、原质量／乘性后质量、q、判决和原因全部输出CSV/JSON。Tmax保留原值，单位与物理温度可靠性未知，本次不新加温度门限。')
    h('4. 覆盖率、状态与旧BetP方法区别')
    p('全部1528帧纳入二分类，decision coverage=100%，无弃权；未知质量不是从统计中删掉帧。空间候选覆盖率只代表工程eligible，不是二分类评价覆盖率。')
    for b in LABELS: p(LABELS[b]+'：'+json.dumps(m['coverage'][b],ensure_ascii=False))
    p('下表为已有公布方法的原输出重新统计，明确不是本次mF判决，也不是本次UV乘法。经典UV DS历史分支仅作为有标签泄漏风险的对照，主方法彻底移除。判决量变化与UV变化不得混为同一增益。')
    prior=[k for k in m['methods'] if k.startswith('prior_')]
    tab(['历史方法（BetP）','TP','FP','FN','TN','Acc'],[[k,*(m['methods'][k][x] for x in ['TP','FP','FN','TN']),f(m['methods'][k]['accuracy'])] for k in prior])
    h('5. 预固定敏感性（不选最佳参数）')
    tab(['分支','λ','TP','FP','FN','TN','Acc'],[[r['branch'],r['lambda_value'],r['TP'],r['FP'],r['FN'],r['TN'],f(r['accuracy'])] for r in sens])
    h('6. 局限、来源与复核')
    p('评价参考是用户指定第四列UV，同时UV参与门控，因此是循环评价／UV符合度，不是独立场景真值准确率，不能得到定位AP。正例占1398/1528，Accuracy与Precision不能替代Recall和Balanced Accuracy；零预测正例Precision标未定义而不伪造100%。UV单独按定义等于标签100%，不算独立性能。')
    p(f'沿用七点候选FULL affine（热→RGB），未做新标定：fit RMSE={cal["fit_rmse_px"]:.4f}px，留一LOO RMSE={cal["leave_one_out_rmse_px"]:.4f}px，max={cal["leave_one_out_max_px"]:.4f}px；候选未独立验证。七点窄带、全视场／跨场景／视差有效性未知，LOO 5.7718px不能证明独立可靠。矩阵：'+json.dumps(cal['affine_matrix_2x3']))
    p('源输入CSV与JSON逐字段一致，1528顺序与paper_sample第四列一致、两模态各1528唯一图名。重新核验3056图像和权重哈希，从选中真实检测分数重新计算视觉BPA／DS验证缓存质量；数值不变量400组覆盖λ端点、UV端点、严格mF边界、空间0.3等号、质量守恒、只减不增、全缺失与UV不得越门。CSV独立重算指标和全部输出公式；DOCX ZIP CRC／读回／SHA256验证；旧文件逐文件哈希保留。')
    p('产物目录：'+str(OUT)+'；fusion_results.csv/json、metrics.json/csv、sensitivity_metrics.csv、frozen_config.json、config_freeze_receipt.json、provenance.json、numerical_invariant_tests.json、deliverable_verification.json、artifact_manifest.json。')
    d.save(REPORT)
if __name__=='__main__': run()
