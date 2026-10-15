from pathlib import Path
import sys,json,csv,hashlib,datetime,zipfile,collections
import numpy as np
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
ROOT=Path('C:/lyk/fire/qualityDetection'); OLD=ROOT/'fusion/supplied320_paper_experiment'; OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(OLD))
from fuse320 import corners,transform,polygon_iou,match,metrics
SOURCE=Path('C:/lyk/paper_data/paper_data_/paper_data/paper_sample.txt')
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
load=lambda p:json.loads(Path(p).read_text(encoding='utf-8-sig'))
def save(name,obj):
 with (OUT/name).open('x',encoding='utf-8') as f:json.dump(obj,f,ensure_ascii=False,indent=2,allow_nan=False)
def csvsave(name,rows,fields=None):
 with (OUT/name).open('x',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields or list(rows[0]));w.writeheader();w.writerows(rows)
def parse(text):
 rows=[];seen=set()
 for n,s in enumerate(text.splitlines(),1):
  if not s.strip():continue
  a=s.split()
  if len(a)!=4 or a[3] not in ('0','1'):raise ValueError(f'Malformed row/UV line {n}: {s}')
  key=(a[0],a[1])
  if key in seen:raise ValueError(f'Ambiguous duplicate pair line {n}')
  if Path(a[0]).name!=a[0] or Path(a[1]).name!=a[1]:raise ValueError('Expected exact basenames')
  if not np.isfinite(float(a[2])):raise ValueError('Invalid Tmax')
  seen.add(key);rows.append(dict(source_line=n,rgb_name=a[0],thermal_name=a[1],Tmax=float(a[2]),UV_reference=int(a[3])))
 return rows
METHODS=['rgb_baseline','thermal_baseline','rgb_operating','thermal_operating','main_fire','optional_uv_fire']
LABELS=['RGB >0.5','热成像 >0.5','RGB ≥0.2（公平候选）','热成像 ≥0.1（公平候选）','空间DS（无UV主方案）','空间DS × UV（λ=0.5）']
def summary(y,p):
 m=metrics(y,p);specificity=m['TN']/(m['TN']+m['FP']) if m['TN']+m['FP'] else None
 m['balanced_accuracy']=(m['recall']+specificity)/2 if m['recall'] is not None and specificity is not None else None
 return m
def targets(results,method):
 m=results[method];out={}
 for name in ['rgb_operating','thermal_operating','rgb_baseline','thermal_baseline']:
  b=results[name];out[name]=dict(FPR_strictly_lower=m['FPR']<b['FPR'],FNR_strictly_lower=m['FNR']<b['FNR'],both_strictly_lower=m['FPR']<b['FPR'] and m['FNR']<b['FNR'])
 out['fair_both_modalities']=all(out[n]['both_strictly_lower'] for n in ['rgb_operating','thermal_operating'])
 out['standard_both_modalities']=all(out[n]['both_strictly_lower'] for n in ['rgb_baseline','thermal_baseline'])
 return out

def run():
 started=datetime.datetime.now(datetime.timezone.utc).isoformat();source_hash=sha(SOURCE)
 protected={str(p):sha(p) for p in OLD.rglob('*') if p.is_file()}
 current=parse(SOURCE.read_text(encoding='utf-8-sig'));assert len(current)==1528
 oldrows=list(csv.DictReader((OLD/'frame_predictions.csv').open(encoding='utf-8-sig')))
 oldmap={(r['rgb_name'],r['thermal_name']):r for r in oldrows};assert len(oldmap)==len(oldrows)==1528
 assert set(oldmap)=={(r['rgb_name'],r['thermal_name']) for r in current},'Pair sets changed'
 oldaudit={(r['rgb'],r['thermal']):r for r in load(ROOT/'audit/test_pairs.json')};assert len(oldaudit)==1528
 protocol=load(OLD/'frozen_protocol.json');assert sha(OLD/'frozen_protocol.json')==load(OLD/'freeze_receipt.json')['protocol_sha256']
 assert protocol['parameters']==dict(tR=.2,tT=.1,iou_gate=.01,rmin=1.,decision_threshold=.3)
 pred={m:load(OLD/(m+'_predictions.json')) for m in ['rgb','thermal']};pmap={m:{r['pair_index']:r for r in rs} for m,rs in pred.items()}
 assert all(len(rs)==1528 and set(rs)==set(range(1528)) for rs in pmap.values())
 oldevidence={r['pair_index']:r for r in load(OLD/'matched_evidence.json')}
 models=load(OLD/'model_metadata.json');model_checks={}
 for m,r in models.items():
  assert sha(r['source'])==r['sha256'];assert r['sha256']==protocol['model_sources'][m]['sha256'];model_checks[m]=dict(path=r['source'],sha256=r['sha256'],verified=True,architecture=r['architecture']['yaml_file'])
 assert sha(protocol['geometry']['source'])==protocol['geometry']['source_sha256']
 A=np.array(protocol['geometry']['affine_matrix_2x3']);changes=[];mapping=[];rows=[];evidence=[];imagehashes=set();reordered=0
 for i,r in enumerate(current):
  key=(r['rgb_name'],r['thermal_name']);old=oldmap[key];oi=int(old['pair_index']);reordered+=i!=oi
  assert r['Tmax']==oldaudit[key]['Tmax'],'Tmax changed';assert int(old['UV_reference'])==int(oldaudit[key]['UV_input'])
  mapping.append(dict(current_pair_index=i,current_source_line=r['source_line'],cached_pair_index=oi,cached_source_line=old['source_line'],rgb_name=key[0],thermal_name=key[1]))
  if r['UV_reference']!=int(old['UV_reference']):changes.append(dict(**mapping[-1],old_UV=int(old['UV_reference']),new_UV=r['UV_reference']))
  for m,folder,name in [('rgb','bgr',key[0]),('thermal','thermal',key[1])]:
   pr=pmap[m][oi];p=SOURCE.parent/folder/name;assert pr['image_name']==name;assert sha(p)==pr['image_sha256'];imagehashes.add(pr['image_sha256'])
   assert all(s[-2:]==[320,320] for s in pr['model_input_shapes']);assert pr['device']=='cuda:0';assert all(b['class_id']==0 for b in pr['detections'])
  rb=pmap['rgb'][oi]['detections'];tb=pmap['thermal'][oi]['detections']
  ious=np.array([[polygon_iou(corners(b['bbox']),transform(corners(t['bbox']),A)) for t in tb] for b in rb])
  matches=match(rb,tb,ious);best=max(matches,key=lambda x:x['mF'],default=None);mf=best['mF'] if best else 0.
  assert matches==oldevidence[oi]['matches'] and best==oldevidence[oi]['selected'],'Geometry/matching mismatch'
  assert abs(mf-float(old['mass_fire']))<1e-12
  row=dict(pair_index=i,cached_pair_index=oi,**r,old_UV_reference=int(old['UV_reference']),rgb_box_count=len(rb),thermal_box_count=len(tb),rgb_candidates=sum(b['score']>=.2 for b in rb),thermal_candidates=sum(b['score']>=.1 for b in tb),matched_count=len(matches),selected_rgb_index=best['r'] if best else -1,selected_thermal_index=best['t'] if best else -1,selected_iou=best['iou'] if best else 0.,mass_fire=mf,mass_nofire=0.,mass_uncertain=1-mf,conflict=0.,main_fire=int(mf>.3),rgb_baseline=int(any(b['score']>.5 for b in rb)),thermal_baseline=int(any(b['score']>.5 for b in tb)),rgb_operating=int(any(b['score']>=.2 for b in rb)),thermal_operating=int(any(b['score']>=.1 for b in tb)),optional_uv_mass=mf*(.5+.5*r['UV_reference']),optional_uv_fire=int(mf*(.5+.5*r['UV_reference'])>.3))
  for method in METHODS[:-1]:assert row[method]==int(old[method])
  rows.append(row);evidence.append(dict(pair_index=i,cached_pair_index=oi,matches=matches,selected=best))
 assert len(imagehashes)==3056
 assert changes,'No corrected UV labels at canonical source: stop and locate modified file'
 oldhash=load(ROOT/'fusion/crossmodal_roi_ds/full_receipt.json')['protected_hashes'][str(SOURCE)]
 assert source_hash!=oldhash
 y=[r['UV_reference'] for r in rows];oy=[r['old_UV_reference'] for r in rows]
 results={m:summary(y,[r[m] for r in rows]) for m in METHODS}
 oldresults={m:summary(oy,[int(oldmap[(r['rgb_name'],r['thermal_name'])][m]) for r in rows]) for m in METHODS}
 for m in METHODS:
  for k in ['TP','FP','FN','TN']:assert oldresults[m][k]==load(OLD/'metrics.json')['results'][m][k]
 audit=dict(source=str(SOURCE),old_source_sha256=oldhash,old_hash_evidence=str(ROOT/'fusion/crossmodal_roi_ds/full_receipt.json'),old320_source_hash_recorded=False,old_hash_note='320 protocol did not directly record txt SHA; prior protected source receipt and identical recorded per-row old labels/Tmax establish original reference.',current_source_sha256=source_hash,old_positive=sum(oy),old_negative=len(oy)-sum(oy),current_positive=sum(y),current_negative=len(y)-sum(y),changed_labels=len(changes),zero_to_one=sum(c['old_UV']==0 for c in changes),one_to_zero=sum(c['old_UV']==1 for c in changes),changed_source_lines=[c['current_source_line'] for c in changes],reordered_rows=reordered,all_Tmax_unchanged=True,image_hashes_verified=3056,models=model_checks,new_forward=False,new_training=False,original_prediction_hashes={m:sha(OLD/(m+'_predictions.json')) for m in pred},frozen_protocol_sha256=sha(OLD/'frozen_protocol.json'),all_matched_geometry_and_mass_reproduced=True,UV_zero_optional_positive=sum(r['UV_reference']==0 and r['optional_uv_fire'] for r in rows),optional_decision_changes_vs_old=sum(r['optional_uv_fire']!=int(oldmap[(r['rgb_name'],r['thermal_name'])]['optional_uv_fire']) for r in rows),started_utc=started)
 save('source_audit.json',audit);save('frozen_protocol_reused.json',protocol);save('matched_evidence.json',evidence)
 payload=dict(reference='Current fourth-column UV reference; not independent scene ground truth',results=results,old_results_obsolete_for_current_reference=oldresults,targets={m:targets(results,m) for m in ['main_fire','optional_uv_fire']})
 save('metrics.json',payload);csvsave('metrics.csv',[dict(method=m,**results[m]) for m in METHODS]);csvsave('frame_predictions.csv',rows);csvsave('row_mapping.csv',mapping);csvsave('old_new_label_diff.csv',changes)
 doc=Document();style=doc.styles['Normal'];style.font.name='宋体';style.font.size=Pt(10.5);style.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
 doc.add_heading('320输入：修正UV标签后的空间D-S冻结协议复验',0)
 doc.add_heading('摘要',1)
 doc.add_paragraph(f'本研究在用户修正paper_sample.txt第四列后，对1528组图像进行冻结协议复验。新参考中UV=1为{sum(y)}组、UV=0为{len(y)-sum(y)}组；共修改{len(changes)}个标签，其中0→1为{audit["zero_to_one"]}个，1→0为{audit["one_to_zero"]}个。复用经过SHA256核验的真实320×320模型前向缓存，重新执行仿射多边形IoU、匈牙利配对和D-S质量计算，并依据当前参考重新统计全部方法。未重训、未重新神经网络前向、未针对新标签调参。')
 doc.add_heading('1 数据与溯源',1)
 doc.add_paragraph(f'来源：{SOURCE}。第一、第二列图像对与历史缓存通过精确文件名对唯一映射，非仅按行号对齐；发生顺序变化的行数为{reordered}。第三列Tmax逐项与原始审计一致；3056张原图SHA256均与实际推理记录一致。所有历史脚本、缓存、报告及用户当前标签均保持只读。')
 doc.add_paragraph(f'旧文本SHA256：{oldhash}；当前SHA256：{source_hash}。320记录没有直接保存文本散列，旧散列由之前的受保护来源回执提供，逐行UV和Tmax另与320缓存和原始审计交叉核验。精确变更行见old_new_label_diff.csv。全部旧指标仅适用于旧标签，不能再作为当前参考结果引用。')
 for m,r in model_checks.items():doc.add_paragraph(f'{m}：{r["path"]}；YOLO26n（非PPYOLOv2）；SHA256={r["sha256"]}。权重与原始320实验一致。')
 doc.add_heading('2 方法与公式',1)
 doc.add_paragraph('原前向：imgsz=320、rect=False、实际张量320×320、raw score>0.05、max_det=300、fire单类。YOLO26n端到端输出绕过常规NMS；原iou=0.7参数不能解释为实际NMS。融合先筛RGB≥0.2与thermal≥0.1，热框四角用冻结2×3仿射矩阵映射至RGB，以凸多边形交并比进行dummy-enabled Hungarian最大总IoU一对一配对，IoU≥0.01。配对后取最大mF候选；同分按原配对顺序。')
 doc.add_paragraph('正类简单支持BPA：mR(F)=cR, mR(N)=0, mR(Θ)=1−cR；热成像同理。rmin=1，无动态可靠性折扣。K=0，mF=1−(1−cR)(1−cT)，mN=0，mΘ=1−mF；无配对时mF=0。主决策mF>0.3，完全不使用UV标签。可选分支mUV=mF×(0.5+0.5×UV)，决策仍为严格>0.3；UV=0仅减半，并非强制否决，mF>0.6时仍可判火。')
 doc.add_paragraph('FPR=FP/(FP+TN)，FNR=FN/(TP+FN)，Accuracy=(TP+TN)/N，Precision=TP/(TP+FP)，Recall=TP/(TP+FN)，F1=2TP/(2TP+FP+FN)，Balanced accuracy=(Recall+TN/(TN+FP))/2。退化分母输出None，不用伪零。比较为同一缓存、同一320输入、同一1528组当前标签。')
 doc.add_heading('3 当前结果：公平候选与标准阈值',1)
 fields=['TP','FP','FN','TN','FPR','FNR','accuracy','precision','recall','F1','balanced_accuracy']
 def table(methods,values):
  t=doc.add_table(rows=1,cols=1+len(fields));t.style='Table Grid'
  for c,s in zip(t.rows[0].cells,['方法']+fields):c.text=s
  for m in methods:
   cells=t.add_row().cells;cells[0].text=LABELS[METHODS.index(m)]
   for c,k in zip(cells[1:],fields):c.text=str(values[m][k]) if k in ['TP','FP','FN','TN'] else ('None' if values[m][k] is None else f'{values[m][k]*100:.4f}%')
 table(['rgb_operating','thermal_operating','main_fire','optional_uv_fire'],results)
 doc.add_paragraph('上表为首要公平比较，单模态保留与融合完全相同的候选置信度阈值；融合的双真实检测门控和几何一致性仍为方法差异。以下>0.5仅为标准阈值补充，不能将其收益全部归因于D-S。')
 table(['rgb_baseline','thermal_baseline','main_fire','optional_uv_fire'],results)
 doc.add_heading('4 标签修正前后比较与目标检验',1)
 t=doc.add_table(rows=1,cols=9);t.style='Table Grid'
 for c,s in zip(t.rows[0].cells,['方法','旧TP','旧FP','旧FN','旧TN','新TP','新FP','新FN','新TN']):c.text=s
 for m in METHODS:
  cells=t.add_row().cells
  for c,s in zip(cells,[LABELS[METHODS.index(m)]]+[str(oldresults[m][k]) for k in ['TP','FP','FN','TN']]+[str(results[m][k]) for k in ['TP','FP','FN','TN']]):c.text=s
 for m in ['main_fire','optional_uv_fire']:
  target=payload['targets'][m];doc.add_paragraph(f'{LABELS[METHODS.index(m)]}：同时严格降低误报率与漏报率，公平候选双单模态目标={target["fair_both_modalities"]}；标准>0.5双单模态目标={target["standard_both_modalities"]}。')
  for b in ['rgb_operating','thermal_operating','rgb_baseline','thermal_baseline']:doc.add_paragraph(f'相对{LABELS[METHODS.index(b)]}：FPR严格降低={target[b]["FPR_strictly_lower"]}，FNR严格降低={target[b]["FNR_strictly_lower"]}。')
 doc.add_paragraph(f'当前UV=0但可选UV分支仍判火：{audit["UV_zero_optional_positive"]}组。可选UV分支相对旧决策改变{audit["optional_decision_changes_vs_old"]}组；主分支和四个视觉单模态分支决策均未改变，变化来自参考标签。')
 doc.add_heading('5 讨论与局限',1)
 overlap=load(OLD/'source_overlap_audit.json')['stats'];doc.add_paragraph('当前源训练/验证内容与测试图像完全散列重合：'+json.dumps(overlap,ensure_ascii=False)+'。这证明当前内容重合，不能证明这些图像曾进入历史权重训练。原冻结参数已经探索性使用同一数据，修正标签复验不是独立测试，也不是重新选参。仿射仅7个窄带点拟合，LOO RMSE约5.771805像素，缺少独立场景全视场标定与视差核验；它仍是候选未验证配准。')
 doc.add_paragraph('第四列按用户要求作为二元参考，并非独立人工场景真值。主方案不读取UV作为输入，因此避免直接UV门控循环，但仍只报告与该参考的一致性。可选UV既作输入又作真值，指标存在循环/泄漏，不能作为独立融合性能证据。简单支持D-S只有OR强化和空间双框门控，K=0且rmin=1，不应声称冲突消解或动态可靠性增益。')
 doc.add_heading('6 结论',1)
 v=results['main_fire'];doc.add_paragraph(f'修正参考后主方案TP={v["TP"]}、FP={v["FP"]}、FN={v["FN"]}、TN={v["TN"]}，FPR={v["FPR"]*100:.4f}%、FNR={v["FNR"]*100:.4f}%。是否同时优于两种公平候选单模态：{payload["targets"]["main_fire"]["fair_both_modalities"]}；是否同时优于两种标准阈值单模态：{payload["targets"]["main_fire"]["standard_both_modalities"]}。据实际统计报告，不为达到目标调参。')
 report=ROOT/'reports/320输入_修正UV标签_空间DS复验报告.docx'
 if report.exists():raise FileExistsError(report)
 doc.save(report)
 with zipfile.ZipFile(report) as z:assert z.testzip() is None
 rd=Document(report);assert len(rd.tables)==3
 assert '6 结论' in '\n'.join(p.text for p in rd.paragraphs)
 assert sha(SOURCE)==source_hash
 assert all(sha(p)==h for p,h in protected.items()),'Historical artifacts changed'
 save('execution_verification.json',dict(success=True,source_unchanged=True,all_old_artifacts_unchanged=True,old_artifacts_verified=len(protected),source_images_verified=3056,all_geometry_recomputed_and_reproduced=True,all_main_visual_decisions_unchanged=True,report_crc_ok=True,report_tables=3,report=str(report),report_sha256=sha(report),completed_utc=datetime.datetime.now(datetime.timezone.utc).isoformat()))
 (OUT/'report_readback.txt').write_text('\n'.join(p.text for p in rd.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in rd.tables for r in t.rows),encoding='utf-8')
 save('artifact_sha256.json',{str(p):sha(p) for p in OUT.iterdir() if p.is_file() and p.name!='artifact_sha256.json'}|{str(report):sha(report)})
 print(json.dumps(dict(audit=audit,results=results,targets=payload['targets'],report=str(report)),ensure_ascii=False,indent=2))
if __name__=='__main__':run()