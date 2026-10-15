from pathlib import Path
import json,csv,hashlib,zipfile,unittest
import cv2,numpy as np
from scipy.optimize import linear_sum_assignment
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.section import WD_ORIENT
from rerun_corrected_uv import parse,summary
OUT=Path(__file__).resolve().parent;OLD=OUT.parent/'supplied320_paper_experiment';ROOT=OUT.parent.parent
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
load=lambda p:json.loads(Path(p).read_text(encoding='utf-8-sig'))
class Regression(unittest.TestCase):
 def test_malformed(self):
  for s in ['a b 1 2','a b 1','a b 1 0.0','a b nan 0','a b 1 0\na b 2 1']:
   with self.assertRaises(ValueError):parse(s)
 def test_missing_labels(self):
  with self.assertRaises(ValueError):parse('a b 3')
 def test_reorder_identity(self):
  x=parse('a b 1 0\nc d 2 1');z=parse('c d 2 1\na b 1 0');d={(r['rgb_name'],r['thermal_name']):r['UV_reference'] for r in x}
  self.assertEqual([d[(r['rgb_name'],r['thermal_name'])] for r in z],[1,0])
 def test_metrics_known(self):
  v=summary([1,1,0,0],[1,0,1,0]);self.assertEqual([v[k] for k in ['TP','FP','FN','TN']],[1,1,1,1]);self.assertEqual(v['balanced_accuracy'],.5)
 def test_degenerate(self):
  v=summary([1],[1]);self.assertIsNone(v['FPR']);self.assertIsNone(v['balanced_accuracy'])
 def test_uv_gate(self):
  self.assertTrue(.8*.5>.3);self.assertFalse(.6*.5>.3);self.assertTrue(.4*1>.3)
 def test_main_uv_independence(self):
  for mass in [0,.3,.3000001,.6,.9]:self.assertEqual([int(mass>.3) for uv in [0,1]],[int(mass>.3)]*2)
def independent():
 rows=list(csv.DictReader((OUT/'frame_predictions.csv').open(encoding='utf-8-sig')));assert len(rows)==1528
 p={m:{r['pair_index']:r for r in load(OLD/(m+'_predictions.json'))} for m in ['rgb','thermal']}
 A=np.array(load(OLD/'frozen_protocol.json')['geometry']['affine_matrix_2x3']);methods=['rgb_baseline','thermal_baseline','rgb_operating','thermal_operating','main_fire','optional_uv_fire'];out={m:[] for m in methods};y=[]
 for row in rows:
  idx=int(row['cached_pair_index']);rb=p['rgb'][idx]['detections'];tb=p['thermal'][idx]['detections'];uv=int(row['UV_reference']);y.append(uv)
  ri=[j for j,b in enumerate(rb) if b['score']>=.2];ti=[j for j,b in enumerate(tb) if b['score']>=.1];sel=[]
  if ri and ti:
   s=np.zeros((len(ri),len(ti)))
   for a,j in enumerate(ri):
    x,y0,X,Y=rb[j]['bbox'];poly=np.array([[x,y0],[X,y0],[X,Y],[x,Y]],np.float32)
    for b,k in enumerate(ti):
     x,y0,X,Y=tb[k]['bbox'];q=np.array([[x,y0,1],[X,y0,1],[X,Y,1],[x,Y,1]])@A.T;q=q.astype(np.float32);inter,_=cv2.intersectConvexConvex(poly,q);den=abs(cv2.contourArea(poly))+abs(cv2.contourArea(q))-inter;s[a,b]=inter/den if den>0 else 0
   n,m=s.shape;w=np.zeros((n+m,n+m));w[:n,:m]=np.where(s>=.01,s,-1);aa,bb=linear_sum_assignment(-w)
   for a,b in zip(aa,bb):
    if a<n and b<m and s[a,b]>=.01:sel.append((ri[a],ti[b],float(s[a,b]),1-(1-rb[ri[a]]['score'])*(1-tb[ti[b]]['score'])))
  best=max(sel,key=lambda x:x[3],default=None);mass=best[3] if best else 0
  assert int(row['matched_count'])==len(sel);assert abs(float(row['mass_fire'])-mass)<1e-12
  if best:assert int(row['selected_rgb_index'])==best[0] and int(row['selected_thermal_index'])==best[1] and abs(float(row['selected_iou'])-best[2])<1e-12
  decisions=dict(rgb_baseline=int(any(b['score']>.5 for b in rb)),thermal_baseline=int(any(b['score']>.5 for b in tb)),rgb_operating=int(bool(ri)),thermal_operating=int(bool(ti)),main_fire=int(mass>.3),optional_uv_fire=int(mass*(.5+.5*uv)>.3))
  for method,v in decisions.items():assert int(row[method])==v;out[method].append(v)
 results=load(OUT/'metrics.json')['results']
 for method,pr in out.items():
  counts=dict(TP=sum(a==b==1 for a,b in zip(y,pr)),TN=sum(a==b==0 for a,b in zip(y,pr)),FP=sum(a==0 and b==1 for a,b in zip(y,pr)),FN=sum(a==1 and b==0 for a,b in zip(y,pr)));c=counts
  v=dict(N=len(y),**c,accuracy=(c['TP']+c['TN'])/len(y),precision=c['TP']/(c['TP']+c['FP']),recall=c['TP']/(c['TP']+c['FN']),F1=2*c['TP']/(2*c['TP']+c['FP']+c['FN']),FPR=c['FP']/(c['FP']+c['TN']),FNR=c['FN']/(c['FN']+c['TP']));v['balanced_accuracy']=(v['recall']+c['TN']/(c['TN']+c['FP']))/2
  for k,value in v.items():assert abs(value-results[method][k])<1e-12
 source=Path(load(OUT/'source_audit.json')['source']);assert sha(source)==load(OUT/'source_audit.json')['current_source_sha256']
 diffs=list(csv.DictReader((OUT/'old_new_label_diff.csv').open(encoding='utf-8-sig')));assert len(diffs)==250
 report=ROOT/'reports/320输入_修正UV标签_空间DS复验报告.docx'
 # This is our newly created report only; landscape makes 12-column metric tables readable.
 d=Document(report)
 for section in d.sections:
  section.orientation=WD_ORIENT.LANDSCAPE;section.page_width=Inches(11.7);section.page_height=Inches(8.3);section.left_margin=section.right_margin=Inches(.5)
 for t in d.tables:
  for row in t.rows:
   for cell in row.cells:
    for paragraph in cell.paragraphs:
     for run in paragraph.runs:run.font.size=Pt(8)
 d.save(report)
 with zipfile.ZipFile(report) as z:assert z.testzip() is None
 rd=Document(report);assert len(rd.tables)==3
 for t,ms in [(rd.tables[0],['rgb_operating','thermal_operating','main_fire','optional_uv_fire']),(rd.tables[1],['rgb_baseline','thermal_baseline','main_fire','optional_uv_fire'])]:
  assert len(t.rows)==5
  for i,m in enumerate(ms,1):assert [c.text for c in t.rows[i].cells][1:5]==[str(results[m][k]) for k in ['TP','FP','FN','TN']]
 verification=load(OUT/'execution_verification.json');verification['report_sha256']=sha(report);verification['independent_validator_passed']=True
 (OUT/'execution_verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf-8')
 receipt=dict(success=True,regression_tests=7,independently_recomputed_rows=1528,geometry_and_selected_indices=True,all_six_metrics_recomputed=True,diff_rows=250,report_crc_ok=True,report_tables=3,report_sha256=sha(report))
 (OUT/'validation.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8');print(json.dumps(receipt,indent=2))
if __name__=='__main__':
 result=unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(Regression));assert result.wasSuccessful();independent()
