from pathlib import Path
import sys,json,csv,hashlib,zipfile,collections
import numpy as np
import cv2
from scipy.optimize import linear_sum_assignment
from scipy.stats import beta
from docx import Document
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(OUT));import run_dual as run
load=run.load

def iou(a,b,A=None):
 def corners(b):
  x,y,X,Y=b;return np.float32([[x,y],[X,y],[X,Y],[x,Y]])
 p=corners(a);q=corners(b)
 if A is not None:q=cv2.transform(q[None],np.asarray(A,dtype=np.float32))[0]
 ap=cv2.contourArea(p);aq=cv2.contourArea(q)
 inter,_=cv2.intersectConvexConvex(p,q);return float(inter/(ap+aq-inter)) if ap+aq-inter>0 else 0.
def metric(y,p):
 tp=sum(a==b==1 for a,b in zip(y,p));fp=sum(a==0 and b==1 for a,b in zip(y,p));fn=sum(a==1 and b==0 for a,b in zip(y,p));tn=sum(a==b==0 for a,b in zip(y,p));return dict(N=len(y),TP=tp,FP=fp,FN=fn,TN=tn,FPR=fp/(fp+tn),FNR=fn/(tp+fn),precision=tp/(tp+fp) if tp+fp else 0.,recall=tp/(tp+fn),F1=2*tp/(2*tp+fp+fn),accuracy=(tp+tn)/len(y))
def unique(boxes):
 z=[]
 for b in sorted(boxes,key=lambda b:-b['score']):
  if all(iou(b['bbox'],a['bbox'])<.8 for a in z):z.append(b)
 return z

def verify():
 m=load(OUT/'metrics.json');initial=load(OUT/'initial_diagnosis.json');reference=run.old.parse();y=[r['UV_reference'] for r in reference];A=run.A
 for p,h in initial['protected'].items():assert run.sha(p)==h,('protected changed',p)
 receipt=load(OUT/'forward_receipt.json');assert receipt['full_images']==3056 and all(s[-2:]==[320,320] for s in receipt['shapes'])
 for modality in ['rgb','thermal']:
  raw=load(OUT/(modality+'_full_predictions.json'));assert len(raw)==1528
  for i,r in enumerate(raw):
   assert r['pair_index']==i and run.sha(r['image_path'])==r['image_sha256'];im=cv2.imread(r['image_path']);assert list(im.shape[:2])==r['orig_shape']
   assert all(b['score']>.005 for b in r['detections'])
  baseline=metric(y,[int(any(b['score']>.5 for b in r['detections'])) for r in raw]);assert baseline==m['baselines'][modality]
 jobs=load(OUT/'roi_receipts.json');assert len(jobs)==receipt['roi_images'];counter=collections.Counter()
 for j in jobs:
  counter[j['pair_index']]+=1;x,y0,X,Y=j['crop'];assert X>x and Y>y0
  for b in j['actual_detections']:
   assert b['score']>.005 and b['crop']==j['crop'] and b['origin']=='crossmodal_roi'
   x1,y1,x2,y2=b['bbox'];assert x1>=x-1e-3 and y1>=y0-1e-3 and x2<=X+1e-3 and y2<=Y+1e-3
 assert max(counter.values())<=6
 recomputed={};changes=[]
 for pool in ['full','augmented']:
  data={k:load(OUT/(k+'_'+pool+'_predictions.json')) for k in ['rgb','thermal']};ev=load(OUT/(pool+'_matched_evidence.json'));rows=list(csv.DictReader((OUT/(pool+'_frame_predictions.csv')).open(encoding='utf-8-sig')));assert len(ev)==len(rows)==1528
  preds={k:[] for k in ['full','oof']}
  for i,row in enumerate(rows):
   assert int(row['pair_index'])==i and int(row['UV_reference'])==reference[i]['UV_reference']
   rgb=unique(data['rgb'][i]['detections']);thermal=unique(data['thermal'][i]['detections'])
   for scope in ['full','oof']:
    cfg=json.loads(row[scope+'_parameters']);ri=[j for j,b in enumerate(rgb) if b['score']>=cfg['tR']];ti=[j for j,b in enumerate(thermal) if b['score']>=cfg['tT']];n=len(ri);nt=len(ti);matches=[]
    if n and nt:
     s=np.array([[iou(rgb[r]['bbox'],thermal[t]['bbox'],A) for t in ti] for r in ri]);w=np.zeros((n+nt,n+nt));w[:n,:nt]=np.where(s>=cfg['gate'],s,-1);aa,bb=linear_sum_assignment(-w)
     for a,b in zip(aa,bb):
      if a<n and b<nt and s[a,b]>=cfg['gate']:
       r,t=ri[a],ti[b];cr=rgb[r]['score'];ct=thermal[t]['score'];rt=cfg['rT']*(min(1.,s[a,b]/.3) if cfg['soft'] else 1.);f=1-(1-cr*cfg['rR']*(1-cfg['u0']))*(1-ct*rt*(1-cfg['u0']));matches.append((r,t,s[a,b],f))
    best=max(matches,key=lambda v:v[3],default=None);mass=best[3] if best else 0.;pred=int(best is not None and mass>cfg['threshold']);preds[scope].append(pred)
    assert pred==int(row[scope+'_fire']);assert abs(mass-float(row[scope+'_mF']))<1e-6
    assert abs(float(row[scope+'_mF'])+float(row[scope+'_mN'])+float(row[scope+'_mU'])-1)<1e-10
    actual=ev[i][scope];assert len(matches)==len(actual['matches'])
    for r,t,s,f in matches:
     matched=[b for b in actual['matches'] if b['r']==r and b['t']==t];assert len(matched)==1;assert abs(s-matched[0]['iou'])<1e-5
    if not best:assert mass==0 and pred==0
  recomputed[pool]={s:metric(y,p) for s,p in preds.items()}
  for scope in ['full','oof']:
   for k,v in recomputed[pool][scope].items():assert abs(v-m['results'][pool][scope][k])<1e-10
  # Independently reconstruct all 32 candidate predictions then verify each fold's FP/TP selection.
  trials=list(csv.DictReader((OUT/(pool+'_all_trials.csv')).open(encoding='utf-8-sig')));assert len(trials)==m['grid_per_pool']
  grid=load(OUT/'pre_evaluation_freeze.json')['grid'];allpred=np.zeros((len(grid),1528),int)
  uu={k:[unique(v['detections']) for v in data[k]] for k in data};ss=[np.array([[iou(r['bbox'],t['bbox'],A) for t in uu['thermal'][i]] for r in uu['rgb'][i]]).reshape(len(uu['rgb'][i]),len(uu['thermal'][i])) for i in range(1528)]
  for k,c in enumerate(grid):
   for i in range(1528):
    rr=[j for j,b in enumerate(uu['rgb'][i]) if b['score']>=c['tR']];tt=[j for j,b in enumerate(uu['thermal'][i]) if b['score']>=c['tT']];n=len(rr);nt=len(tt);mass=0.
    if n and nt:
     s=ss[i][np.ix_(rr,tt)];w=np.zeros((n+nt,n+nt));w[:n,:nt]=np.where(s>=c['gate'],s,-1);aa,bb=linear_sum_assignment(-w)
     for a,b in zip(aa,bb):
      if a<n and b<nt and s[a,b]>=c['gate']:
       cr=uu['rgb'][i][rr[a]]['score'];ct=uu['thermal'][i][tt[b]]['score'];rt=c['rT']*(min(1.,s[a,b]/.3) if c['soft'] else 1.);mass=max(mass,1-(1-cr*c['rR']*(1-c['u0']))*(1-ct*rt*(1-c['u0'])))
    allpred[k,i]=mass>c['threshold']
   trialmetric=metric(y,allpred[k].tolist())
   for key in ['TP','FP','FN','TN']:assert trialmetric[key]==int(trials[k][key])
  group=np.array(run.P['all_groups']);yy=np.array(y)
  for f in m['results'][pool]['folds']:
   train=np.flatnonzero(group!=f['fold']);test=np.flatnonzero(group==f['fold']);assert train.tolist()==f['train_indices'] and test.tolist()==f['test_indices'];assert not set(group[train])&set(group[test])
   counts=[(int(sum(allpred[k,train][yy[train]==0])), -int(sum(allpred[k,train][yy[train]==1])), k) for k in range(len(grid))];selected=min(counts)[2];assert selected==f['index'] and grid[selected]==f['config']
   for i in test:assert int(rows[i]['oof_fire'])==allpred[selected,i]
 # Coverage and actual new-frame gains.
 raw={k:load(OUT/(k+'_full_predictions.json')) for k in ['rgb','thermal']};aug={k:load(OUT/(k+'_augmented_predictions.json')) for k in ['rgb','thermal']}
 for i in range(1528):
  def maximum(data):
   rr=unique(data['rgb'][i]['detections']);tt=unique(data['thermal'][i]['detections']);return max((iou(r['bbox'],t['bbox'],A) for r in rr for t in tt),default=0.)
  b,a=maximum(raw),maximum(aug)
  if b<.001 and a>=.001:changes.append(dict(pair_index=i,UV_reference=y[i],before_iou=b,after_iou=a))
 for j in changes[:3]:
  i=j['pair_index'];im=cv2.imread(raw['rgb'][i]['image_path']);th=cv2.imread(raw['thermal'][i]['image_path']);overlay=im.copy()
  for b in aug['rgb'][i]['detections']:cv2.rectangle(overlay,tuple(np.int32(b['bbox'][:2])),tuple(np.int32(b['bbox'][2:])),(0,255,0),1)
  for b in aug['thermal'][i]['detections']:cv2.polylines(overlay,[np.int32(run.project(run.old.corners(b['bbox']),A))],True,(0,0,255),1)
  cv2.imwrite(str(OUT/f'roi_overlay_{i}.jpg'),run.montage([im,th,overlay]))
 run.save('roi_new_coverage.json',changes)
 point_coverage={}
 for modality,key in [('rgb','rgb_points'),('thermal','thermal_points')]:
  hull=cv2.convexHull(np.float32(run.P['geometry'][key]));H,W=raw[modality][0]['orig_shape'];inside=total=0
  for r in raw[modality]:
   for b in r['detections']:
    x,y0,X,Y=b['bbox'];total+=1;inside+=cv2.pointPolygonTest(hull,((x+X)/2,(y0+Y)/2),False)>=0
  point_coverage[modality]=dict(hull_area_px=float(cv2.contourArea(hull)),image_area_px=H*W,field_fraction=float(cv2.contourArea(hull)/(H*W)),candidate_centers_inside=int(inside),candidate_centers_total=total,original_dimensions=[H,W])
 run.save('point_coverage.json',point_coverage)
 path=run.ROOT/'reports/双基线约束_空间DS验证报告.docx'
 if path.exists():
  with zipfile.ZipFile(path) as z:assert z.testzip() is None
  rd=Document(path);assert len(rd.tables)>=1
  table=rd.tables[0];assert len(table.rows)==9
  expected=[m['baselines']['rgb'],m['baselines']['thermal']]
  for pool in ['full','augmented']:expected.extend([m['results'][pool]['full'],m['results'][pool]['oof'],m['results'][pool]['uv_secondary']])
  for row,v in zip(table.rows[1:],expected):
   for cell,k in zip(row.cells[1:],['TP','FP','FN','TN','FPR','FNR','precision','recall','F1','accuracy']):assert float(cell.text)==v[k]
 report=dict(verified_utc=run.now(),all_1528_rows_independently_recomputed=True,polygon_hungarian_BPA_recomputed=True,source_images_verified=3056,protected_files_verified=len(initial['protected']),roi_forwards=len(jobs),maximum_roi_per_frame=max(counter.values()),metrics=recomputed,new_coverage=changes,point_coverage=point_coverage,zero_of_240_two_sided_95_upper=float(beta.ppf(.975,1,240)),targets=m['targets'],report_crc_verified=path.exists(),primary_goal_pass=m['targets']['augmented']['oof']['pass_stricter'])
 run.save('verification.json',report);run.save('artifact_sha256.json',{p.name:run.sha(p) for p in OUT.iterdir() if p.is_file() and p.name!='artifact_sha256.json' and p.suffix!='.log'});print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':verify()
