from pathlib import Path
import sys,json,csv,hashlib,datetime,collections,itertools,time,zipfile
import numpy as np
import cv2
OUT=Path(__file__).resolve().parent;ROOT=OUT.parent.parent;PREV=ROOT/'fusion/corrected_uv_optimized_ds'
sys.path.insert(0,str(PREV))
import run_optimization as old
from optimized_core import frame,select
from dual_core import project,inverse,offset,montage
load=old.load;sha=old.sha;now=old.now
P=load(PREV/'pre_search_protocol.json');A=np.array(P['geometry']['affine_matrix_2x3'])
def save(name,obj):
 with (OUT/name).open('w',encoding='utf-8') as f:json.dump(obj,f,ensure_ascii=False,indent=2,allow_nan=False)
def csvsave(name,rows):
 with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def diagnose():
 rows=old.parse();r=load(PREV/'rgb_predictions.json');t=load(PREV/'thermal_predictions.json');ds=[];counts=collections.Counter();miss=[]
 protected={str(f):sha(f) for f in PREV.rglob('*') if f.is_file()}
 for m,v in P['models'].items():assert sha(v['source'])==v['sha256'];protected[v['source']]=sha(v['source'])
 protected[str(old.SOURCE)]=sha(old.SOURCE);protected[P['geometry']['source']]=sha(P['geometry']['source'])
 for i,(rr,tt,row) in enumerate(zip(r,t,rows)):
  s=old.geom(rr['detections'],tt['detections'],A);best=float(s.max()) if s.size else 0
  reason='no_both' if not rr['detections'] and not tt['detections'] else 'no_RGB' if not rr['detections'] else 'no_thermal' if not tt['detections'] else 'zero_overlap' if best==0 else 'tiny_overlap' if best<.01 else 'covered'
  if row['UV_reference']==1 and best<.01:counts[reason]+=1;miss.append(i)
  ds.append(dict(pair_index=i,UV_reference=row['UV_reference'],rgb_count=len(rr['detections']),thermal_count=len(tt['detections']),max_iou=best,reason=reason,rgb_shape=rr['orig_shape'],thermal_shape=tt['orig_shape']))
 csvsave('prior_candidate_diagnosis.csv',ds)
 save('initial_diagnosis.json',dict(prior_positive_uncovered=len(miss),reasons=dict(counts),missing_indices=miss,coordinate_dimensions=dict(rgb=dict(collections.Counter(str(x['orig_shape']) for x in r)),thermal=dict(collections.Counter(str(x['orig_shape']) for x in t))),geometry=P['geometry'],protected=protected,source_sha256=sha(old.SOURCE)))
 selected=[]
 for reason in counts:
  ix=[i for i in miss if ds[i]['reason']==reason];selected+=ix[:1]+ix[-1:]
 selected=list(dict.fromkeys(selected))[:6]
 for i in selected:
  im=cv2.imread(r[i]['image_path']);th=cv2.imread(t[i]['image_path']);overlay=im.copy()
  for b in r[i]['detections']:cv2.rectangle(overlay,tuple(np.int32(b['bbox'][:2])),tuple(np.int32(b['bbox'][2:])),(0,255,0),1)
  for b in t[i]['detections']:cv2.polylines(overlay,[np.int32(project(old.corners(b['bbox']),A))],True,(0,0,255),1)
  cv2.putText(overlay,f"pair {i} {ds[i]['reason']}",(5,15),0,.4,(255,255,255),1)
  cv2.imwrite(str(OUT/f'prior_overlay_{i}.jpg'),montage([im,th,overlay]))
 # Registration audit, finite TWO variants, first three geometrically chosen missing frames, no UV/boxes in estimator.
 audits=[]
 for i in selected[:3]:
  im=cv2.imread(r[i]['image_path'],0);th=cv2.imread(t[i]['image_path'],0)
  th=cv2.copyMakeBorder(th,0,max(0,im.shape[0]-th.shape[0]),0,max(0,im.shape[1]-th.shape[1]),cv2.BORDER_CONSTANT,value=0)
  def grad(x):
   x=x.astype(np.float32)/255.;return cv2.magnitude(cv2.Sobel(x,cv2.CV_32F,1,0),cv2.Sobel(x,cv2.CV_32F,0,1))
  for variant in ['gradient','inverted_intensity']:
   x,y=(grad(im),grad(th)) if variant=='gradient' else (im.astype(np.float32)/255.,1-th.astype(np.float32)/255.)
   yy,xx=np.indices(im.shape);train=((xx//32+yy//32)%2==0).astype(np.uint8)*255;valid=train==0
   def quality(B):
    wy=cv2.warpAffine(y,B,(im.shape[1],im.shape[0]),flags=cv2.INTER_LINEAR|cv2.WARP_INVERSE_MAP)
    ok=valid&(cv2.warpAffine(np.ones_like(y),B,(im.shape[1],im.shape[0]),flags=cv2.WARP_INVERSE_MAP)>.99)
    return float(np.corrcoef(x[ok],wy[ok])[0,1]) if sum(ok.ravel())>100 else -1.
   initial=inverse(A).astype(np.float32);before=quality(initial)
   try:
    cc,B=cv2.findTransformECC(x,y,initial.copy(),cv2.MOTION_AFFINE,(cv2.TERM_CRITERIA_COUNT|cv2.TERM_CRITERIA_EPS,100,1e-5),train,5)
    after=quality(B);delta=float(np.max(np.linalg.norm(project(old.corners([0,0,320,320]),inverse(B))-project(old.corners([0,0,320,320]),A),axis=1)))
    accepted=after>.3 and after>before+.03 and delta<20
    audits.append(dict(pair_index=i,variant=variant,before=before,after=after,corner_displacement=delta,accepted=bool(accepted),candidate_matrix=inverse(B).tolist(),note='Smoke-only no full-field independent correspondence validation; never promoted automatically'))
   except cv2.error as e:audits.append(dict(pair_index=i,variant=variant,accepted=False,error=str(e)))
 save('registration_smoke.json',audits)
 save('protocol.json',dict(created_utc=now(),source_sha256=old.EXPECTED,primary_UV=False,reference_exposure='All 1528 previously seen; OOF internal exploratory only',groups=P['all_groups'],group_summary=P['groups'],models={k:{'source':v['source'],'sha256':v['sha256']} for k,v in P['models'].items()},geometry=A.tolist(),registration_policy='Retain seven-point affine; ECC smoke never automatically promoted without independent scene validation',conf=.005,imgsz=320,rect=False,roi=dict(trigger='Every frame, top3 genuine detections per modality; opposite modality projected crop; cap6; no UV triggers',padding=.5,minimum_crop_size=32,score_strictly_greater=.005),grid=dict(gate=[.001,.005,.01,.02],tR=[.005,.1,.3,.5],tT=[.005,.05,.1],rR=[1.],rT=[.1,.3,.6],soft=[False,True],u0=[0.],threshold=[.3,.4,.5,.6]),selection='Calibration FP ascending TP descending deterministic order',target='FP=0 FN<=25',tiny_iou_warning='IoU .001 is weak nonzero consistency, not strong overlap',baseline='Original full-image320 same checkpoint score>.5; no ROI baseline substitution'))
 print('DIAGNOSIS',counts,'missing',len(miss),flush=True)
def predict(model,ims):
 shapes=[];h=model.model.register_forward_pre_hook(lambda m,args:shapes.append(list(args[0].shape)))
 res=list(model.predict(ims,imgsz=320,rect=False,conf=.005,device=0,half=False,batch=32,classes=[0],max_det=300,iou=.7,augment=False,verbose=False,save=False));h.remove()
 assert shapes and all(s[-2:]==[320,320] for s in shapes)
 out=[]
 for result,im in zip(res,ims):
  np.testing.assert_array_equal(result.orig_img,im)
  out.append([dict(bbox=[x,y,X,Y],score=c,class_id=int(k),origin='full') for x,y,X,Y,c,k in result.boxes.data.cpu().tolist() if c>.005])
 assert len(out)==len(ims)
 return out,shapes
def crop_for(box,B,shape):
 q=project(old.corners(box),B);lo=q.min(0);hi=q.max(0);center=(lo+hi)/2;half=np.maximum((hi-lo),32)/2*2
 lo=np.maximum(np.floor(center-half),0).astype(int);hi=np.minimum(np.ceil(center+half),[shape[1],shape[0]]).astype(int)
 return [int(lo[0]),int(lo[1]),int(hi[0]),int(hi[1])] if np.all(hi-lo>=4) else None
def dedup(boxes):
 out=[]
 for b in sorted(boxes,key=lambda b:-b['score']):
  if all(old.polygon_iou(old.corners(b['bbox']),old.corners(c['bbox']))<.8 for c in out):out.append(b)
 return out
def fresh_roi():
 import torch
 from ultralytics import YOLO
 assert torch.cuda.is_available();rows=old.parse();models={k:YOLO(v['source']) for k,v in P['models'].items()};previous={k:load(PREV/(k+'_predictions.json')) for k in models};shapes=[];start=now();tic=time.perf_counter()
 # Real THREE paired full-image smoke and six symmetric ROI proposals before all full forwards.
 smoke=[]
 for m in models:
  ims=[cv2.imread(x['image_path']) for x in previous[m][:3]];z,s=predict(models[m],ims);shapes+=s
  for i,boxes in enumerate(z):
   B=inverse(A) if m=='rgb' else A;target='thermal' if m=='rgb' else 'rgb';source=previous[target][i];b=boxes[0]['bbox'] if boxes else [100,100,180,180];crop=crop_for(b,B,source['orig_shape']);im=cv2.imread(source['image_path']);x,y,X,Y=crop
   zz,ss=predict(models[target],[im[y:Y,x:X]]);shapes+=ss;smoke.append(dict(pair_index=i,source=m,target=target,crop=crop,proposal_is_not_evidence=True,real_detections=[dict(b,bbox=offset(b['bbox'],crop)) for b in zz[0]]))
 save('fresh_roi_smoke.json',smoke)
 full={}
 for m in models:
  out=[]
  for first in range(0,len(rows),32):
   originals=previous[m][first:first+32];ims=[cv2.imread(x['image_path']) for x in originals];z,s=predict(models[m],ims);shapes+=s
   for v,boxes,im in zip(originals,z,ims):
    assert sha(v['image_path'])==v['image_sha256'] and list(im.shape[:2])==v['orig_shape'];out.append(dict(v,detections=boxes))
  full[m]=out;save(m+'_full_predictions.json',out);print('FULL',m,len(out),flush=True)
 augmented={k:[dict(v,detections=list(v['detections'])) for v in full[k]] for k in models};receipts=[]
 for target in models:
  source='thermal' if target=='rgb' else 'rgb';B=A if target=='rgb' else inverse(A);jobs=[];crops=[]
  def flush():
   if not crops:return
   z,s=predict(models[target],crops);shapes.extend(s)
   for job,boxes in zip(jobs,z):
    i,crop,anchor=job;actual=[dict(b,bbox=offset(b['bbox'],crop),origin='crossmodal_roi',crop=crop,source_anchor=anchor) for b in boxes];augmented[target][i]['detections'].extend(actual);receipts.append(dict(pair_index=i,target=target,crop=crop,source_anchor=anchor,actual_detections=actual))
   jobs.clear();crops.clear()
  for i in range(len(rows)):
   im=cv2.imread(full[target][i]['image_path'])
   for anchor,b in sorted(enumerate(full[source][i]['detections']),key=lambda v:-v[1]['score'])[:3]:
    crop=crop_for(b['bbox'],B,im.shape); 
    if crop:
     x,y,X,Y=crop;jobs.append((i,crop,anchor));crops.append(im[y:Y,x:X])
     if len(crops)==32:flush()
  flush()
  for row in augmented[target]:row['detections']=dedup(row['detections'])
  save(target+'_augmented_predictions.json',augmented[target]);print('ROI',target,'jobs',sum(x['target']==target for x in receipts),flush=True)
 save('roi_receipts.json',receipts)
 save('forward_receipt.json',dict(start_utc=start,end_utc=now(),seconds=time.perf_counter()-tic,full_images=3056,roi_images=len(receipts),smoke_full_images=6,smoke_roi_images=6,gpu=torch.cuda.get_device_name(0),shapes=shapes,model_sha256={k:sha(v['source']) for k,v in P['models'].items()},source_sha256=sha(old.SOURCE),conf=.005,all_orig_images_hash_verified=True))
def candidate_bounds():
 y=np.array([r['UV_reference'] for r in old.parse()]);bounds=[];reasons={}
 for kind in ['full','augmented']:
  rgb=load(OUT/('rgb_'+kind+'_predictions.json'));thermal=load(OUT/('thermal_'+kind+'_predictions.json'));maximum=[];counts=collections.Counter()
  for i,(r,t) in enumerate(zip(rgb,thermal)):
   rr=dedup(r['detections']);tt=dedup(t['detections']);s=old.geom(rr,tt,A);v=float(s.max()) if s.size else 0.;maximum.append(v)
   if y[i]==1 and v<.001:counts['no_RGB' if not rr else 'no_thermal' if not tt else 'zero_overlap' if v==0 else 'tiny_overlap']+=1
  s=np.array(maximum);reasons[kind]=dict(counts)
  for gate in [.001,.005,.01,.02]:bounds.append(dict(pool=kind,gate=gate,positive_covered=int(sum((s>=gate)&(y==1))),positive_uncovered=int(sum((s<gate)&(y==1))),negative_covered=int(sum((s>=gate)&(y==0)))))
 stamp=load(OUT/'pre_optimization_candidate_bounds.json')['utc'] if (OUT/'pre_optimization_candidate_bounds.json').exists() else now()
 save('pre_optimization_candidate_bounds.json',dict(utc=stamp,rechecked_utc=now(),bounds=bounds,uncovered_positive_reasons=reasons));print('BOUNDS',bounds,'REASONS',reasons,flush=True)
def evaluate():
 rows=old.parse();y=np.array([r['UV_reference'] for r in rows]);groups=np.array(P['all_groups']);proto=load(OUT/'protocol.json');grid=[dict(zip(proto['grid'],v)) for v in itertools.product(*proto['grid'].values())]
 bound=load(OUT/'pre_optimization_candidate_bounds.json');limited=all(b['positive_covered']<1263 for b in bound['bounds'])
 if limited:grid=[c for c in grid if c['tR'] in [.3,.5] and c['tT'] in [.005,.05] and c['gate'] in [.001,.01] and c['rT'] in [.1,.3] and not c['soft'] and c['threshold'] in [.4,.5]]
 save('pre_evaluation_freeze.json',dict(utc=now(),protocol_sha256=sha(OUT/'protocol.json'),grid=grid,coverage_impossible=limited,reason='Candidate floor exceeds target: skip exhaustive futile optimization, finite 32-grid diagnostic only; no claim of repaired recall'))
 datasets={kind:{m:load(OUT/(m+'_'+kind+'_predictions.json')) for m in ['rgb','thermal']} for kind in ['full','augmented']};bounds=[];best_by_kind={}
 for data in datasets.values():
  for modality in data.values():
   for row in modality:row['detections']=dedup(row['detections'])
 for kind,data in datasets.items():
  ious=[old.geom(r['detections'],t['detections'],A) for r,t in zip(data['rgb'],data['thermal'])];maxs=np.array([s.max() if s.size else 0 for s in ious])
  for gate in proto['grid']['gate']:bounds.append(dict(pool=kind,gate=gate,positive_covered=int(sum((maxs>=gate)&(y==1))),positive_uncovered=int(sum((maxs<gate)&(y==1))),negative_covered=int(sum((maxs>=gate)&(y==0)))))
  preds=np.zeros((len(grid),len(rows)),bool);scores=np.zeros_like(preds,dtype=float)
  # Compute finite grids via eligible one-to-one matches once per score/gate combination.
  keys=list(dict.fromkeys((c['tR'],c['tT'],c['gate']) for c in grid))
  for no,key in enumerate(keys):
   tr,tt,gate=key;cfg=dict(grid[0],tR=tr,tT=tt,gate=gate);matches=[]
   for i in range(len(rows)):
    o=frame(data['rgb'][i]['detections'],data['thermal'][i]['detections'],ious[i],cfg)
    matches.extend((i,b['rgb_score'],b['thermal_score'],b['iou']) for b in o['matches'])
   a=np.array(matches);ids=a[:,0].astype(int) if len(a) else np.array([],int)
   for k,c in enumerate(grid):
    if (c['tR'],c['tT'],c['gate'])!=key:continue
    if len(a):
     rt=c['rT']*(np.clip(a[:,3]/.3,0,1) if c['soft'] else 1);s=1-(1-a[:,1])*(1-rt*a[:,2]);np.maximum.at(scores[k],ids,s)
    preds[k]=scores[k]>c['threshold']
  fullidx=select(preds,y,np.arange(len(y)));oof=np.zeros(len(y),bool);assignment=np.zeros(len(y),int);folds=[]
  for g in np.unique(groups):
   train=np.flatnonzero(groups!=g);test=np.flatnonzero(groups==g);k=select(preds,y,train);freeze=dict(utc=now(),fold=int(g),config=grid[k],index=k,train_indices=train.tolist(),test_indices=test.tolist(),calibration=old.metrics(y[train].tolist(),preds[k,train].astype(int).tolist()))
   save(kind+'_fold_'+str(g)+'_freeze.json',freeze);oof[test]=preds[k,test];assignment[test]=k;folds.append(dict(**freeze,heldout=old.metrics(y[test].tolist(),oof[test].astype(int).tolist())))
  best_by_kind[kind]=dict(full=old.metrics(y.tolist(),preds[fullidx].astype(int).tolist()),oof=old.metrics(y.tolist(),oof.astype(int).tolist()),full_config=grid[fullidx],folds=folds)
  details=[];evidence=[]
  for i,row in enumerate(rows):
   outputs={scope:frame(data['rgb'][i]['detections'],data['thermal'][i]['detections'],ious[i],grid[k]) for scope,k in [('full',fullidx),('oof',assignment[i])]}
   z=dict(pair_index=i,**row,proxy_group=int(groups[i]))
   for scope,o in outputs.items():
    assert o['fire']==int(preds[fullidx,i] if scope=='full' else oof[i]);z.update({scope+'_'+key:o[key] for key in ['fire','mF','mN','mU','K']});b=o['selected'];z[scope+'_iou']=b['iou'] if b else 0;z[scope+'_reason']='matched_score' if b else 'no_real_dual_match';z[scope+'_parameters']=json.dumps(grid[fullidx if scope=='full' else assignment[i]])
   z['uv_secondary_fire']=int(outputs['full']['mF']*(.5+.5*y[i])>.5 and outputs['full']['selected'] is not None)
   evidence.append(dict(pair_index=i,**outputs));details.append(z)
  csvsave(kind+'_frame_predictions.csv',details);save(kind+'_matched_evidence.json',evidence)
  csvsave(kind+'_all_trials.csv',[dict(index=k,**c,**old.metrics(y.tolist(),preds[k].astype(int).tolist())) for k,c in enumerate(grid)])
  best_by_kind[kind]['uv_secondary']=old.metrics(y.tolist(),[r['uv_secondary_fire'] for r in details]);print('RESULT',kind,best_by_kind[kind]['full'],best_by_kind[kind]['oof'],flush=True)
 full=datasets['full'];baselines={m:old.metrics(y.tolist(),[int(any(b['score']>.5 for b in r['detections'])) for r in full[m]]) for m in full}
 assert [baselines['rgb'][k] for k in ['TP','FP','FN','TN']]==[1262,0,26,240]
 assert [baselines['thermal'][k] for k in ['TP','FP','FN','TN']]==[716,180,572,60]
 targets={kind:{scope:dict(pass_stricter=bool(v[scope]['FP']==0 and v[scope]['FN']<=25),FP=v[scope]['FP'],FN=v[scope]['FN']) for scope in ['full','oof']} for kind,v in best_by_kind.items()}
 save('candidate_bounds.json',bounds);save('metrics.json',dict(baselines=baselines,results=best_by_kind,targets=targets,source_sha256=sha(old.SOURCE),independent_test=False,grid_per_pool=len(grid)))
def report():
 from docx import Document
 from docx.shared import Pt,Inches
 from docx.oxml.ns import qn
 m=load(OUT/'metrics.json');d=load(OUT/'initial_diagnosis.json');receipt=load(OUT/'forward_receipt.json');bounds=load(OUT/'candidate_bounds.json');doc=Document();style=doc.styles['Normal'];style.font.name='宋体';style.font.size=Pt(10.5);style.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
 doc.add_heading('双基线约束：空间IoU与D-S验证报告',0)
 doc.add_paragraph('结论：'+json.dumps(m['targets'],ensure_ascii=False)+'。主验收为留组OOF同时FP=0、FN≤25；全数据最佳不能替代OOF。本报告仅呈现实际执行结果，不承诺达到不存在的性能。')
 sections=[('1 目标与参考','当前paper_sample.txt第四列UV参考共1288正例、240负例。目标FPR≤两单模态且FNR严格小于两单模态，即FP=0且FN≤25。全部1528组已反复查看，六组按既有时间间隔规则冻结；OOF仅内部探索，非独立泛化。源SHA256：'+m['source_sha256']),('2 缺失候选归因','旧>0.01缓存中IoU≥0.01无法覆盖78正例，具体：'+json.dumps(d['reasons'],ensure_ascii=False)+'。真实双检测覆盖1282，不可将无重叠与无检测混为一谈。原始尺寸：'+json.dumps(d['coordinate_dimensions'])+'。七点位于狭窄纵向带，留一RMSE=5.771805像素，不支持全视场或跨场景有效性保证。框均为原图坐标，不把320坐标错误放大。'),('3 实际方法与有限替代','原始YOLO26n best.pt和best-2.pt不变（不是PPYOLOv2），全图3056张真实CUDA前向，conf>0.005，imgsz=320、rect=False，实测输入320×320；端到端检测绕过传统NMS。每帧每模态真实前三框投影至对方原图，以0.5框宽高边距且最小原框32像素形成裁剪，至多6裁剪/帧。裁剪真实模型仍320输入，框仅加原图裁剪偏移，投影本身不是检测证据。所有帧统一触发，不输入UV或图像编号判火。ROI作同模态IoU0.8去重；无ROI控制也同样去重后比较。'),('4 配准与质量','在代表图像上进行两种预定结构方案：梯度幅值ECC和反转强度ECC。估计仅使用图像像素，初始七点仿射；棋盘空间掩膜留出相关性验证，记录收敛、改进和四角位移。少数帧ECC烟测不足认证全场景，因此不擅自推广新矩阵，也不根据UV或火框对齐拟合。详见registration_smoke.json。'),('5 BPA及判定','先按分数筛框，用带dummy匈牙利最大IoU一对一匹配。必须真实双模态且IoU≥0.001/0.005/0.01/0.02之一。0.001仅极弱非零重叠，不能称强一致。mF=rc、mN=0、mΘ=1−rc；RGB r=1，热成像r=0.1/0.3/0.6并可乘clip(IoU/0.3)。D-S为正类简单支持，K=0且mF=1−(1−mRF)(1−mTF)，不是冲突消解或概率校准。缺模态/未匹配保持未知质量且不判火；UV也不能补造空间证据。分数门槛、可靠性和IoU有限格点仅按其余组FP优先、TP次优选择，冻结后评估本组。'),('6 可解释性及风险','修改可靠性或阈值不能弥补真实候选上界。UV乘法仅次要控制，λ=0.5、严格>0.5能利用参考本身清除负例，属于循环一致性，不构成达标证据。当前训练/验证源与测试图像已知存在散列重合，不能证明历史权重训练成员，也不能声称独立测试。0/240误报不保证部署误报为零；双侧95% Clopper–Pearson上界在verification.json记录。未进行盲目训练或替换更强模型以弱化基线。')]
 for heading,text in sections:doc.add_heading(heading,1);doc.add_paragraph(text)
 doc.add_heading('7 实测指标',1)
 cov=load(OUT/'pre_optimization_candidate_bounds.json');points=load(OUT/'point_coverage.json');visual=load(OUT/'visual_audit.json');ver=load(OUT/'verification.json')
 doc.add_paragraph('候选上界先于参数选择确定：降低全图候选阈值及7333次真实ROI前向后，最佳仅覆盖1213/1288正例，仍有75个空间未覆盖（'+json.dumps(cov['uncovered_positive_reasons']['augmented'],ensure_ascii=False)+'），故FN理论下限75>25，原定大格点检索被终止，仅执行预先冻结32项诊断配置/候选池，六折各自按训练组选择。不能据此声称达成强目标。')
 doc.add_paragraph('七点凸包覆盖RGB视场'+str(points['rgb']['field_fraction'])+'、热视场'+str(points['thermal']['field_fraction'])+'；RGB候选中心仅'+str(points['rgb']['candidate_centers_inside'])+'/'+str(points['rgb']['candidate_centers_total'])+'位于凸包内，热成像仅'+str(points['thermal']['candidate_centers_inside'])+'/'+str(points['thermal']['candidate_centers_total'])+'。原图RGB320×320、热成像192×256；显示拼图只补黑边不改变坐标。上述几何跨域风险不能靠拟合误差小来排除。')
 doc.add_paragraph('人工视觉审阅：60号原热框明显映射在RGB候选上方；490/491号人物遮挡，缺乏清晰RGB火焰；817号RGB有亮火焰样区域但原映射热框偏高。新覆盖491、503、817仅证明新增候选重叠；部分ROI框覆盖人体或工作台，不能认证真实火焰对应。部分热源图已烧录红框及Fire:0.660/0.758文本，构成检测域污染和已有模型标注风险。本次未清除或改写原图，也不把烧录框作为新真值。详见visual_audit.json。')
 doc.add_paragraph('零误报分支0/240的双侧95% Clopper–Pearson上界='+str(ver['zero_of_240_two_sided_95_upper'])+'，不表示部署零风险。独立复算全部1528行空间匹配、BPA及折选择，同时检查3056图像SHA、原产物和权重不变。')
 t=doc.add_table(rows=1,cols=11);t.style='Table Grid';fields=['TP','FP','FN','TN','FPR','FNR','precision','recall','F1','accuracy']
 for c,s in zip(t.rows[0].cells,['方法']+fields):c.text=s
 entries=[('RGB原始320 >0.5',m['baselines']['rgb']),('热成像原始320 >0.5',m['baselines']['thermal'])]
 for pool,v in m['results'].items():
  entries.extend([(pool+' 全拟合描述性',v['full']),(pool+' 留组OOF',v['oof']),(pool+' UV循环控制',v['uv_secondary'])])
 for name,v in entries:
  cells=t.add_row().cells;cells[0].text=name
  for c,k in zip(cells[1:],fields):c.text=str(v[k])
 doc.add_heading('8 候选逻辑上界与分组',1)
 for b in bounds:doc.add_paragraph(json.dumps(b,ensure_ascii=False))
 for pool,v in m['results'].items():
  doc.add_paragraph(pool+'全拟合参数：'+json.dumps(v['full_config']))
  for f in v['folds']:doc.add_paragraph(pool+'组'+str(f['fold'])+' '+json.dumps(f['heldout'])+' 参数'+json.dumps(f['config']))
 doc.add_paragraph('真实执行回执：'+json.dumps({k:receipt[k] for k in ['start_utc','end_utc','seconds','full_images','roi_images','gpu','model_sha256']},ensure_ascii=False))
 for p in sorted(OUT.glob('prior_overlay_*.jpg'))[:3]:doc.add_picture(str(p),width=Inches(6));doc.add_paragraph(p.name+'：左RGB，中热成像，右RGB真实框（绿）及仿射热框（红）。仅定性检查，非人工真值。')
 for p in sorted(OUT.glob('roi_overlay_*.jpg'))[:3]:doc.add_picture(str(p),width=Inches(6));doc.add_paragraph(p.name+'：新增ROI候选重叠，不是独立确认的真实火焰对应。')
 path=ROOT/'reports/双基线约束_空间DS验证报告.docx';doc.save(path)
 with zipfile.ZipFile(path) as z:assert z.testzip() is None
 rd=Document(path);txt='\n'.join(p.text for p in rd.paragraphs)+'\n'+'\n'.join('|'.join(c.text for c in r.cells) for t in rd.tables for r in t.rows);(OUT/'report_readback.txt').write_text(txt,encoding='utf-8');assert 'FN≤25' in txt and '留组OOF' in txt
 save('report_receipt.json',dict(path=str(path),sha256=sha(path),crc_ok=True,tables=len(rd.tables)))
if __name__=='__main__':
 action=sys.argv[1];globals()[action]()
