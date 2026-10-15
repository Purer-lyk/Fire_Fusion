"""Execute real predictions through user's C++ rules; never manufacture labels."""
import csv
import hashlib
import json
import re
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
import cv2
import numpy as np
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import core as c

ROOT=Path('C:/lyk/fire/qualityDetection')
OUT=ROOT/'fusion'
REPORT=ROOT/'reports/PPYOLOv2_DS融合实际测试报告.docx'
SOURCE=Path('C:/lyk/paper_data/paper_data_/paper_data')
CPP=Path('C:/lyk/paper/bpa_test')
ANNOTATOR=Path('C:/lyk/pipline/fire_annotation/data/annotations.json')
METHODS=['rgb_baseline','thermal_baseline','legacy_fixed_01','classic_ds','candidate_geometry_repeated','recommended_spatial_then_ds']
NAMES={'rgb_baseline':'RGB置信度基线','thermal_baseline':'热成像置信度基线','legacy_fixed_01':'原C++ fuse3：固定0.1两次空间规则（主要结果）','classic_ds':'经典D-S三证据','candidate_geometry_repeated':'候选几何s两次空间规则（探索扩展）','recommended_spatial_then_ds':'RGB/热空间+UV经典D-S（可选建议变体）'}

def sha(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        for b in iter(lambda:f.read(1048576),b''): h.update(b)
    return h.hexdigest()

def save(name,obj):
    (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')

def write_csv(name,rows,fields=None):
    with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=fields or list(rows[0])); w.writeheader(); w.writerows(rows)

def stats(v):
    a=np.asarray(v,float)
    return dict(min=float(a.min()),p25=float(np.quantile(a,.25)),median=float(np.median(a)),p75=float(np.quantile(a,.75)),max=float(a.max()),mean=float(a.mean()))

def calibration():
    path=CPP/'main/homo_main.cpp'
    text=re.sub(r'//[^\n]*','',path.read_text(encoding='utf-8'))
    pts=[]
    for name in ['pts1','pts2']:
        block=re.search(r'vector<Point2f>\s+'+name+r'\s*=\s*\{(.*?)\};',text,re.S).group(1)
        pts.append(np.array([(float(x),float(y)) for x,y in re.findall(r'Point2f\(([^,]+),\s*([^\)]+)\)',block)],float))
    rgb,thermal=pts
    assert rgb.shape==thermal.shape==(7,2)
    cv2.setRNGSeed(7)
    h,mask=cv2.findHomography(thermal,rgb,cv2.RANSAC,3.0)
    assert h is not None and np.isfinite(h).all()
    def errors(h,src,dst):
        return np.linalg.norm(cv2.perspectiveTransform(src.reshape(-1,1,2),h).reshape(-1,2)-dst,axis=1)
    fit=errors(h,thermal,rgb)
    loo=[]
    for i in range(7):
        keep=np.arange(7)!=i
        hi,mi=cv2.findHomography(thermal[keep],rgb[keep],cv2.RANSAC,3.0)
        loo.append(float(errors(hi,thermal[i:i+1],rgb[i:i+1])[0]))
    data=dict(status='candidate/exploratory, unverified across test scenes',source=str(path),source_sha256=sha(path),direction='thermal pts2 -> RGB pts1',thermal_points=thermal.tolist(),rgb_points=rgb.tolist(),H=h.tolist(),estimator='cv2.findHomography RANSAC threshold=3 original RGB pixels, rng seed=7',inlier_mask=mask.ravel().tolist(),fit_errors_px=fit.tolist(),fit_rmse_px=float(np.sqrt(np.mean(fit**2))),fit_max_px=float(fit.max()),leave_one_out_errors_px=loo,leave_one_out_rmse_px=float(np.sqrt(np.mean(np.array(loo)**2))),leave_one_out_max_px=max(loo),validation_note='Seven points are concentrated in a narrow y band; fit/LOO only measure these points, not full field or independent scenes. No rectangle offsets applied.')
    save('candidate_calibration.json',data)
    return h,data

def geometry(rgb,thermal,h):
    rd=sorted(rgb['detections'],key=lambda d:d['score'],reverse=True)
    td=sorted(thermal['detections'],key=lambda d:d['score'],reverse=True)
    if not rd or not td: return 0.,dict(status='missing_detection',rgb_detection=None,thermal_detection=None,top_score_iou=0.)
    # Confidence evidence is frame maximum; geometry anchors THAT RGB box, then best IoU thermal box.
    a=c.box_polygon(rd[0]['bbox_xyxy'])
    best=-1.; best_j=None; invalid=0; top_iou=0.
    for j,d in enumerate(td):
        try: s=c.polygon_iou(a,c.project_polygon(d['bbox_xyxy'],h))
        except ValueError: invalid+=1; continue
        if j==0: top_iou=s
        if s>best: best=s; best_j=j
    if best_j is None: return 0.,dict(status='invalid_projection',rgb_detection=rd[0],thermal_detection=None,top_score_iou=top_iou,invalid_projections=invalid)
    return best,dict(status='candidate_geometry',rgb_detection=rd[0],thermal_detection=td[best_j],thermal_score_rank=best_j,top_score_iou=top_iou,invalid_projections=invalid)

def current_truth_audit():
    manifest=json.loads((ROOT/'audit/test_manifest.json').read_text(encoding='utf-8'))
    exact=[]
    for m in manifest:
        if m['modality']!='thermal' or m.get('gt_status')!='exact_hash_existing_annotation': continue
        assert sha(m['path'])==m['sha256']
        for cand in m['exact_annotation_candidates']:
            assert sha(cand['source'])==m['sha256']
            xml=ET.parse(cand['xml']).getroot()
            boxes=[]
            for obj in xml.findall('object'):
                assert obj.findtext('name').lower()=='fire'
                b=obj.find('bndbox'); boxes.append([float(b.findtext(k)) for k in ['xmin','ymin','xmax','ymax']])
            assert boxes==m['boxes'], 'Source XML changed: cannot reuse old audit'
        exact.append(dict(image_name=m['basename'],image_sha256=m['sha256'],boxes=m['boxes'],source_xml=m['annotation_provenance']))
    # Current source annotation candidates are NOT silently upgraded to reviewed scene truth.
    raw=[]
    for p in sorted((SOURCE/'annotations').rglob('*')):
        if not p.is_file(): continue
        item=dict(path=str(p),sha256=sha(p),bytes=p.stat().st_size)
        if p.suffix.lower()=='.txt' and p.name!='classes.txt':
            lines=p.read_text(encoding='utf-8-sig').splitlines(); item['nonempty_lines']=len([l for l in lines if l.strip()])
            item['status']='unreviewed_candidate; empty file not assumed negative'
            item['image_exists']=(SOURCE/p.parent.name/(p.stem+'.jpg')).exists()
            valid=True
            for line in lines:
                if not line.strip(): continue
                try:
                    a=list(map(float,line.split())); valid &= len(a)==5 and a[0]==0 and all(0<=v<=1 for v in a[1:]) and a[3]>0 and a[4]>0 and a[1]-a[3]/2>=-1e-6 and a[1]+a[3]/2<=1+1e-6 and a[2]-a[4]/2>=-1e-6 and a[2]+a[4]/2<=1+1e-6
                except ValueError: valid=False
            item['format_valid']=bool(valid)
        else: item['text']=p.read_text(encoding='utf-8-sig')[:200]
        raw.append(item)
    annotations=json.loads(ANNOTATOR.read_text(encoding='utf-8'))
    reviewed=[a for a in annotations.values() if a.get('reviewed')]
    assert all(0<=a['pair_id']<1528 for a in reviewed)
    data=dict(checked_utc=datetime.now(timezone.utc).isoformat(),thermal_exact_verified_images=len(exact),thermal_exact_verified_boxes=sum(len(m['boxes']) for m in exact),rgb_old_audit_gt_covered=0,reviewed_scene_records=reviewed,annotator_path=str(ANNOTATOR),annotator_sha256=sha(ANNOTATOR),annotator_mapping_source='app.py Dataset: zero-based text line index, DEFAULT_SOURCE same paper_sample.txt',annotator_app_sha256=sha(ANNOTATOR.parent.parent/'app.py'),current_source_annotation_files=raw,raw_annotation_candidate_count=sum(p['path'].endswith('.txt') and not p['path'].endswith('classes.txt') for p in raw),raw_nonempty_candidate_count=sum(p.get('nonempty_lines',0)>0 for p in raw),full_independent_scene_truth_available=False,truth_policy='Temperature and UV are inputs only. Raw RGB annotations are candidate object boxes with unverified review/provenance; not full scene truth. Reviewed record(s) are not representative full-set metrics.')
    save('current_truth_audit.json',data); save('thermal_verified_subset.json',exact)
    return data,exact

def thermal_metrics(images,truth):
    byname={x['image_name']:x for x in images}; tp=fp=fn=0; details=[]
    for g in truth:
        pred=byname[g['image_name']]; assert pred['image_sha256']==g['image_sha256']
        detections=sorted((d for d in pred['detections'] if d['score']>=.5),key=lambda d:d['score'],reverse=True)
        matched=set(); itp=0
        for d in detections:
            available=[(c.box_iou(d['bbox_xyxy'],b),j) for j,b in enumerate(g['boxes']) if j not in matched]
            if available:
                iou,j=max(available)
                if iou>=.5: matched.add(j); itp+=1
        ifp=len(detections)-itp; ifn=len(g['boxes'])-itp
        tp+=itp; fp+=ifp; fn+=ifn
        details.append(dict(image_name=g['image_name'],image_sha256=g['image_sha256'],ground_truth_boxes=len(g['boxes']),predicted_boxes_score_05=len(detections),TP=itp,FP=ifp,FN=ifn))
    p=tp/(tp+fp) if tp+fp else None; r=tp/(tp+fn) if tp+fn else None
    f=2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else None
    result=dict(scope='ONLY exact-hash and current source-XML verified thermal subset; target-level, not full fusion or scene-level',images=len(truth),gt_boxes=tp+fn,score_threshold=.5,iou_threshold=.5,matching='descending score, one-to-one greedy best unmatched GT IoU',TP=tp,FP=fp,FN=fn,precision=p,recall=r,F1=f,accuracy=None,accuracy_reason='Object detection has no well-defined TN universe',details=details)
    save('thermal_subset_metrics.json',result); write_csv('thermal_subset_metrics.csv',details)
    return result

def run():
    OUT.mkdir(exist_ok=True); REPORT.parent.mkdir(exist_ok=True)
    original_hashes={str(p):sha(p) for p in list((CPP/'src').glob('*.cpp'))+list((CPP/'main').glob('*.cpp'))+list((CPP/'include').glob('*.h'))}
    sample_path=SOURCE/'paper_sample.txt'; sample_sha=sha(sample_path)
    samples=[c.parse_sample_line(l) for l in sample_path.read_text(encoding='utf-8-sig').splitlines()]
    assert len(samples)==1528
    predictions={}; verification={}
    for mode in ['rgb','thermal']:
        p=OUT/(mode+'_predictions.json'); data=json.loads(p.read_text(encoding='utf-8'))
        images=data['images']; meta=data['metadata']; summary=json.loads((OUT/(mode+'_summary.json')).read_text(encoding='utf-8'))
        assert len(images)==1528 and {x['pair_index'] for x in images}==set(range(1528))
        assert summary['complete'] and summary['processed_images']==1528 and summary['missing_images']==[]
        assert sha(meta['checkpoint'])==meta['checkpoint_sha256'] and sample_sha==meta['paper_sample_sha256']
        byindex={x['pair_index']:x for x in images}
        for i,row in enumerate(samples):
            image=byindex[i]; name=row.rgb if mode=='rgb' else row.thermal
            assert image['image_name']==name and image['source_line']==i+1
            assert sha(image['image_path'])==image['image_sha256']
        predictions[mode]=byindex
        verification[mode]=dict(processed_images=len(images),missing_images=[],unique_pair_indices=len(byindex),checkpoint=meta['checkpoint'],checkpoint_sha256=meta['checkpoint_sha256'],prediction_sha256=sha(p),retained_boxes=sum(len(x['detections']) for x in images),minimum_retained_score=meta['minimum_score'],config_sha256=meta['config_sha256'],current_image_hashes_verified=True)
    h,cal=calibration(); truth,exact=current_truth_audit()
    tm=thermal_metrics(list(predictions['thermal'].values()),exact)
    state=c.SensorState(); prior=None; rows=[]; polygons=[]; inputs=[]; expected=[]; resets=[]; sequence=0
    for i,sample in enumerate(samples):
        ts=datetime.strptime(sample.rgb.removesuffix('_bgr.jpg'),'%Y_%m_%d_%H_%M_%S_%f')
        delta=None if prior is None else (ts-prior).total_seconds()
        reset=prior is None or delta<=0 or delta>120
        if reset:
            state=c.SensorState(); sequence+=1
            resets.append(dict(pair_index=i,source_line=i+1,gap_seconds=delta,reason='initial' if prior is None else 'timestamp_discontinuity_gap_gt_120s_or_nonincreasing'))
        prior=ts
        rgb=predictions['rgb'][i]; thermal=predictions['thermal'][i]
        rc=max((d['score'] for d in rgb['detections']),default=0.)
        tc=max((d['score'] for d in thermal['detections']),default=0.)
        sr=c.sensor_step(state,sample,rc,tc); state=sr.state
        spatial,pair=geometry(rgb,thermal,h)
        fused={'legacy_fixed_01':c.fuse3(sr.rgb,sr.thermal,sr.uv,'spatial',.1),'classic_ds':c.fuse3(sr.rgb,sr.thermal,sr.uv,'ds'),'candidate_geometry_repeated':c.fuse3(sr.rgb,sr.thermal,sr.uv,'spatial',spatial),'recommended_spatial_then_ds':c.fuse_rgb_thermal_uv(sr.rgb,sr.thermal,sr.uv,'spatial',spatial)}
        row=dict(pair_index=i,source_line=i+1,rgb_name=sample.rgb,thermal_name=sample.thermal,Tmax=sample.temperature,UV_input=sample.uv,sequence_id=sequence,state_reset=int(reset),gap_seconds=delta,rgb_confidence=rc,thermal_confidence=tc,rgb_baseline_decision='fire' if rc>=.5 else 'nofire',thermal_baseline_decision='fire' if tc>=.5 else 'nofire',candidate_spatial_iou=spatial,top_score_pair_iou=pair['top_score_iou'],geometry_status=pair['status'])
        for name,m in [('rgb_bpa',sr.rgb),('thermal_bpa',sr.thermal),('uv_bpa',sr.uv)]:
            for k,v in zip(['fire','nofire','uncertain'],m.values()): row[name+'_'+k]=v
        for name,result in fused.items():
            for k,v in zip(['fire','nofire','uncertain'],result.mass.values()): row[name+'_'+k]=v
            row[name+'_decision']=c.decide(result.mass,.95)
            row[name+'_binary05']='fire' if result.mass.fire>=.5 else 'nofire'
            row[name+'_conflict_ab']=result.conflict
            row[name+'_degenerate']=int(result.degenerate)
        rows.append(row); polygons.append(dict(pair_index=i,spatial_consistency=spatial,**pair))
        inputs.append(f'{int(reset)} {rc:.17g} {tc:.17g} {sample.temperature:.17g} {sample.uv}\n')
        expected.append([*sr.rgb.values(),*sr.thermal.values(),*sr.uv.values(),*fused['legacy_fixed_01'].mass.values(),*fused['classic_ds'].mass.values()])
    # Actual native C++ reference, compiled from the user's unchanged five .cpp files.
    proc=subprocess.run([str(OUT/'cpp_reference_harness.exe')],input=''.join(inputs),text=True,capture_output=True,check=True)
    ref=np.array([list(map(float,l.split())) for l in proc.stdout.splitlines()]); exp=np.array(expected)
    assert ref.shape==exp.shape==(1528,15) and np.isfinite(ref).all()
    maxerr=float(np.abs(ref-exp).max()); assert maxerr<1e-10
    (OUT/'cpp_reference_results.txt').write_text(proc.stdout,encoding='utf-8')
    (OUT/'cpp_reference_inputs.txt').write_text(''.join(inputs),encoding='utf-8')
    save('cpp_parity_verification.json',dict(rows=1528,values_per_row=15,compared='RGB/thermal/UV BPA + original fixed0.1 fuse3 + classic DS',maximum_absolute_error=maxerr,tolerance=1e-10,reference_exit_code=proc.returncode,source_cpp_unchanged=all(sha(p)==v for p,v in original_hashes.items())))
    assert all(sha(p)==v for p,v in original_hashes.items())
    write_csv('fusion_results.csv',rows); save('fusion_results.json',dict(primary='legacy_fixed_01',rows=rows))
    save('geometry_pairing.json',polygons)
    label_rows=[dict(pair_index=i,source_line=i+1,rgb_name=s.rgb,thermal_name=s.thermal,scene_truth='',reviewed='',reviewer='',evidence_exposed='',sequence_id='',rgb_boxes_xyxy='',thermal_boxes_xyxy='',notes='') for i,s in enumerate(samples)]
    write_csv('blank_scene_label_template_1528.csv',label_rows)
    counts={}
    for name in METHODS:
        ctr=Counter(r[name+'_decision'] for r in rows)
        counts[name]=dict(fire=ctr['fire'],nofire=ctr['nofire'],uncertain=ctr['uncertain'],decision_coverage=(1528-ctr['uncertain'])/1528,uncertainty_rate=ctr['uncertain']/1528,accuracy=None,precision=None,recall=None,F1=None)
        if name not in ['rgb_baseline','thermal_baseline']:
            counts[name]['binary05_fire_count']=sum(r[name+'_binary05']=='fire' for r in rows)
            counts[name]['degenerate_count']=sum(r[name+'_degenerate'] for r in rows)
    reviewed_results=[]
    for a in truth['reviewed_scene_records']:
        r=rows[a['pair_id']]
        reviewed_results.append(dict(pair_index=a['pair_id'],scene_truth=a['scene_truth'],evidence_exposed=a.get('evidence_exposed'),predictions={name:r[name+'_decision'] for name in METHODS},interpretation='Single/current reviewed records only; not full-set accuracy or representative performance'))
    summary=dict(primary_algorithm='Original C++ fuse3 spatial s=0.1 applied twice; original random inputs replaced real NN maximum scores and text Tmax/UV',total_rows=1528,inference_verification=verification,test_source=str(sample_path),test_source_sha256=sample_sha,algorithm_source_sha256=original_hashes,parameters=dict(window=10,time_lambda=.7,rgb_reliability=.3,T0=50,uv_alpha=.01,uv_beta=.1,baseline_threshold=.5,fusion_fire_nofire_threshold=.95),temporal_policy=dict(order='paper_sample.txt exact order',reset_rule='initial or nonincreasing timestamp or gap >120 seconds',resets=resets,sequence_count=sequence,sequence_boundaries_validated=False,note='Filename timestamps and 120s threshold identify acquisition discontinuities heuristically, not verified scene boundaries. Fraction parsed with %f; physical timestamp scale unverified. State reset is an explicit adaptation from original uninterrupted while loop.'),calibration=cal,geometry_policy='Anchor max-score RGB detection; select largest polygon IoU thermal detection among all retained >=0.01 outputs (tie: higher thermal score). Confidence BPA uses independent frame maxima. Also export max-score/max-score IoU. No one-to-one multiobject assignment. Missing/no overlap or invalid geometry => s=0; absence not proof of nofire. No clipping or offsets.',counts=counts,distributions={key:stats([r[key] for r in rows]) for key in ['rgb_confidence','thermal_confidence','Tmax','candidate_spatial_iou','top_score_pair_iou','uv_bpa_fire','legacy_fixed_01_fire','classic_ds_fire','candidate_geometry_repeated_fire','recommended_spatial_then_ds_fire']},geometry_nonzero_rows=sum(r['candidate_spatial_iou']>0 for r in rows),reviewed_scene_sanity=reviewed_results,thermal_subset_metrics={k:v for k,v in tm.items() if k!='details'},truth_audit=dict(thermal_verified_images=truth['thermal_exact_verified_images'],thermal_verified_boxes=truth['thermal_exact_verified_boxes'],reviewed_scene_count=len(truth['reviewed_scene_records']),raw_rgb_candidate_count=truth['raw_annotation_candidate_count'],raw_nonempty_rgb_candidate_count=truth['raw_nonempty_candidate_count'],full_scene_truth=False),cpp_reference_max_abs_error=maxerr,metrics_note='Full 1528-row Accuracy/Precision/Recall/F1 unavailable. Never use Tmax or UV as labels. Output counts are not accuracy.',created_utc=datetime.now(timezone.utc).isoformat())
    save('fusion_summary.json',summary)
    report(summary,rows,truth,tm)
    with zipfile.ZipFile(REPORT) as z:
        assert z.testzip() is None
        assert all(n in z.namelist() for n in ['[Content_Types].xml','word/document.xml','_rels/.rels'])
    doc=Document(REPORT)
    readback='\n'.join([p.text for p in doc.paragraphs]+[cell.text for t in doc.tables for row in t.rows for cell in row.cells])
    assert '1528' in readback and 'N/A' in readback and '固定0.1' in readback and '未验证' in readback
    (ROOT/'reports/PPYOLOv2_DS融合实际测试报告_readback.txt').write_text(readback,encoding='utf-8')
    csv_rows=list(csv.DictReader((OUT/'fusion_results.csv').open(encoding='utf-8-sig')))
    assert len(csv_rows)==1528 and len({r['pair_index'] for r in csv_rows})==1528
    labels=list(csv.DictReader((OUT/'blank_scene_label_template_1528.csv').open(encoding='utf-8-sig')))
    assert len(labels)==1528 and all(not r['scene_truth'] for r in labels)
    verified=dict(output_csv_rows=len(csv_rows),unique_pair_indices=1528,blank_label_rows=1528,blank_labels_verified=True,docx_crc_ok=True,docx_readback_ok=True,docx_paragraphs=len(doc.paragraphs),docx_tables=len(doc.tables),docx_sha256=sha(REPORT),fusion_csv_sha256=sha(OUT/'fusion_results.csv'),original_cpp_unchanged=True,cpp_reference_rows=1528,cpp_reference_max_abs_error=maxerr,report_path=str(REPORT))
    save('fusion_deliverable_verification.json',verified)
    print(json.dumps(dict(verification=verified,counts=counts,thermal_metrics=summary['thermal_subset_metrics'],calibration=dict(fit_rmse=cal['fit_rmse_px'],loo_rmse=cal['leave_one_out_rmse_px']),truth_audit=summary['truth_audit'],temporal_resets=resets),ensure_ascii=False,indent=2))

def report(s,rows,truth,tm):
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7)
    st=d.styles['Normal']; st.font.name='Microsoft YaHei'; st.font.size=Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    d.add_heading('PPYOLOv2 多模态 D-S 融合实际测试报告',0)
    d.add_paragraph('主要交付：忠实复现用户 C:/lyk/paper/bpa_test 的原始 fuse3 固定0.1两次空间规则；候选几何仅列为探索扩展。')
    d.add_paragraph('真实执行1528对RGB/热图及文本温度、UV输入。完整融合性能指标不可计算；本报告不声称融合优于单模态。生成时间（UTC）：'+s['created_utc'])
    def heading(t): d.add_heading(t,1)
    def p(t): d.add_paragraph(t)
    def table(headers,data):
        t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
        for cell,text in zip(t.rows[0].cells,headers): cell.text=str(text)
        for row in data:
            for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
    heading('1. 实际输入、模型及完整性')
    p('测试文件：'+s['test_source']+'。每行依次为RGB文件名、thermal文件名、最高温度、UV电平。温度单位与测量校准未独立验证。UV与温度仅为算法输入，绝不是真值。')
    p('paper_sample.txt SHA256：'+s['test_source_sha256'])
    table(['模态','实际图像数','缺失数','保留框数（≥0.01）'],[[k,v['processed_images'],len(v['missing_images']),v['retained_boxes']] for k,v in s['inference_verification'].items()])
    for mode,v in s['inference_verification'].items():
        p(mode+'模型：'+v['checkpoint']); p('模型SHA256：'+v['checkpoint_sha256']); p('预测JSON SHA256：'+v['prediction_sha256'])
    p('两模态各1528个唯一pair_index，逐行文件名/原始行号一致，当前原图SHA256全部复核；模型文件重新计算SHA256与推理元数据一致。PPYOLOv2为ResNet50-vd/DCNv2、PPYOLOPAN、IoU-aware head、MatrixNMS；配置architecture键YOLOv3是PaddleDetection实现方式，不是换用普通YOLOv3。UV未进入神经网络。')
    heading('2. 原始C++算法及端到端一致性')
    p('算法权威来源：src/vision_bpa.cpp、temp_bpa.cpp、uv_bpa.cpp、ds_fusion.cpp和main/main.cpp；源代码完全只读。main.cpp中随机置信度与随机温度被真实NN最大框置信度和paper_sample文本输入替代；若无框，置信度取0。窗口10、λ=0.7、RGB可靠性0.3、T0=50、UV α=0.01、β=0.1均保持原main设置。')
    p('稳定性 S=exp(-std(history_confidence))；RGB当前BPA=(c·S·0.3, (1-c)·S, 1−前两项)。热PC=clamp(1−exp(−10c),0,1)，PT=clamp(1−exp(−0.1(Tmax−50)),0,1)，BPA=(PC·PT·S, (1−c)(1−PT)S, 1−前两项)。thermal历史记录的是置信度而不是温度，忠实保留。所有BPA经λ·previous+(1−λ)·current后归一化。')
    p('UV初始q=0.001，UV=1时q←q+0.01q(1−q)，UV=0时q←q−0.1q(1−q)，再clamp。当前BPA=(q,0.8(1−q),0.2(1−q))；原始previous=(0.1,0.7,0.05)总和0.85，不擅自修正初始化，第一次平滑后才归一化。该更新可能令UV火焰支持长期较低，这不构成预测正确性的证据。')
    p('经典DS：K=aF·bN+aN·bF；F=aF·bF+aF·bU+aU·bF，N=aN·bN+aN·bU+aU·bN，U=aU·bU；除以1−K再normalize。原空间规则：Ksp=(1−s)K，Fsp=s·F/(1−Ksp)，Nsp=N/(1−Ksp)，Usp=U/(1−Ksp)，最后normalize。归一化将共同分母取消，因此调整K并未独立改变最终归一化比例；实际为(sF,N,U)/(sF+N+U)。固定s=0.1只是源码常数，不是实测空间一致性。')
    p('原fuse3先融合RGB与thermal，再把结果与UV融合，两次均用s=0.1，返回abc_sp；尽管内部计算PCR和DS，最终返回的是空间分支。本报告主要结果保留该行为。候选几何重复空间分支也仍把同一s施加UV，是为保持原分支结构的探索替换；UV没有位置，不能据此声称UV具有几何一致性。可选建议变体仅RGB/thermal用空间，随后UV用经典DS，明确不同于用户原fuse3。未启用源码中未调用的UV gate。')
    p('实际编译链接用户原始5个BPA/DS .cpp文件（未改），同1528行真实输入运行C++参考程序；比较每行RGB/thermal/UV三项BPA、原fuse3和经典DS共15个数值，最大绝对误差='+format(s['cpp_reference_max_abs_error'],'.3e')+'，阈值1e−10，参考程序exit0。总/近总冲突或零支持时Python显式输出(0,0,1)并标记；这是避免原源码normalize零除的安全差异，逐方法实际发生数见结果表。')
    heading('3. 时序处理边界及限制')
    p('严格按paper_sample原文本顺序处理；文件名时间戳非递增或相邻间隔>120秒时重置全部BPA历史。此阈值是显式采集间断启发式，不是已确认场景分割。小于阈值也可能跨场景；每行间隔变化，原λ是按样本而非物理时间。原main持续循环无重置，本次重置是明确的真实采集适配。文件名小数部分按%f解析，物理时间尺度仍未验证。')
    table(['重置pair_index','文本行','相邻间隔秒','原因'],[[r['pair_index'],r['source_line'],r['gap_seconds'],r['reason']] for r in s['temporal_policy']['resets']])
    heading('4. 候选配准与真实polygon IoU（未验证）')
    cal=s['calibration']; p('来源：main/homo_main.cpp当前未注释的七组对应点，方向thermal pts2→RGB pts1。cv2.findHomography(RANSAC,3像素)估计；没有使用compute_homography_rectangle中−5、+35、+15等人为偏移。状态：candidate/exploratory, unverified across test scenes（候选/探索，跨测试场景未验证）。')
    table(['thermal pts2','RGB pts1','拟合误差(px)','留一误差(px)'],[[a,b,format(e,'.4f'),format(l,'.4f')] for a,b,e,l in zip(cal['thermal_points'],cal['rgb_points'],cal['fit_errors_px'],cal['leave_one_out_errors_px'])])
    p('H（原图像素坐标）：\n'+np.array2string(np.array(cal['H']),precision=9)); p('拟合RMSE='+format(cal['fit_rmse_px'],'.4f')+'px，最大='+format(cal['fit_max_px'],'.4f')+'px；留一RMSE='+format(cal['leave_one_out_rmse_px'],'.4f')+'px，最大='+format(cal['leave_one_out_max_px'],'.4f')+'px。留一仍来自同一组局部对应点，绝不等价于独立场景验证。')
    p('七点集中在很窄的y带，向其他区域外推存在风险；不同场景、平面深度、镜头位姿和视差均未验证。不能宣称该H已验证或配准完成。需要覆盖全视场、多独立场景的对应点和外部验证残差。')
    p('配对：以最大置信度RGB框为锚，将各保留thermal框四角用H投影成四边形；cv2.intersectConvexConvex求多边形交并比，选IoU最大thermal框（相同时置信度高者）。BPA置信度仍使用各模态独立最大分数，所以几何最佳thermal框可能不是thermal最高分框；导出其rank和分数，同时导出最高分对IoU。此为帧级单锚最佳IoU，不是多目标一对一匹配；低分框最佳匹配可能偏乐观。输出框原像素坐标，不以轴对齐外接框代替投影多边形，不进行额外裁剪。缺框/无重叠/投影非法时s=0并记录原因；s=0并不等于无火。')
    p('实际非零几何一致性行数：'+str(s['geometry_nonzero_rows'])+' / 1528。')
    heading('5. 实际输出分布与方法对比')
    p('单模态基线：最大检测分数≥0.5判fire，否则nofire。融合主决策：mF≥0.95判fire，mN≥0.95判nofire，否则uncertain；该0.95沿用原main对fire的阈值，nofire对称与uncertain输出是显式报告策略。CSV另给mF≥0.5二分类结果便于未来统一阈值评估。不同阈值统计不是性能优劣比较。')
    table(['方法','Fire','NoFire','Uncertain','覆盖率','不确定率','退化数'],[[NAMES[name],v['fire'],v['nofire'],v['uncertain'],format(v['decision_coverage'],'.4%'),format(v['uncertainty_rate'],'.4%'),v.get('degenerate_count','—')] for name,v in s['counts'].items()])
    table(['变量','最小','中位','均值','最大'],[[k,*[format(v[a],'.6g') for a in ['min','median','mean','max']]] for k,v in s['distributions'].items()])
    heading('6. 真值现状与准确率/查准率/召回率/F1')
    p('当前重新检查了测试annotations目录、人工复核数据库及旧源XML/原图。旧audit的RGB真值覆盖0不是当前全部文件现状：现发现RGB YOLO候选标签'+str(truth['raw_annotation_candidate_count'])+'份，其中非空'+str(truth['raw_nonempty_candidate_count'])+'份；标签内容和格式已记录，人工复核/来源独立性未确认，空文件不自动认定负样本。人工数据库有'+str(len(truth['reviewed_scene_records']))+'条reviewed场景记录，其索引映射由app.py和同一paper_sample确认。这些零散记录不能补齐1528对独立场景真值。')
    table(['方法/范围','Accuracy','Precision','Recall','F1'],[[NAMES[n]+'（全1528对）','N/A','N/A','N/A','N/A'] for n in METHODS]+[['thermal精确复核89图目标级子集','N/A',format(tm['precision'],'.6f') if tm['precision'] is not None else 'N/A',format(tm['recall'],'.6f') if tm['recall'] is not None else 'N/A',format(tm['F1'],'.6f') if tm['F1'] is not None else 'N/A']])
    p('N/A原因：缺少覆盖全测试集、与输入证据独立的fire/nofire场景真值，未标注不能当无火；温度高低、UV值、网络输出均不能当真值。Precision=TP/(TP+FP)，Recall=TP/(TP+FN)，F1=2TP/(2TP+FP+FN)，Accuracy=(TP+TN)/(TP+FP+FN+TN)。应明确uncertain是否弃权及覆盖率，不把弃权算准确。')
    p('热成像严格子集：当前逐图重新计算SHA256，与源训练数据同图哈希匹配；重新解析源XML并核对原审计框，'+str(tm['images'])+'张、'+str(tm['gt_boxes'])+'框。仅以score≥0.5、IoU≥0.5、按分数贪心一对一匹配计算目标级：TP='+str(tm['TP'])+'，FP='+str(tm['FP'])+'，FN='+str(tm['FN'])+'。没有负样本且目标级TN无定义，Accuracy不可算。该偏选子集不代表全部1528图，更不能代表融合场景级性能；这些测试同哈希图在旧训练数据审计中被排除。')
    for r in s['reviewed_scene_sanity']:
        p('人工已复核记录核查（仅个案，不汇总为性能）：pair_index='+str(r['pair_index'])+'，scene_truth='+r['scene_truth']+'，evidence_exposed='+str(r['evidence_exposed'])+'，预测='+json.dumps(r['predictions'],ensure_ascii=False))
    heading('7. 可追溯样例（仅输出，不伪造真值）')
    indices=sorted(set([0,1,60,127,130,911,1348,1527,int(np.argmax([r['candidate_spatial_iou'] for r in rows]))]))
    table(['pair/行','RGB分数','热分数','s','原固定0.1 mF','经典DS mF','几何重复 mF'],[[str(i)+'/'+str(i+1),*[format(rows[i][k],'.6f') for k in ['rgb_confidence','thermal_confidence','candidate_spatial_iou','legacy_fixed_01_fire','classic_ds_fire','candidate_geometry_repeated_fire']]] for i in indices])
    heading('8. 产物、验证与后续必要条件')
    p('fusion/fusion_results.csv和fusion_results.json：1528行完整真实融合输出；fusion_summary.json：计数、分布、哈希、边界；candidate_calibration.json：H和拟合/留一残差；geometry_pairing.json：实际配对框；cpp_reference_results.txt、cpp_parity_verification.json：原C++逐行参考；current_truth_audit.json：当前真值检查；thermal_subset_metrics.json/csv：严格热目标级子集；blank_scene_label_template_1528.csv：1528行空白标签模板，所有scene_truth保留空值。')
    p('必须补齐：独立人工场景真值（包含无火场景及复核记录），完整且明确来源的模态框标签，已确认的采集序列边界和时间单位，多场景配准验证，预先定义且同尺度的决策阈值与弃权处理。之后才可计算完整Accuracy/P/R/F1并比较融合收益。当前交付是实际算法执行报告，而不是已证实性能提升的论文结论。')
    heading('附录：原算法源文件SHA256')
    for path,v in s['algorithm_source_sha256'].items(): p(path+'\n'+v)
    footer=sec.footer.paragraphs[0]; footer.text='真实执行报告 · 候选配准未验证 · 完整性能指标N/A'
    d.save(REPORT)

if __name__=='__main__': run()
