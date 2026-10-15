from pathlib import Path
import json,csv,hashlib,zipfile,collections
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path('C:/lyk/fire/qualityDetection');OUT=ROOT/'fusion/calibrated_spatial_ds';REPORT=ROOT/'reports/校准空间DS_优化验证报告.docx'
def load(name):return json.loads((OUT/name).read_text(encoding='utf-8'))
def build():
 assert not REPORT.exists(),'Do not overwrite any previous report.'
 m=load('metrics.json');split=load('split_metadata.json');receipt=load('execution.json');frozen=load('frozen_selected.json');audit=load('source_audit.json')
 old=json.loads((ROOT/'fusion/crossmodal_roi_ds/full_results.json').read_text());heldgroups=[g for g in split['groups'] if g['split']=='heldout' and g['label']==1];ix=[i for g in heldgroups for i in range(g['start'],g['stop_exclusive'])]
 diag={'heldout_positive':len(ix),'rgb_candidates_present':sum(bool(old[i]['rgb_detections']) for i in ix),'thermal_candidates_present':sum(bool(old[i]['thermal_detections']) for i in ix),'dual_present':sum(bool(old[i]['rgb_detections']) and bool(old[i]['thermal_detections']) for i in ix),'roi_attempts':sum(len(old[i]['roi_attempts']) for i in ix),'rgb_roi_real_detections':sum(len(a['detections']) for i in ix for a in old[i]['roi_attempts'] if a['review_modality']=='rgb'),'diagnostic_only':'No heldout tuning; derives failure cause from frozen .1 candidate pool. Cannot claim bound for fresh .05 forward.'}
 (OUT/'heldout_failure_diagnostic.json').write_text(json.dumps(diag,indent=2),encoding='utf-8')
 d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.7);sec.left_margin=sec.right_margin=Cm(1.8)
 normal=d.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10);normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 d.add_heading('校准空间DS：优化与内部保留验证报告',0)
 d.add_paragraph('结论：全量UV参考一致性达到目标；内部保留子集未满足目标。因此“全量与保留均通过”的验收结果为失败，不能声称独立泛化通过。')
 d.add_paragraph('全部1528帧此前均已被查看。本次只能称为“内部保留评估、探索”，并非独立验证。标签来自paper_sample.txt第四列UV参考，而非独立人工场景真值；空间仿射配准仍为未独立验证的候选几何。')
 d.add_heading('1. 真数据、模型与复用边界',1)
 d.add_paragraph('真实配对1528组：UV阳性1398、阴性130。逐文件重新计算3056张源图SHA256，与缓存记录逐项一致，3056个哈希唯一。源文件、原始权重、历史预测、旧实验输出及既有DOCX均保持不变。没有新增硬件、没有重训检测器。')
 d.add_paragraph('模型均为用户提供YOLO26n，单fire类，参数2504190；不得称作PPYOLOv2。RGB best.pt与热成像best-2.pt路径及完整SHA见source_audit.json。缓存的真实前向输入640，full-frame conf=.1，ROI conf=.1、触发conf=.25、每模态至多3次，映射回原图的真实ROI检测才可作为证据，投影框仅用于裁剪提议。')
 d.add_paragraph('本次检测候选阈值在已缓存的真实.1输出上重新过滤，所有搜索阈值均不低于.1；未假装拥有更低阈值预测，未运行新前向。上次GPU前向及完整ROI provenance见源目录full_execution.json/full_results.json，本次numerical search使用该真实缓存。')
 d.add_heading('2. 搜索之前固定的时间分组',1)
 d.add_paragraph('先检查文件名时间与标签分布：阴性集中在2026年5月，阳性集中在2026年8月；全局严格按时间60/20/20会使部分集合缺类。因此使用类别分层、类内按时间顺序的整组分配。大于10秒的时间间隙/时间回退/UV标签变化产生15个时间块代理；不是已验证场景，不能虚构scene ID。四位时间尾数按小数秒解释，来源单位尚未独立验证。')
 d.add_paragraph('目标为每类累计帧数接近60%与80%的整组切分。组大小不均，无法严格60/20/20；实际920/406/202。未随机打散相邻帧，同组不跨集合；未知同视频或近重复泄漏仍可能存在。训练部分未拟合后验校准器，仅作描述性评估。')
 t=d.add_table(rows=1,cols=5);t.style='Table Grid'
 for cell,s in zip(t.rows[0].cells,['集合','帧数','UV阳性','UV阴性','组数']):cell.text=s
 for key,label in [('train','训练/描述'),('calibration','校准/选参'),('heldout','内部保留')]:
  z=m['results'][key];cells=t.add_row().cells
  for cell,v in zip(cells,[label,z['N'],z['positive'],z['negative'],sum(g['split']==key for g in split['groups'])]):cell.text=str(v)
 d.add_paragraph('split_allocation.csv与split_metadata.json均在新搜索开始前写出。分配SHA256：'+split['allocation_sha256'])
 d.add_heading('3. BPA、可靠性、空间匹配与决策',1)
 d.add_paragraph('对真实火焰候选，使用fire-only support：m=(r*c, 0, 1-r*c)，而不是把1-c误当非火证据。r=q_A*[rmin+(1-rmin)*s]，s为仿射四角多边形IoU，q_A=1为名义值，未通过独立数据校准。无检测时为未知；最终必须同时存在两模态真实匹配检测，否则判定score=0。')
 d.add_paragraph('先按tR/tT过滤真实检测，再重做带dummy的Hungarian最大IoU一对一匹配；不沿用过滤前索引。仅s≥gate的匹配可参与经典DS。此BPA非火质量为0，K=0；融合等价mF=1-(1-r*cR)(1-r*cT)，是OR式正证据增强，而非冲突感知负证据融合。取单个支持匹配的最大mF，不累计融合相关重复框。')
 d.add_paragraph('主分支完全不使用UV；决策为mF>threshold。最大温度不参与决策，单位/物理标定未知。UV可选次分支保持mF′=mF*[λ+(1−λ)*gUV]，λ=.5，gUV∈{0,1}。UV=1只保留视觉分数，不能强行触发火警；UV=0只乘.5。其UV输入/UV参考评价是循环一致性，不能替代主分支视觉收益证明。')
 d.add_heading('4. 有限网格、只在校准集选参、冻结',1)
 d.add_paragraph('共480组：12组(tR,tT)覆盖.1,.15,.2,.25,.3,.4,.5；IoU gate∈{.01,.05,.1,.2,.3}；rmin∈{.5,1}；decision_threshold∈{.3,.4,.5,.6}。具体完整网格在search_plan.json。目标按校准集FP最小、FN次小、固定trial_id再次排序。未用全量/保留结果排序，未使用UV次分支选参。')
 d.add_paragraph('冻结参数：'+json.dumps(m['selected_parameters'],ensure_ascii=False)+'。选择trial_id='+str(frozen['selected_trial']['trial_id'])+'。冻结时间='+frozen['created_utc']+'；冻结前保留评价次数0，冻结后正式评价1次，未因保留失败再调参。')
 d.add_paragraph('注意：获选rmin=1使可靠性r=1，不再随IoU动态折扣；空间约束仅体现在真实双模态匹配和IoU门槛。不得声称该获选配置证明了动态可靠性机制优于固定可靠性。置信度与质量仍为名义支持值，报告标题的“校准”指决策阈值校准，不是概率校准。')
 d.add_paragraph('冻结配置SHA256：'+m['freeze_sha256'])
 d.add_heading('5. 同一集合基线与实测指标',1)
 names={'primary_vision':'主：视觉空间DS','secondary_uv':'次：视觉DS+UV','rgb_baseline':'原RGB基线','thermal_baseline':'原热成像基线'}
 for key,label in [('full','全量1528'),('train','训练/描述920'),('calibration','校准406（选参集）'),('heldout','内部保留202')]:
  d.add_heading(label,2);z=m['results'][key];table=d.add_table(rows=1,cols=9);table.style='Table Grid'
  for cell,s in zip(table.rows[0].cells,['方法','TP','FP','FN','TN','FPR%','FNR%','P/R%','F1%']):cell.text=s
  for name,vals in z['methods'].items():
   cells=table.add_row().cells;values=[names[name],vals['TP'],vals['FP'],vals['FN'],vals['TN'],f"{100*vals['FPR']:.4f}",f"{100*vals['FNR']:.4f}",f"{100*vals['precision']:.2f}/{100*vals['recall']:.2f}",f"{100*vals['F1']:.2f}"]
   for cell,v in zip(cells,values):cell.text=str(v)
  for row in table.rows:
   for cell in row.cells:
    for p in cell.paragraphs:
     for run in p.runs:run.font.size=Pt(8)
  d.add_paragraph('主分支目标状态：'+('达到' if z['methods']['primary_vision']['meets_target_same_subset'] else '未满足')+'。阈值比较采用本集合原热成像FNR与RGB FPR，不用全局基线替代子集基线。')
 d.add_paragraph('全量主分支：TP1032/FP0/FN366/TN130，FPR=0%，FNR=26.1802575%，低于原热成像47.63948498%；原RGB FPR=96.92307692%。但内部保留：TP0/FP0/FN180/TN22，FNR=100%，明显劣于该子集原热成像TP176/FN4、FNR=2.22222222%。双重验收未满足。')
 d.add_paragraph('次UV分支在四个集合与主分支指标完全相同，本次没有UV新增收益；不能用UV参考循环评价声称真实泛化改进。')
 d.add_heading('6. 保留失败与候选上界诊断',1)
 d.add_paragraph('保留阳性来自最后时间块（原行1348–1527），180帧热成像均有候选，但RGB真实候选为0；已有360次ROI尝试也未产生RGB真实检测。因此当前conf≥.1缓存下，任何坚持双真实模态候选的BPA/阈值组合都不可能召回这180帧，乐观候选召回上界也是0。此结论只限现有缓存，不代表新.05前向也没有候选。不得用仿射投影框冒充检测，或删除难例/强行UV触发来修饰结果。')
 d.add_paragraph('全量纯空间候选覆盖诊断：IoU≥.01或.05时，1398阳性中1184存在真实匹配，130阴性中0存在真实匹配；这只是标签事后诊断的候选上界，不是分类器、不可用于选参或宣称可达准确率。已保存candidate_coverage_diagnostic.json与heldout_failure_diagnostic.json。')
 d.add_heading('7. 文件、验证与局限',1)
 d.add_paragraph('新目录：'+str(OUT)+'。关键交付：calibrated_core.py、run_calibration.py、split_allocation.csv、search_plan.json、calibration_trials.csv、frozen_selected.json、frame_predictions.csv、metrics.json、source_audit.json、execution.json、单元测试与RED/GREEN日志。运行入口遇到冻结配置已存在会拒绝再次搜索，防止默默复用保留集调参。')
 d.add_paragraph('数值指标从逐帧CSV重新独立计数核对，冻结SHA与选择trial核对；DOCX生成后ZIP、XML及内容读回验证。所有旧输出和权重SHA再次核对。主要未解决问题：现有RGB模型/ROI在最后正样本时间块没有真实候选；时间块不等于场景且近重复未验证；配准未独立验证；没有独立人工真值或全新场景。需要新、未使用的成对数据及有效RGB检测覆盖，才能检验部署泛化。不能把本次全量成功包装为新的独立验证通过。')
 d.add_paragraph('来源冻结审计与执行结果：'+json.dumps(receipt,ensure_ascii=False))
 footer=sec.footer.paragraphs[0];footer.text='内部保留探索 · 非独立验证 · UV参考一致性'
 REPORT.parent.mkdir(parents=True,exist_ok=True);d.save(REPORT)
 with zipfile.ZipFile(REPORT) as z:assert z.testzip() is None
 read=Document(REPORT);text='\n'.join(p.text for p in read.paragraphs)+'\n'+'\n'.join(c.text for t in read.tables for r in t.rows for c in r.cells)
 (OUT/'report_readback.txt').write_text(text,encoding='utf-8');print(str(REPORT))
if __name__=='__main__':build()
