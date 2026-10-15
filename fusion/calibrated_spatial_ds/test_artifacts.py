import unittest,json,csv,hashlib,io
from pathlib import Path
from docx import Document
ROOT=Path('C:/lyk/fire/qualityDetection');OUT=ROOT/'fusion/calibrated_spatial_ds'
class ArtifactTest(unittest.TestCase):
 def test_csv_metrics_and_frozen_receipt(self):
  m=json.loads((OUT/'metrics.json').read_text()); f=json.loads((OUT/'frozen_selected.json').read_text());rows=list(csv.DictReader(io.StringIO((OUT/'frame_predictions.csv').read_text(encoding='utf-8-sig'))))
  self.assertEqual(len(rows),1528);self.assertEqual(len({r['pair_index'] for r in rows}),1528)
  self.assertEqual(hashlib.sha256((OUT/'frozen_selected.json').read_bytes()).hexdigest(),m['freeze_sha256'])
  trials=list(csv.DictReader(io.StringIO((OUT/'calibration_trials.csv').read_text(encoding='utf-8-sig'))));self.assertEqual(len(trials),480)
  win=min(trials,key=lambda r:(int(r['FP']),int(r['FN']),int(r['trial_id'])));self.assertEqual(int(win['trial_id']),f['selected_trial']['trial_id'])
  for split,result in m['results'].items():
   rr=rows if split=='full' else [r for r in rows if r['split']==split];self.assertEqual(len(rr),result['N'])
   for name,key in [('primary_vision','primary_vision_fire'),('secondary_uv','secondary_uv_fire'),('rgb_baseline','rgb_baseline_fire'),('thermal_baseline','thermal_baseline_fire')]:
    tp=sum(int(r['uv_reference']) and int(r[key]) for r in rr);fp=sum(not int(r['uv_reference']) and int(r[key]) for r in rr);fn=sum(int(r['uv_reference']) and not int(r[key]) for r in rr);tn=sum(not int(r['uv_reference']) and not int(r[key]) for r in rr)
    expected=result['methods'][name];self.assertEqual((tp,fp,fn,tn),tuple(expected[k] for k in ['TP','FP','FN','TN']));self.assertAlmostEqual(expected['FPR'],fp/(fp+tn));self.assertAlmostEqual(expected['FNR'],fn/(fn+tp))
  allocation=list(csv.DictReader(io.StringIO((OUT/'split_allocation.csv').read_text(encoding='utf-8-sig'))));groups={}
  for r in allocation:groups.setdefault(r['group_id'],set()).add(r['split'])
  self.assertTrue(all(len(s)==1 for s in groups.values()))
  self.assertEqual(f['heldout_evaluations_before_freeze'],0);self.assertEqual(m['heldout_evaluations'],1)
 def test_docx_readback(self):
  p=ROOT/'reports/校准空间DS_优化验证报告.docx';self.assertTrue(p.exists(),'new report not yet produced')
  d=Document(p);text='\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
  for x in ['全部1528帧此前均已被查看','内部保留评估、探索','并非独立验证','1032','366','180','480','未满足','K=0','UV']:self.assertIn(x,text)
if __name__=='__main__':unittest.main()
