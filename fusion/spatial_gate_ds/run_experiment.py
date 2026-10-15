"""Run frozen user-defined strict consensus policy on all genuine cached pairs."""
from pathlib import Path
import sys,json,csv,hashlib,zipfile
from datetime import datetime,timezone
from collections import Counter
import numpy as np
import cv2
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
import gate_core as gate
s=gate.s
OUT=Path(__file__).resolve().parent
ROOT=OUT.parent.parent
OLD=OUT.parent/'supplied_models_uv_reference'
AFF=OUT.parent/'supplied_models_affine_iou'
PRIOR=OUT.parent/'spatial_reliability_ds'
SOURCE=Path('C:/lyk/paper_data/paper_data_/paper_data')
REPORT=ROOT/'reports/spatial_gate_DS_1528_最终报告.docx'
sys.path.insert(0,str(OUT.parent))
import core as legacy
NAMES={'rgb':'RGB检测','thermal':'热成像检测','prior_ordinary_ds':'此前普通DS（历史候选）','ordinary_ds':'同候选普通DS（无空间、无UV）','ordinary_uv':'同候选普通DS＋瞬时UV','raw_spatial_ds':'IoU折扣DS（原始判决、未硬门控）','raw_spatial_uv':'IoU折扣DS＋UV（原始判决、未硬门控）','gate_visual':'严格空间门控DS（无UV）','gate_visual_uv':'严格空间门控DS＋瞬时UV','uv_alone':'UV单独（参考恒等）'}

def sha(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()
def load(p): return json.loads(Path(p).read_text(encoding='utf-8'))
def save(name,value): (OUT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def metric(y,p):
    y=np.asarray(y,int); p=np.asarray(p,int)
    assert len(y)==len(p)>0 and set(y)<= {0,1} and set(p)<= {0,1}
    tp=int(np.sum((y==1)&(p==1))); fp=int(np.sum((y==0)&(p==1)))
    fn=int(np.sum((y==1)&(p==0))); tn=int(np.sum((y==0)&(p==0)))
    div=lambda a,b:a/b if b else None
    recall=div(tp,tp+fn); specificity=div(tn,tn+fp)
    return dict(N=len(y),TP=tp,FP=fp,FN=fn,TN=tn,accuracy=(tp+tn)/len(y),precision=div(tp,tp+fp),recall=recall,F1=div(2*tp,2*tp+fp+fn),specificity=specificity,balanced_accuracy=(recall+specificity)/2 if recall is not None and specificity is not None else None,false_positive_rate=div(fp,fp+tn),false_negative_rate=div(fn,fn+tp),predicted_fire_rate=float(np.mean(p)))

def overlays(rows,details,pred,a):
    folder=OUT/'examples'; folder.mkdir(exist_ok=True)
    choices=[]
    conditions=[lambda r:r['semantic_state']=='NoFire' and (r['rgb_count']==0 or r['thermal_count']==0),lambda r:r['semantic_state']=='NoFire' and r['rgb_count']>0 and r['thermal_count']>0,lambda r:r['semantic_state']=='Transition',lambda r:r['semantic_state']=='FireCandidate' and r['UV_reference_label']==1,lambda r:r['semantic_state']=='FireCandidate' and r['UV_reference_label']==0]
    for cond in conditions:
        available=[r for r in rows if cond(r) and r['pair_index'] not in choices]
        if available: choices.append(max(available,key=lambda r:r['polygon_iou'])['pair_index'])
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',17); examples=[]
    for i in choices:
        im=Image.open(pred['rgb'][i]['image_path']).convert('RGB'); draw=ImageDraw.Draw(im)
        for j,b in enumerate(pred['rgb'][i]['detections']):
            draw.rectangle(b['bbox_xyxy'],outline='lime',width=2); draw.text(tuple(b['bbox_xyxy'][:2]),f'R{j}:{b["score"]:.3f}',font=font,fill='lime')
        for j,b in enumerate(pred['thermal'][i]['detections']):
            poly=[tuple(x) for x in s.affine(b['bbox_xyxy'],a)]
            draw.line(poly+[poly[0]],fill='magenta',width=2); draw.text(poly[0],f'T{j}:{b["score"]:.3f}',font=font,fill='magenta')
        im=im.resize((960,round(im.height*960/im.width))); canvas=Image.new('RGB',(960,im.height+130),(20,20,20)); canvas.paste(im,(0,130)); draw=ImageDraw.Draw(canvas); r=rows[i]
        texts=[f'pair {i} 状态={r["semantic_state"]} IoU={r["polygon_iou"]:.4f} g={r["g"]:.4f} UV={r["UV_input"]}',f'硬标签 无UV={r["gate_visual_binary"]} 含UV={r["gate_visual_uv_binary"]}；原始BetP 无UV={r["raw_spatial_ds_betp"]:.4f} 含UV={r["raw_spatial_uv_betp"]:.4f}',f'选中 R{r["rgb_index"]}/T{r["thermal_index"]}；绿色RGB、紫色热框四角仿射多边形',f'τL=0.1、τH=0.3；过渡态二分类无火；UV不得越过门控；非独立真值样例']
        for j,t in enumerate(texts): draw.text((8,4+30*j),t,font=font,fill='white')
        p=folder/f'pair_{i}.png'; canvas.save(p); examples.append(dict(pair_index=i,path=str(p),state=r['semantic_state'],sha256=sha(p)))
    return examples

def report(summary):
    d=Document(); st=d.styles['Normal']; st.font.name='Microsoft YaHei'; st.font.size=Pt(9); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    d.sections[0].left_margin=d.sections[0].right_margin=Inches(.65)
    def p(t): d.add_paragraph(str(t))
    def h(t): d.add_heading(t,1)
    def percent(v): return '未定义' if v is None else f'{100*v:.4f}%'
    def table(headers,data):
        t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
        for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
        for row in data:
            for c,v in zip(t.add_row().cells,row): c.text=str(v)
    met=summary['metrics']['methods']; cfg=summary['config']; cal=summary['calibration']
    d.add_heading('双模态严格空间一致性门控D-S：1528对最终报告',0)
    p('最新用户规则：空间一致性需要RGB与热成像两种模态一致；双模态不一致＝最终二分类无火；IoU负责平滑过渡及有火资格。本次已按此工程决策策略修改并实际运行全部1528对缓存真实预测。仅新增策略与报告，未重新推理或训练，原C++与历史文件不改。UTC：'+summary['created_utc'])
    p('默认数值τL=0.1、τH=0.3可配置。用户没有给出数值，本次在计算指标前固定工程默认，不按UV测试结果选择最优阈值。只有匹配IoU≥0.3且融合BetP(F)>0.5才能最终有火。')
    h('1. 实测指标：不掩盖召回损失')
    table(['方法','TP','FP','FN','TN','Acc','P','R','F1'],[[NAMES[k],v['TP'],v['FP'],v['FN'],v['TN'],percent(v['accuracy']),percent(v['precision']),percent(v['recall']),percent(v['F1'])] for k,v in met.items()])
    table(['方法','平衡准确率','特异度','假阳性率','漏检率','判火比例'],[[NAMES[k],percent(v['balanced_accuracy']),percent(v['specificity']),percent(v['false_positive_rate']),percent(v['false_negative_rate']),percent(v['predicted_fire_rate'])] for k,v in met.items()])
    for name in ['gate_visual','gate_visual_uv']:
        v=met[name]; ref=met['prior_ordinary_ds']; p(f'{NAMES[name]}：Accuracy={percent(v["accuracy"])}，F1={percent(v["F1"])}，TP/FP/FN/TN={v["TP"]}/{v["FP"]}/{v["FN"]}/{v["TN"]}；相对此前普通DS准确率差={100*(v["accuracy"]-ref["accuracy"]):+.4f}个百分点，F1差={100*(v["F1"]-ref["F1"]):+.4f}个百分点。')
    p('结论由上表真实计算给出，不将硬门控宣传成必然准确率改善。严格一致性会排除单模态、错位及低IoU检测，满足“不一致＝无火”的工程策略，同时增加漏检风险；无火输出是用户授权的决策政策，不是物理无火证明。高正例比例下必须同时看Recall、Specificity及Balanced Accuracy。')
    p('UV参考标签：1398有火／130无火，全部1528帧纳入。第四列同时作为含UV分支输入与评价标签，故含UV结果是循环评价、UV符合度，不是独立场景真值验证、不是目标级AP；无UV分支不直接输入UV。UV单独100%为参考恒等，不能据此证明融合泛化。')
    h('2. 冻结的状态与决策边界')
    table(['条件','空间语义状态','g/r','最终二分类'],[['缺任一模态／无正IoU匹配','NoFire','0','无火，不允许UV推翻'],['匹配且s≤0.1','NoFire','0','无火'],['0.1<s<0.3','Transition','(s−0.1)/0.2','无火，暴露过渡状态'],['匹配且s≥0.3','FireCandidate','1','仅融合BetP(F)>0.5判火，否则无火']])
    p('空间状态与融合判决分开记录：FireCandidate只表示空间合格，不等于最终Fire；Transition仍为binary0。判决顺序是先空间资格再D-S/BetP。CSV保存原始三质量、g、IoU、gate_pass、semantic_state及各最终hard label，不声称硬无火全来自D-S质量本身。')
    p('全量空间状态计数：'+json.dumps(summary['state_counts'],ensure_ascii=False)+'；按UV参考标签分层：'+json.dumps(summary['state_by_reference'],ensure_ascii=False))
    p('被门控阻止的原始含UV判火数='+str(summary['raw_uv_fire_blocked'])+'；其中UV参考有火数='+str(summary['raw_uv_fire_blocked_reference_positive'])+'。这证明UV不能跨越门控。')
    h('3. 几何与一对一配对／候选选择')
    p('热框真实四角通过FULL affine 2×3矩阵映射到RGB原像素坐标：q=A[:,0:2]p+A[:,2]；矩阵='+json.dumps(cal['affine_matrix_2x3'])+'。使用cv2.intersectConvexConvex计算四边形与RGB矩形真实交并比，不使用映射外包矩形、不加人为偏移、不裁剪四边形。')
    p('Hungarian最大总polygon IoU一对一匹配，dummy允许未匹配；仅正IoU≥1e−12进入配对，无交叠不强行配对。每个RGB／thermal索引最多用一次。未匹配目标保留审计但绝不能独自触发判火。')
    p('帧代表：优先最大匹配IoU；相同IoU用无UV的视觉BetP、最低在场置信度，再候选顺序确定。先选候选再UV融合，绝不用UV标签选目标。多目标之间不累计D-S，以免相关重复框虚增置信度。由于本缓存所有在场火焰分数>0.5，合格视觉配对均BetP>0.5；已核验“任一合格视觉配对判火”与最大IoU代表本数据输出完全一致。')
    p('候选仿射留一RMSE='+f'{cal["leave_one_out_rmse_px"]:.4f}'+'px，最大误差='+f'{cal["leave_one_out_max_px"]:.4f}'+'px；fit RMSE='+f'{cal["fit_rmse_px"]:.4f}'+'px，内点6/7。候选配准未验证：仅7点且纵向覆盖窄，同点拟合／留一不能证明跨场景、全视场及视差可靠。严格门控也会放大配准错误造成的漏检。')
    p('选中候选状态：'+json.dumps(summary['selected_status_counts'])+'；双缺框帧='+str(summary['both_missing_frames'])+'，单模态缺框帧='+str(summary['single_missing_frames'])+'，双模态在场但无任何正IoU配对帧='+str(summary['both_present_no_overlap_frames'])+'。')
    h('4. 证据构造、折扣与经典D-S')
    p('视觉在场BPA=(0.8c,0.8(1−c),0.2)，缺失BPA=(0,0,1)。c是检测置信度，不是独立校准概率。IoU同时影响证据和最终资格：g=clip((s−τL)/(τH−τL),0,1)，对两种视觉证据均折扣mF′=g mF、mN′=g mN、mU′=1−g+g mU，然后经典D-S组合。低IoU不直接生成“物理无火”证据，硬无火由显式政策门控给出。')
    p('经典D-S：K=aF bN+aN bF；F支持=aF bF+aF bU+aU bF；N支持=aN bN+aN bU+aU bN；U支持=aU bU；除以支持总和（1−K）。总冲突回退全未知并标记。普通DS无空间对照保留同候选／同基础BPA，单独列此前历史候选普通DS以避免悄改63.4817%的旧结果。')
    p('瞬时UV BPA预先固定：UV1=(0.9,0,0.1)，UV0=(0,0.9,0.1)；先视觉D-S再组合UV。UV无空间框，不参与IoU和候选选择。原始DS＋UV可在缺失视觉时BetP>0.5，但最终hard label仍被空间门控置0。不采用空间BPA=(g,1−g,0)，避免高IoU构成独断火质量强行覆盖置信度。')
    p('Tmax原值保存；单位及校准未知，本新规则不增加温度阈值。瞬时UV与原C++时序UV是不同BPA；本报告不将新公式描述为原C++忠实复现。旧报告与旧C++保持原状，新策略是明确扩展。')
    h('5. 预先冻结的阈值敏感性，不选择最优')
    p('预先列τH=0.1、0.2、0.3、0.5；各τL=τH/3。每组重新计算g、折扣及硬标签，全部结果只作诊断；主结果始终τH=0.3，不根据本测试参考标签重新选择。')
    table(['τL','τH','方法','TP','FP','FN','TN','Acc','R','F1'],[[f'{v["tau_l"]:.4f}',f'{v["tau_h"]:.1f}',NAMES[k],m['TP'],m['FP'],m['FN'],m['TN'],percent(m['accuracy']),percent(m['recall']),percent(m['F1'])] for v in summary['sensitivity'] for k,m in v['methods'].items()])
    h('6. 真实模型、源文件保留及可复算性')
    for mode,pr in summary['provenance'].items():
        p(mode+'：'+pr['metadata']['architecture']+'（不是PPYOLOv2）；checkpoint SHA256='+pr['metadata']['checkpoint_sha256']+'；预测SHA256='+pr['sha256']+'；1528图，框数='+str(pr['boxes'])+'，无框帧='+str(pr['missing_frames'])+'。')
    p('严格复用用户权重YOLO26n真实fire检测score>0.5。每图文件名／pair_id／source_line／SHA256／尺寸、两份权重SHA256、仿射来源SHA256全部核验；3056图像未伪造或替换。模型训练／测试重叠未独立审计。')
    p('冻结UTC='+summary['freeze']['utc']+'；配置SHA256='+summary['freeze']['sha256']+'。计算指标前冻结。旧fusion所有文件、已有reports及原C++源全部逐文件SHA256前后核验。新增输出均在fusion/spatial_gate_ds及指定新报告路径，源数据只读。')
    h('7. 实际叠加图与交付')
    p('真实样例显示RGB框、映射thermal四边形、NoFire／Transition／FireCandidate、g、原始BetP及硬判决；不是独立定位真值。')
    for ex in summary['examples']: p('pair '+str(ex['pair_index'])+' '+ex['state']); d.add_picture(ex['path'],width=Inches(6.5))
    p('交付：全1528逐帧CSV／JSON、指标CSV／JSON、敏感性CSV、全部候选几何与证据JSON、冻结配置／回执、代码／TDD RED-GREEN日志、叠加图、核验清单及本中文DOCX。CSV回读独立重算混淆矩阵与指标，DOCX ZIP CRC／文本表格／嵌图回读与SHA256核验。尚缺独立场景标签及全视场配准验证；不把循环评价或工程硬无火称为物理证明。')
    d.save(REPORT)

def run():
    cfg=load(OUT/'frozen_config.json'); receipt=load(OUT/'config_freeze_receipt.json'); assert sha(OUT/'frozen_config.json')==receipt['sha256'] and receipt['before_metrics']
    files=[p for p in OUT.parent.rglob('*') if p.is_file() and OUT not in p.parents]
    files += [p for p in (ROOT/'reports').rglob('*') if p.is_file() and p!=REPORT]
    files += [p for p in Path('C:/lyk/paper/bpa_test').rglob('*') if p.is_file() and p.suffix in ['.cpp','.h']]
    files += [SOURCE/'paper_sample.txt']
    protected={str(p):sha(p) for p in files}; save('protected_before_hashes.json',protected)
    cal=load(AFF/'candidate_affine_calibration.json'); assert sha(cal['source'])==cal['source_sha256']; a=cal['affine_matrix_2x3']
    samples=[legacy.parse_sample_line(l) for l in (SOURCE/'paper_sample.txt').read_text(encoding='utf-8-sig').splitlines()]
    assert len(samples)==1528 and len({x.rgb for x in samples})==len({x.thermal for x in samples})==1528
    pred={}; provenance={}
    for mode in ['rgb','thermal']:
        path=OLD/(mode+'_predictions.json'); data=load(path); records=data['images']; assert len(records)==1528 and len({r['pair_index'] for r in records})==1528
        for i,(r,sm) in enumerate(zip(records,samples)):
            assert r['pair_index']==i and r['source_line']==i+1 and r['image_name']==getattr(sm,mode)
            assert sha(r['image_path'])==r['image_sha256']; im=cv2.imread(r['image_path']); assert im is not None and im.shape[:2]==(r['height'],r['width'])
            for b in r['detections']: assert b['score']>.5 and b['class_id']==0; s.polygon(b['bbox_xyxy'])
        assert sha(data['metadata']['checkpoint'])==data['metadata']['checkpoint_sha256']
        pred[mode]=records; provenance[mode]=dict(metadata=data['metadata'],sha256=sha(path),boxes=sum(len(r['detections']) for r in records),missing_frames=sum(not r['detections'] for r in records))
    with (PRIOR/'fusion_results.csv').open(encoding='utf-8-sig') as f: prior=list(csv.DictReader(f))
    assert len(prior)==1528
    rows=[]; details=[]
    for i,sm in enumerate(samples):
        assert int(prior[i]['pair_index'])==i and int(prior[i]['UV_reference_label'])==sm.uv
        det=gate.frame(pred['rgb'][i]['detections'],pred['thermal'][i]['detections'],a,sm.uv,cfg['nominal_tau_l'],cfg['nominal_tau_h'],cfg['base_ignorance'],cfg['uv_instant_reliability'])
        candidates=det['candidates']; matched=[c for c in candidates if c['status']=='matched']
        assert len({c['rgb_index'] for c in matched})==len({c['thermal_index'] for c in matched})==len(matched)
        assert det['final_visual']==int(any(c['final_visual'] for c in matched))
        assert det['final_visual_uv']==int(any(c['final_visual_uv'] for c in matched))
        ordinary=np.asarray(det['ordinary_mass']); um=np.asarray(det['instant_uv_mass']); ordinary_uv=s.ds(ordinary,um)[0]
        row=dict(pair_index=i,source_line=i+1,rgb_name=sm.rgb,thermal_name=sm.thermal,Tmax=sm.temperature,UV_input=sm.uv,UV_reference_label=sm.uv,rgb_count=len(pred['rgb'][i]['detections']),thermal_count=len(pred['thermal'][i]['detections']),selected_candidate=det['selected_candidate'],rgb_index=det['rgb_index'],thermal_index=det['thermal_index'],polygon_iou=det['iou'],matching_status=det['status'],g=det['g'],reliability=det['g'],semantic_state=det['state'],gate_pass=int(det['gate_pass']),tau_l=cfg['nominal_tau_l'],tau_h=cfg['nominal_tau_h'],rgb_confidence=det['rgb_confidence'],thermal_confidence=det['thermal_confidence'],rgb_binary=int(bool(pred['rgb'][i]['detections'])),thermal_binary=int(bool(pred['thermal'][i]['detections'])),prior_ordinary_ds_binary=int(prior[i]['ordinary_ds_binary']),gate_visual_binary=det['final_visual'],gate_visual_uv_binary=det['final_visual_uv'],uv_alone_binary=sm.uv)
        for name,m in dict(ordinary_ds=ordinary,ordinary_uv=ordinary_uv,raw_spatial_ds=det['visual_mass'],raw_spatial_uv=det['visual_uv_mass']).items():
            assert np.isfinite(m).all() and min(m)>=0 and np.isclose(sum(m),1)
            for key,val in zip(['fire','nofire','uncertain'],m): row[name+'_'+key]=float(val)
            row[name+'_betp']=float(m[0]+.5*m[2]); row[name+'_binary']=int(row[name+'_betp']>.5)
        if not det['gate_pass']: assert row['gate_visual_binary']==row['gate_visual_uv_binary']==0
        rows.append(row); details.append(dict(pair_index=i,**det))
    with (OUT/'fusion_results.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    save('fusion_results.json',dict(rows=rows)); save('geometry_evidence.json',details)
    y=[sm.uv for sm in samples]; methods={name:metric(y,[r[name+'_binary'] for r in rows]) for name in NAMES}
    sensitivity=[]; sensrows=[]
    for thresholds in cfg['sensitivity']:
        tl,th=thresholds['tau_l'],thresholds['tau_h']; result=[]
        for i,sm in enumerate(samples):
            det=gate.frame(pred['rgb'][i]['detections'],pred['thermal'][i]['detections'],a,sm.uv,tl,th,cfg['base_ignorance'],cfg['uv_instant_reliability'])
            result.append(det)
        mm={name:metric(y,[r[field] for r in result]) for name,field in [('gate_visual','final_visual'),('gate_visual_uv','final_visual_uv')]}
        sensitivity.append(dict(tau_l=tl,tau_h=th,methods=mm,state_counts=dict(Counter(r['state'] for r in result))))
        for name,m in mm.items(): sensrows.append(dict(tau_l=tl,tau_h=th,method=name,**m))
        if th==cfg['nominal_tau_h']: assert mm['gate_visual']==methods['gate_visual'] and mm['gate_visual_uv']==methods['gate_visual_uv']
    metrics=dict(reference='UV-reference agreement; UV input branches circular, not independent truth',label_counts=dict(Counter(y)),methods=methods,thresholds=dict(tau_l=.1,tau_h=.3,betp_threshold=.5),config_sha256=receipt['sha256'],sensitivity=sensitivity)
    save('metrics.json',metrics)
    for name,records in [('metrics.csv',[dict(method=k,method_cn=NAMES[k],**v) for k,v in methods.items()]),('sensitivity_metrics.csv',sensrows)]:
        with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f:
            w=csv.DictWriter(f,fieldnames=list(records[0])); w.writeheader(); w.writerows(records)
    summary=dict(created_utc=datetime.now(timezone.utc).isoformat(),total_pairs=1528,source_sample_sha256=sha(SOURCE/'paper_sample.txt'),config=cfg,freeze=receipt,calibration=cal,provenance=provenance,metrics=metrics,sensitivity=sensitivity,state_counts=dict(Counter(r['semantic_state'] for r in rows)),state_by_reference={str(label):dict(Counter(r['semantic_state'] for r in rows if r['UV_reference_label']==label)) for label in [0,1]},selected_status_counts=dict(Counter(r['matching_status'] for r in rows)),both_missing_frames=sum(r['rgb_count']==r['thermal_count']==0 for r in rows),single_missing_frames=sum((r['rgb_count']==0)!=(r['thermal_count']==0) for r in rows),both_present_no_overlap_frames=sum(r['rgb_count']>0 and r['thermal_count']>0 and all(c['status']!='matched' for c in details[r['pair_index']]['candidates']) for r in rows),raw_uv_fire_blocked=sum(r['raw_spatial_uv_binary']==1 and r['gate_visual_uv_binary']==0 for r in rows),raw_uv_fire_blocked_reference_positive=sum(r['raw_spatial_uv_binary']==1 and r['gate_visual_uv_binary']==0 and r['UV_reference_label']==1 for r in rows))
    summary['examples']=overlays(rows,details,pred,a); save('execution_summary.json',summary); report(summary)
    with (OUT/'fusion_results.csv').open(encoding='utf-8-sig') as f: loaded=list(csv.DictReader(f))
    assert len(loaded)==1528 and [int(r['pair_index']) for r in loaded]==list(range(1528))
    # Independently count CSV pairs rather than trusting in-memory masks/labels.
    for name in NAMES:
        counts=Counter((int(r['UV_reference_label']),int(r[name+'_binary'])) for r in loaded)
        mm=methods[name]; assert [mm['TP'],mm['FP'],mm['FN'],mm['TN']]==[counts[(1,1)],counts[(0,1)],counts[(1,0)],counts[(0,0)]]
        assert metric([int(r['UV_reference_label']) for r in loaded],[int(r[name+'_binary']) for r in loaded])==mm
    with zipfile.ZipFile(REPORT) as z: assert z.testzip() is None; images=sum(n.startswith('word/media/') for n in z.namelist())
    doc=Document(REPORT); text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for rr in t.rows for c in rr.cells]); assert all(x in text for x in ['1528','循环','未验证','Hungarian','Transition','不一致＝','0.3']); assert images==len(summary['examples'])
    (OUT/'report_readback.txt').write_text(text,encoding='utf-8')
    assert all(sha(p)==h for p,h in protected.items()); assert sha(OUT/'frozen_config.json')==receipt['sha256']
    verification=dict(csv_rows=1528,unique_pair_ids=1528,current_image_hashes_verified=3056,metrics_recomputed_from_csv=True,independent_confusion_counter=True,docx_zip_crc_ok=True,docx_readback=True,embedded_images=images,old_files_unchanged=True,protected_files_count=len(protected),config_frozen_before_metrics=True,config_unchanged=True,gate_no_uv_override_verified=True,one_to_one_and_real_polygon_iou=True,any_eligible_positive_equals_max_iou_candidate_on_real_cache=True,report_path=str(REPORT),report_sha256=sha(REPORT),csv_sha256=sha(OUT/'fusion_results.csv'),metrics_sha256=sha(OUT/'metrics.json'))
    save('deliverable_verification.json',verification)
    print(json.dumps(dict(methods=methods,state_counts=summary['state_counts'],state_by_reference=summary['state_by_reference'],raw_uv_fire_blocked=summary['raw_uv_fire_blocked'],sensitivity=sensitivity,verification=verification),ensure_ascii=False,indent=2))

if __name__=='__main__': run()
