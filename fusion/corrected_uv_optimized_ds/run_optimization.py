from pathlib import Path
import sys,json,csv,hashlib,datetime,itertools,collections,zipfile,time
import numpy as np
OUT=Path(__file__).resolve().parent;ROOT=OUT.parent.parent
OLD=ROOT/'fusion/supplied320_paper_experiment';CORRECTED=ROOT/'fusion/supplied320_corrected_uv'
SOURCE=Path('C:/lyk/paper_data/paper_data_/paper_data/paper_sample.txt')
sys.path.insert(0,str(OLD))
from fuse320 import corners,transform,polygon_iou,metrics,aligned
from optimized_core import frame,select
EXPECTED='7fa9997f14d65213063651f8448dc56344f797911185a2d6170147f1960bd483'
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
now=lambda:datetime.datetime.now(datetime.timezone.utc).isoformat()
load=lambda p:json.loads(Path(p).read_text(encoding='utf-8-sig'))
def save(name,obj):
 with (OUT/name).open('x',encoding='utf-8') as f:json.dump(obj,f,indent=2,ensure_ascii=False,allow_nan=False)
def csvsave(name,rows):
 with (OUT/name).open('x',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def temporal_groups(names):
 times=[]
 for name in names:
  x=name.split('_');times.append(datetime.datetime.strptime('_'.join(x[:6]),'%Y_%m_%d_%H_%M_%S').timestamp()+int(x[6])/10000)
 groups=[0]
 for a,b in zip(times,times[1:]):groups.append(groups[-1]+int(b-a>300 or b<a))
 return groups
def folds(groups):
 for g in np.unique(groups):yield np.flatnonzero(groups!=g),np.flatnonzero(groups==g)
def parse():
 from importlib.util import spec_from_file_location,module_from_spec
 sp=spec_from_file_location('corrected_reader',CORRECTED/'rerun_corrected_uv.py');m=module_from_spec(sp);sp.loader.exec_module(m)
 rows=m.parse(SOURCE.read_text(encoding='utf-8-sig'))
 assert sha(SOURCE)==EXPECTED and len(rows)==1528 and collections.Counter(r['UV_reference'] for r in rows)=={1:1288,0:240}
 return rows

def prepare():
 rows=parse();groups=np.array(temporal_groups([r['rgb_name'] for r in rows]));grid=[]
 axes=dict(tR=[.2,.35,.5],tT=[.01,.05,.1],gate=[.01,.05,.1,.3],rR=[.7,1.],rT=[.1,.3,.5,1.],u0=[0.,.1],soft=[False,True],threshold=[.4,.5,.6])
 for v in itertools.product(*axes.values()):grid.append(dict(zip(axes,v)))
 protected={str(p):sha(p) for d in [OLD,CORRECTED] for p in d.rglob('*') if p.is_file()}
 proto=load(OLD/'frozen_protocol.json')
 for m,v in proto['model_sources'].items(): assert sha(v['source'])==v['sha256'];protected[v['source']]=sha(v['source'])
 protected[str(SOURCE)]=sha(SOURCE);protected[proto['geometry']['source']]=sha(proto['geometry']['source'])
 group_rows=[]
 for g in np.unique(groups):
  ix=np.flatnonzero(groups==g);yy=[rows[i]['UV_reference'] for i in ix]
  group_rows.append(dict(group=int(g),first=int(ix[0]),last=int(ix[-1]),N=len(ix),positive=sum(yy),negative=len(yy)-sum(yy),first_name=rows[ix[0]]['rgb_name'],last_name=rows[ix[-1]]['rgb_name']))
 save('pre_search_protocol.json',dict(created_utc=now(),source_sha256=EXPECTED,source=str(SOURCE),N=1528,positive=1288,negative=240,group_rule='Filename timestamp chronological gap >300 seconds or backward reset. Defined WITHOUT UV labels. Leave one entire proxy group out.',groups=group_rows,all_groups=groups.tolist(),search_grid=grid,search_trials=len(grid),objective='Lexicographic minimize calibration FP, then maximize calibration TP, then first deterministic grid index. No heldout labels in select().',exposure='All 1528 previously inspected; internal exploratory diagnosis, NOT independent test. Group proxy not certified scenes.',UV_runtime_input=False,BPA='Positive-only: mF=(1-u0)*r*c; mN=0; mU=1-mF. Raw c not calibrated probability. DS normalized combination K=0. Thermal reliability optionally rT*clip(IoU/.3,0,1); RGB remains asymmetric strong reference. Both REAL detector boxes and IoU>=gate mandatory; missing consensus binary0 with mass ignorance.',decision='max selected matched mF > threshold (strict); dummy Hungarian maximum total eligible IoU, one-to-one.',fresh_inference=dict(conf=.01,imgsz=320,rect=False,half=False,classes=[0],device='cuda:0',batch=32,max_det=300,iou_argument=.7,end2end=True,nms='end2end conventional NMS bypassed'),geometry=proto['geometry'],models=proto['model_sources'],protected_hashes=protected,roi=False,retraining=False))
 save('pre_search_freeze_receipt.json',dict(utc=now(),protocol_sha256=sha(OUT/'pre_search_protocol.json'),before_any_search=True))
 print('PREPARED',len(grid),'trials; proxy groups',group_rows,flush=True)

def geom(rb,tb,A):
 return np.array([[polygon_iou(corners(r['bbox']),transform(corners(t['bbox']),A)) for t in tb] for r in rb]).reshape(len(rb),len(tb))

def fresh():
 import torch,cv2
 from ultralytics import YOLO
 rows=parse();p=load(OUT/'pre_search_protocol.json');A=np.array(p['geometry']['affine_matrix_2x3']);cfg=p['search_grid'][0]
 cached={m:load(OLD/(m+'_predictions.json')) for m in ['rgb','thermal']}
 smoke=[]
 for i in range(3):
  rb=cached['rgb'][i]['detections'];tb=cached['thermal'][i]['detections'];s=geom(rb,tb,A);o=frame(rb,tb,s,cfg);assert abs(o['mF']+o['mN']+o['mU']-1)<1e-12;smoke.append(dict(pair_index=i,output=o))
 save('cached_smoke_3pairs.json',dict(real_cached=True,results=smoke,utc=now()))
 assert torch.cuda.is_available();models={m:YOLO(v['source']) for m,v in p['models'].items()}
 assert all(m.model.end2end and len(m.names)==1 and m.names[0].lower()=='fire' for m in models.values())
 shapes=[]
 def forward(model,paths,ids):
  ims=[cv2.imread(str(x)) for x in paths];assert all(x is not None for x in ims)
  seen=[];hook=model.model.register_forward_pre_hook(lambda m,args:seen.append(list(args[0].shape)))
  results=list(model.predict(ims,imgsz=320,rect=False,conf=.01,iou=.7,max_det=300,classes=[0],device=0,batch=32,half=False,augment=False,verbose=False,save=False));hook.remove()
  assert seen and all(s[-2:]==[320,320] for s in seen);shapes.extend(seen);out=[]
  for result,im,path,i in zip(results,ims,paths,ids):
   np.testing.assert_array_equal(result.orig_img,im);ds=[]
   for x,y,X,Y,c,k in result.boxes.data.cpu().tolist():
    assert int(k)==0 and c>.01 and X>x and Y>y;ds.append(dict(bbox=[x,y,X,Y],score=c,class_id=0))
   out.append(dict(pair_index=i,image_name=path.name,image_path=str(path),image_sha256=sha(path),orig_shape=list(im.shape[:2]),detections=ds,model_input_shapes=seen,device='cuda:0'))
  assert len(out)==len(paths);return out
 paths={m:[SOURCE.parent/d/r[k] for r in rows] for m,d,k in [('rgb','bgr','rgb_name'),('thermal','thermal','thermal_name')]}
 start=time.perf_counter();sm={m:forward(models[m],paths[m][:3],list(range(3))) for m in models};save('fresh_smoke_3pairs.json',sm)
 for m in models:
  out=[]
  for first in range(0,len(rows),32):
   ids=list(range(first,min(first+32,len(rows))));out.extend(forward(models[m],paths[m][first:first+32],ids))
   if first%320==0:print('FRESH',m,len(out),flush=True)
  aligned(out,[x.name for x in paths[m]]);save(m+'_predictions.json',out)
  assert all(x['image_sha256']==y['image_sha256'] for x,y in zip(out,cached[m]))
 save('fresh_forward_receipt.json',dict(utc=now(),full_images=3056,smoke_images=6,seconds=time.perf_counter()-start,gpu=torch.cuda.get_device_name(0),torch=torch.__version__,shapes=shapes,weight_sha256={m:sha(v['source']) for m,v in p['models'].items()},source_hash=sha(SOURCE),input_size=320,conf=.01,actual_full_forward=True))

def search():
 p=load(OUT/'pre_search_protocol.json');assert sha(OUT/'pre_search_protocol.json')==load(OUT/'pre_search_freeze_receipt.json')['protocol_sha256']
 rows=parse();y=np.array([r['UV_reference'] for r in rows]);groups=np.array(p['all_groups']);grid=p['search_grid'];A=np.array(p['geometry']['affine_matrix_2x3'])
 rgb=load(OUT/'rgb_predictions.json');thermal=load(OUT/'thermal_predictions.json');ious=[geom(r['detections'],t['detections'],A) for r,t in zip(rgb,thermal)]
 predictions=np.zeros((len(grid),len(rows)),bool);scores=np.zeros((len(grid),len(rows)),float)
 geometry_keys=list(dict.fromkeys((c['tR'],c['tT'],c['gate']) for c in grid));coverage=[]
 for no,(tr,tt,gate) in enumerate(geometry_keys):
  cfg=dict(grid[0],tR=tr,tT=tt,gate=gate);pairs=[]
  for i in range(len(rows)):
   o=frame(rgb[i]['detections'],thermal[i]['detections'],ious[i],cfg)
   pairs.extend((i,x['rgb_score'],x['thermal_score'],x['iou']) for x in o['matches'])
  a=np.array(pairs);ix=a[:,0].astype(int) if len(a) else np.array([],int)
  present=np.zeros(len(rows),bool);present[ix]=True
  coverage.append(dict(tR=tr,tT=tt,gate=gate,matched_frames=int(sum(present)),positive_matched=int(sum(present[y==1])),negative_matched=int(sum(present[y==0])),positive_uncovered=int(sum(~present[y==1])),coverage_FNR_lower_bound=float(sum(~present[y==1])/sum(y))))
  for k,c in enumerate(grid):
   if (c['tR'],c['tT'],c['gate'])!=(tr,tt,gate):continue
   s=np.zeros(len(rows));rr=(1-c['u0'])*c['rR']
   if len(a):
    rt=(1-c['u0'])*c['rT']*(np.clip(a[:,3]/.3,0,1) if c['soft'] else 1)
    f=1-(1-rr*a[:,1])*(1-rt*a[:,2]);np.maximum.at(s,ix,f)
   scores[k]=s;predictions[k]=present&(s>c['threshold'])
  print('GEOMETRY',no+1,'/',len(geometry_keys),'matched',sum(present),flush=True)
 np.savez_compressed(OUT/'grid_predictions.npz',predictions=predictions,mF=scores)
 save('candidate_coverage_upper_bounds.json',coverage)
 full_index=select(predictions,y,np.arange(len(rows)))
 oof=np.zeros(len(rows),bool);folds_out=[];trials=[];assigned=np.full(len(rows),-1)
 for g,(train,test) in zip(np.unique(groups),folds(groups)):
  assert not set(groups[train])&set(groups[test]);k=select(predictions,y,train)
  freeze=dict(utc=now(),fold=int(g),selected_grid_index=k,parameters=grid[k],train_indices=train.tolist(),test_indices=test.tolist(),selection_metrics=metrics(y[train].tolist(),predictions[k,train].astype(int).tolist()),selection_only_train=True)
  save('fold_'+str(g)+'_frozen.json',freeze)
  freeze_sha=sha(OUT/('fold_'+str(g)+'_frozen.json'))
  # Heldout labels first enter metric evaluation AFTER freezing this fold's parameters.
  oof[test]=predictions[k,test];assigned[test]=k
  folds_out.append(dict(fold=int(g),config=grid[k],grid_index=k,freeze_sha256=freeze_sha,train=freeze['selection_metrics'],test=metrics(y[test].tolist(),oof[test].astype(int).tolist())))
  for j,c in enumerate(grid):trials.append(dict(scope='calibration_fold_'+str(g),grid_index=j,**c,**metrics(y[train].tolist(),predictions[j,train].astype(int).tolist())))
 save('full_fit_frozen.json',dict(utc=now(),parameters=grid[full_index],grid_index=full_index,claim='ALL-LABEL exploratory descriptive fit, NOT heldout validation'))
 full=predictions[full_index]
 for j,c in enumerate(grid):trials.append(dict(scope='full_descriptive',grid_index=j,**c,**metrics(y.tolist(),predictions[j].astype(int).tolist())))
 csvsave('finite_complete_trials.csv',trials);save('fold_metrics.json',folds_out)
 details=[];evidence=[]
 for i,r in enumerate(rows):
  fo=frame(rgb[i]['detections'],thermal[i]['detections'],ious[i],grid[full_index]);oo=frame(rgb[i]['detections'],thermal[i]['detections'],ious[i],grid[assigned[i]])
  assert fo['fire']==int(full[i]) and oo['fire']==int(oof[i]);assert abs(fo['mF']-scores[full_index,i])<1e-12
  best=fo['selected'];ob=oo['selected']
  details.append(dict(pair_index=i,**r,proxy_group=int(groups[i]),full_grid_index=full_index,oof_grid_index=int(assigned[i]),full_fire=int(full[i]),oof_fire=int(oof[i]),rgb_baseline=int(any(b['score']>.5 for b in rgb[i]['detections'])),thermal_baseline=int(any(b['score']>.5 for b in thermal[i]['detections'])),rgb_operating=int(any(b['score']>=.2 for b in rgb[i]['detections'])),thermal_operating=int(any(b['score']>=.1 for b in thermal[i]['detections'])),full_mass_fire=fo['mF'],full_mass_nofire=fo['mN'],full_mass_uncertain=fo['mU'],full_conflict=fo['K'],full_selected_iou=best['iou'] if best else 0.,full_selected_rgb_score=best['rgb_score'] if best else 0.,full_selected_thermal_score=best['thermal_score'] if best else 0.,full_rR_effective=best['rR_effective'] if best else 0.,full_rT_effective=best['rT_effective'] if best else 0.,oof_mass_fire=oo['mF'],oof_mass_nofire=oo['mN'],oof_mass_uncertain=oo['mU'],oof_selected_iou=ob['iou'] if ob else 0.))
  evidence.append(dict(pair_index=i,full=fo,oof=oo,thermal_mapped_polygon=transform(corners(best['thermal_bbox']),A).tolist() if best else None))
 csvsave('frame_predictions.csv',details);save('matched_evidence.json',evidence)
 results={k:metrics(y.tolist(),[r[k] for r in details]) for k in ['rgb_baseline','thermal_baseline','rgb_operating','thermal_operating','full_fire','oof_fire']}
 old=load(CORRECTED/'metrics.json')['results']
 for k in ['rgb_baseline','thermal_baseline','rgb_operating','thermal_operating']:
  assert all(results[k][v]==old[k][v] for v in ['TP','FP','FN','TN']),(k,results[k],old[k])
 targets={}
 for k in ['full_fire','oof_fire']:
  v=results[k];targets[k]=dict(literal_strict_RGB_FPR_improvement_possible=False,literal_strict_RGB_FPR_improvement_achieved=v['FPR']<results['rgb_baseline']['FPR'],thermal_goal=v['FPR']<=results['thermal_baseline']['FPR'] and v['FNR']<results['thermal_baseline']['FNR'],stronger_RGB_goal=v['FPR']<=results['rgb_baseline']['FPR'] and v['FNR']<results['rgb_baseline']['FNR'],fair_low_RGB_FPR_improvement=v['FPR']<results['rgb_operating']['FPR'])
 save('metrics.json',dict(reference='Corrected fourth-column UV, primary NO UV input',source_sha256=EXPECTED,results=results,targets=targets,full_parameters=grid[full_index],folds=folds_out))
 csvsave('metrics.csv',[dict(method=k,**v) for k,v in results.items()])
 groupmetrics=[]
 for g in np.unique(groups):
  ix=np.flatnonzero(groups==g)
  for k in results:groupmetrics.append(dict(group=int(g),method=k,**metrics(y[ix].tolist(),[details[i][k] for i in ix])))
 csvsave('group_metrics.csv',groupmetrics)
 save('search_execution.json',dict(utc=now(),trials=len(grid),all_trials_finite=True,calibration_scopes=6,full_scope=1,trials_csv_rows=len(trials),N=1528,all_fold_params_frozen_before_test_metrics=True,primary_UV_runtime=False,fresh_baselines_match_corrected_cache=True,full_fit_descriptive_only=True))
 print(json.dumps(dict(results=results,targets=targets,full_parameters=grid[full_index]),indent=2),flush=True)


def report():
 from docx import Document
 from docx.shared import Pt,Inches
 from docx.oxml.ns import qn
 p=load(OUT/'pre_search_protocol.json');m=load(OUT/'metrics.json');v=m['results'];doc=Document()
 st=doc.styles['Normal'];st.font.name='宋体';st.font.size=Pt(10.5);st.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
 doc.add_heading('修正UV标签的空间IoU与D-S优化验证报告',0)
 doc.add_heading('摘要',1)
 doc.add_paragraph('使用用户指定best.pt（RGB）及best-2.pt（热成像）的YOLO26n冻结权重，以真实320×320输入对1528组图像执行新一轮完整前向（3056张），候选阈值从旧缓存>0.05降低为>0.01。当前第四列UV为1288个正例与240个负例，原始文本SHA256='+EXPECTED+'。主方案不输入UV，不使用标签、时间或文件名作为判火特征。参数检索限于预先冻结的3456项；标签仅用于训练折参数选择及评价。')
 doc.add_heading('1 目标解释与数据边界',1)
 doc.add_paragraph('标准RGB >0.5已FP=0，严格降低RGB误报率在数学上不可能。因此单独核验：（a）相对标准热成像FPR不升且FNR严格下降；（b）更强目标：保持标准RGB FPR=0并将FNR降至26/1288以下；（c）相对RGB ≥0.2公平低阈值候选的误报率改善。不能将（c）冒充相对标准RGB的严格改善。全部1528组已在历史实验中被查看，本次留组法仅是内部探索性诊断，不是新独立测试。')
 doc.add_paragraph('时间代理组在调参之前，仅按图像文件名时间戳相邻间隔>300秒或时间倒退划分，不按UV标签连续段划分。所得6组见下表；含单一类别的组FPR或FNR为未定义None，汇总用 pooled OOF。训练折包含其余完整组，先冻结选定参数，再一次评价该保留组；不得据保留组结果重新调整。')
 t=doc.add_table(rows=1,cols=6);t.style='Table Grid'
 for cell,s in zip(t.rows[0].cells,['组','首行索引','末行索引','N','有火','无火']):cell.text=s
 for g in p['groups']:
  for cell,k in zip(t.add_row().cells,['group','first','last','N','positive','negative']):cell.text=str(g[k])
 doc.add_heading('2 检测、配准与证据建模',1)
 doc.add_paragraph('模型为YOLO26n而非PPYOLOv2，不替换安全训练模型、不重新训练。实际新前向：conf=0.01、imgsz=320、rect=False、half=False、batch=32、max_det=300、fire唯一类、CUDA RTX5070。端到端输出绕过常规NMS，传入的iou=0.7不是实际NMS阈值。每张原图散列与受保护320缓存核对，权重散列和历史产物保持不变。')
 doc.add_paragraph('热成像框四角经冻结2×3仿射矩阵映射到RGB原坐标，使用实际凸多边形IoU。该矩阵仅由7点拟合，留一RMSE=5.771805像素，窄带点覆盖及视差未独立验证，仍为候选标定。筛分数后用带dummy的匈牙利算法最大化合格IoU总和，保证一对一。仅实际双检测且IoU≥门槛可判火；单模态、无框、无空间匹配均输出二分类0，质量保留未知态，不把缺框物理解释为无火。')
 doc.add_paragraph('BPA正类简单支持：m(F)=(1−u0)rc，m(N)=0，m(Θ)=1−m(F)。c为原始检测置信度，不声称校准概率。可靠性非对称：RGB rR，热成像rT；可选软IoU令热成像有效rT=rT×clip(IoU/0.3,0,1)，RGB不随弱空间折扣共同衰减。标准D-S归一化组合；本模型K=0，mF=1−(1−mR(F))(1−mT(F))，不能称冲突消解。最终取配对中最大mF并严格>决策门槛。该方案是空间门控、证据折扣和BPA变更的联合优化，不将全部收益归因于IoU或D-S。')
 doc.add_paragraph('预定义范围：'+json.dumps({k:sorted(set(c[k] for c in p['search_grid'])) for k in p['search_grid'][0]},ensure_ascii=False)+'。目标为校准折FP最小、再TP最大、同分按固定检索顺序。无监督概率校准器/分类器、无UV运行分支、无ROI与温度模型。')
 doc.add_heading('3 同口径结果',1)
 labels={'rgb_baseline':'RGB >0.5','thermal_baseline':'热成像 >0.5','rgb_operating':'RGB ≥0.2','thermal_operating':'热成像 ≥0.1','full_fire':'全数据调优（描述性）','oof_fire':'留组OOF（内部诊断）'}
 for fields in [['TP','FP','FN','TN'],['FPR','FNR','precision','recall','F1','accuracy']]:
  t=doc.add_table(rows=1,cols=1+len(fields));t.style='Table Grid'
  for cell,s in zip(t.rows[0].cells,['方法']+fields):cell.text=s
  for method in labels:
   cells=t.add_row().cells;cells[0].text=labels[method]
   for cell,k in zip(cells[1:],fields):cell.text=str(v[method][k]) if k in ['TP','FP','FN','TN'] else ('未定义' if v[method][k] is None else f'{100*v[method][k]:.4f}%')
 doc.add_paragraph('FPR=FP/(FP+TN)，FNR=FN/(TP+FN)，Precision=TP/(TP+FP)，Recall=TP/(TP+FN)，F1=2TP/(2TP+FP+FN)，Accuracy=(TP+TN)/N。单模态阈值固定，不通过弱化基线制造提升。新前向标准与公平候选四行混淆矩阵均与当前修正UV320缓存一致。')
 doc.add_paragraph('全数据描述性最佳参数：'+json.dumps(m['full_parameters'],ensure_ascii=False))
 for method,tg in m['targets'].items():doc.add_paragraph(labels[method]+'目标判定：'+json.dumps(tg,ensure_ascii=False))
 doc.add_heading('4 留组结果及能力上界',1)
 t=doc.add_table(rows=1,cols=7);t.style='Table Grid'
 for cell,s in zip(t.rows[0].cells,['组','TP','FP','FN','TN','FPR','FNR']):cell.text=s
 for f in m['folds']:
  cells=t.add_row().cells;cells[0].text=str(f['fold'])
  for cell,k in zip(cells[1:],['TP','FP','FN','TN','FPR','FNR']):cell.text=str(f['test'][k])
  doc.add_paragraph('组'+str(f['fold'])+'冻结参数：'+json.dumps(f['config'],ensure_ascii=False)+'；校准统计：'+json.dumps(f['train'],ensure_ascii=False))
 cov=load(OUT/'candidate_coverage_upper_bounds.json');best=min(cov,key=lambda x:x['positive_uncovered']);doc.add_paragraph('所有检索候选中最宽覆盖能力上界（不是标签驱动判火）：'+json.dumps(best,ensure_ascii=False)+'。无法形成真实双模态匹配的正例不能由修改BPA恢复；该逻辑上界与分类性能不同，必须同时报告。')
 doc.add_heading('5 有效性风险与结论',1)
 overlap=load(OLD/'source_overlap_audit.json')['stats'];doc.add_paragraph('已知当前TorchDetection训练/验证源与测试图像散列重合：'+json.dumps(overlap,ensure_ascii=False)+'。只能证明源内容重合，不能据此认证历史权重训练成员；因此本报告不具有独立泛化认证效力。时间代理组不能证明场景独立；UV是用户指定参考而非人工框真值，不能计算目标AP。未以UV强制清除误报，无UV循环分支。')
 for k in ['full_fire','oof_fire']:
  z=v[k];doc.add_paragraph(f'{labels[k]}：TP={z["TP"]} FP={z["FP"]} FN={z["FN"]} TN={z["TN"]}；FPR={z["FPR"]*100:.4f}% FNR={z["FNR"]*100:.4f}%。实际目标通过与否见第3节，不作保证。RGB基线误报率为零，任何严格改善RGB FPR的说法均不成立。')
 doc.add_paragraph('可复核产物：pre_search_protocol.json、每折冻结JSON、finite_complete_trials.csv（全部试验）、grid_predictions.npz、frame_predictions.csv、matched_evidence.json、group_metrics.csv、candidate_coverage_upper_bounds.json、fresh_forward_receipt.json及verification.json。全数据最佳只作描述性，不作为OOF通过的替代。')
 path=ROOT/'reports/修正UV_空间DS优化验证报告.docx'
 if path.exists():raise FileExistsError(path)
 doc.save(path)
 with zipfile.ZipFile(path) as z:assert z.testzip() is None
 rd=Document(path);txt='\n'.join(x.text for x in rd.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in rd.tables for r in t.rows)
 (OUT/'report_readback.txt').write_text(txt,encoding='utf-8');assert '留组OOF' in txt and '5 有效性风险与结论' in txt
 save('report_receipt.json',dict(path=str(path),sha256=sha(path),crc_ok=True,tables=len(rd.tables),paragraphs=len(rd.paragraphs)))
 print('REPORT',path,flush=True)

if __name__=='__main__':
 action=sys.argv[1] if len(sys.argv)>1 else 'all'
 if action=='prepare':prepare()
 elif action=='fresh':fresh()
 elif action=='search':search()
 elif action=='report':report()
 elif action=='all':prepare();fresh();search();report()
 else:raise ValueError(action)
