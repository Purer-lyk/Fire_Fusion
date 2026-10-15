from pathlib import Path
import json,csv,hashlib,zipfile,datetime,collections
import numpy as np
from run_optimization import OUT,ROOT,OLD,CORRECTED,SOURCE,EXPECTED,load,sha,temporal_groups,geom
from optimized_core import frame

def independent_metrics(y,p):
 y=np.array(y,int);p=np.array(p,int)
 assert y.shape==p.shape and set(y)<= {0,1} and set(p)<={0,1}
 tp=int(np.count_nonzero((y==1)&(p==1)));fp=int(np.count_nonzero((y==0)&(p==1)));fn=int(np.count_nonzero((y==1)&(p==0)));tn=int(np.count_nonzero((y==0)&(p==0)))
 div=lambda a,b:a/b if b else None
 return dict(N=len(y),TP=tp,FP=fp,FN=fn,TN=tn,FPR=div(fp,fp+tn),FNR=div(fn,fn+tp),precision=div(tp,tp+fp),recall=div(tp,tp+fn),F1=div(2*tp,2*tp+fp+fn),accuracy=div(tp+tn,len(y)))

def main():
 from docx import Document
 protocol=load(OUT/'pre_search_protocol.json');m=load(OUT/'metrics.json');A=np.array(protocol['geometry']['affine_matrix_2x3']);grid=protocol['search_grid']
 assert sha(SOURCE)==EXPECTED==protocol['source_sha256'];assert sha(OUT/'pre_search_protocol.json')==load(OUT/'pre_search_freeze_receipt.json')['protocol_sha256']
 textrows=[s.split() for s in SOURCE.read_text(encoding='utf-8-sig').splitlines() if s.strip()];y=[int(s[3]) for s in textrows];assert collections.Counter(y)=={1:1288,0:240}
 rows=list(csv.DictReader((OUT/'frame_predictions.csv').open(encoding='utf-8-sig')));assert len(rows)==1528
 recomputed={}
 for method,result in m['results'].items():
  v=independent_metrics(y,[int(r[method]) for r in rows]);assert v==result,(method,v,result);recomputed[method]=v
 historical_manifest_mismatches=[]
 for olddir in [OLD,CORRECTED]:
  manifest=load(olddir/'artifact_sha256.json')
  for path,h in manifest.items():
   target=Path(path) if Path(path).is_absolute() else olddir/path
   if sha(target)!=h:
    historical_manifest_mismatches.append(dict(path=str(target),historical_sha256=h,current_sha256=sha(target),pre_search_sha256=protocol['protected_hashes'].get(str(target))))
    if target.suffix=='.docx':assert sha(target)==load(olddir/'validation.json')['report_sha256'],'Previous report differs from latest validation receipt'
 for path,h in protocol['protected_hashes'].items():assert sha(path)==h,('Protected mismatch',path)
 groups=np.array(temporal_groups([x[0] for x in textrows]));assert groups.tolist()==protocol['all_groups']
 pred={mode:load(OUT/(mode+'_predictions.json')) for mode in ['rgb','thermal']};oldpred={mode:load(OLD/(mode+'_predictions.json')) for mode in pred}
 seenhashes=set();evidence=load(OUT/'matched_evidence.json');z=np.load(OUT/'grid_predictions.npz');pp=z['predictions'];scores=z['mF'];assert pp.shape==scores.shape==(3456,1528);assert np.isfinite(scores).all()
 trialrows=list(csv.DictReader((OUT/'finite_complete_trials.csv').open(encoding='utf-8-sig')));assert len(trialrows)==3456*7
 for scope in set(r['scope'] for r in trialrows):assert sorted(int(r['grid_index']) for r in trialrows if r['scope']==scope)==list(range(3456))
 train_checks=0
 for f in m['folds']:
  frozen=load(OUT/('fold_'+str(f['fold'])+'_frozen.json'));tr=np.array(frozen['train_indices']);te=np.array(frozen['test_indices']);yy=np.array(y)
  assert len(tr)+len(te)==1528 and not set(tr)&set(te) and not set(groups[tr])&set(groups[te]);assert sha(OUT/('fold_'+str(f['fold'])+'_frozen.json'))==f['freeze_sha256']
  # Independent lexicographic select; no heldout labels read for parameter selection.
  fp=np.count_nonzero(pp[:,tr][:,yy[tr]==0],axis=1);tp=np.count_nonzero(pp[:,tr][:,yy[tr]==1],axis=1)
  best=min(range(3456),key=lambda k:(int(fp[k]),-int(tp[k]),k));assert best==frozen['selected_grid_index']
  assert independent_metrics(yy[tr],pp[best,tr])==f['train'];assert independent_metrics(yy[te],pp[best,te])==f['test'];train_checks+=1
 fulli=load(OUT/'full_fit_frozen.json')['grid_index'];yy=np.array(y);fp=np.count_nonzero(pp[:,yy==0],axis=1);tp=np.count_nonzero(pp[:,yy==1],axis=1)
 assert fulli==min(range(3456),key=lambda k:(int(fp[k]),-int(tp[k]),k))
 raw_coverage={key:[] for key in ['both_real_detection','raw_iou_at_least_001']};miss=[]
 for i,(source,r,e) in enumerate(zip(textrows,rows,evidence)):
  assert int(r['pair_index'])==i and r['rgb_name']==source[0] and r['thermal_name']==source[1] and int(r['UV_reference'])==y[i]
  for mode,folder,col in [('rgb','bgr',0),('thermal','thermal',1)]:
   rec=pred[mode][i];assert rec['pair_index']==i and rec['image_name']==source[col];assert sha(SOURCE.parent/folder/source[col])==rec['image_sha256']==oldpred[mode][i]['image_sha256'];seenhashes.add(rec['image_sha256'])
   assert all(s[-2:]==[320,320] for s in rec['model_input_shapes']);assert all(b['score']>.01 and b['class_id']==0 for b in rec['detections'])
  rb=pred['rgb'][i]['detections'];tb=pred['thermal'][i]['detections'];s=geom(rb,tb,A)
  raw_coverage['both_real_detection'].append(int(bool(rb and tb)));raw_coverage['raw_iou_at_least_001'].append(int(bool(s.size and (s>=.01).any())))
  for scope,idxcol,firecol in [('full','full_grid_index','full_fire'),('oof','oof_grid_index','oof_fire')]:
   k=int(r[idxcol]);o=frame(rb,tb,s,grid[k]);assert o==e[scope];assert o['fire']==int(r[firecol])==int(pp[k,i]);assert abs(o['mF']-scores[k,i])<1e-12
   assert abs(o['mF']+o['mN']+o['mU']-1)<1e-12
   assert len(set(x['r'] for x in o['matches']))==len(o['matches'])==len(set(x['t'] for x in o['matches']))
   for x in o['matches']:
    assert x['iou']>=grid[k]['gate'] and np.isfinite(x['iou'])
    for name in ['bpa_rgb','bpa_thermal']:assert abs(sum(x[name])-1)<1e-12 and min(x[name])>=0
  if y[i] and not int(r['full_fire']):
   miss.append(dict(pair_index=i,rgb_name=r['rgb_name'],thermal_name=r['thermal_name'],reason='no_rgb_raw_detection' if not rb else 'no_thermal_raw_detection' if not tb else 'no_raw_spatial_iou001_match' if not s.size or not (s>=.01).any() else 'operating_filter_or_mass_decision',rgb_max=max((b['score'] for b in rb),default=0),thermal_max=max((b['score'] for b in tb),default=0),raw_max_iou=float(s.max()) if s.size else 0))
 assert len(seenhashes)==3056
 bounds={k:dict(positive_covered=sum(v[i] for i in range(1528) if y[i]),positive_uncovered=sum(1-v[i] for i in range(1528) if y[i]),minimum_possible_FNR=sum(1-v[i] for i in range(1528) if y[i])/1288) for k,v in raw_coverage.items()}
 bounds['scope']='Diagnostic logical bounds from all real full-image detections score>.01; no labels used as runtime gate, no new search or OOF retune.'
 with (OUT/'raw_candidate_bounds.json').open('x',encoding='utf-8') as f:json.dump(bounds,f,indent=2)
 with (OUT/'full_fit_false_negative_reasons.csv').open('x',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(miss[0]));w.writeheader();w.writerows(miss)
 overlap=load(OLD/'source_overlap_audit.json');assert all(sha(r['path'])==r['sha256'] for r in overlap['records'])
 rec=load(OUT/'report_receipt.json');report=Path(rec['path']);assert sha(report)==rec['sha256']
 with zipfile.ZipFile(report) as f:assert f.testzip() is None
 d=Document(report);assert len(d.tables)==4
 readback='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows)
 assert readback==(OUT/'report_readback.txt').read_text(encoding='utf-8')
 for target in ['全数据调优（描述性）','留组OOF（内部诊断）','严格降低RGB误报率在数学上不可能',EXPECTED]:assert target in readback
 results=dict(success=True,utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),current_source_verified=True,image_hashes_verified=3056,weights_verified=True,old_artifacts_unchanged_vs_pre_search=True,old_reports_match_latest_validation_hashes=True,historical_manifest_mismatches=historical_manifest_mismatches,protected_files_verified=len(protocol['protected_hashes']),live_training_source_hashes_verified=len(overlap['records']),all_1528_geometry_and_final_masses_recomputed=True,all_confusion_metrics_independently_recomputed=True,all_fold_calibration_selection_verified=train_checks,grid_trials=3456,trials_csv_rows=len(trialrows),trial_scope_count=7,all_grid_scores_finite=True,report_crc_ok=True,report_text_and_tables_readback_ok=True,report_sha256=sha(report),raw_candidate_bounds=bounds,false_negative_reason_counts=dict(collections.Counter(r['reason'] for r in miss)),results=recomputed,targets=m['targets'])
 with (OUT/'verification.json').open('x',encoding='utf-8') as f:json.dump(results,f,indent=2,ensure_ascii=False,allow_nan=False)
 with (OUT/'artifact_sha256.json').open('x',encoding='utf-8') as f:json.dump({str(x):sha(x) for x in OUT.rglob('*') if x.is_file() and x.name!='artifact_sha256.json'}|{str(report):sha(report)},f,indent=2,ensure_ascii=False)
 print(json.dumps(results,indent=2,ensure_ascii=False))
if __name__=='__main__':main()
