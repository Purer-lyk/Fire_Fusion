import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent.parent))
import core as c
import execute_fusion_report as existing
import json, csv, hashlib, time, inspect, subprocess, zipfile, shutil
from datetime import datetime, timezone
from collections import Counter
import numpy as np
import cv2
import torch
import ultralytics
from ultralytics import YOLO
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

ROOT=Path('C:/lyk/fire/qualityDetection')
OUT=Path(__file__).resolve().parent
SOURCE=Path('C:/lyk/paper_data/paper_data_/paper_data')
ATT=Path('C:/Users/Administrator/AppData/Local/hermes/attachments')
REPORT=ROOT/'reports/uploadedmodels_UVreference_1528.docx'
METHODS=['rgb_baseline','thermal_baseline','legacy_fixed_01','classic_ds','candidate_geometry_repeated','recommended_spatial_then_ds']
NAMES={'rgb_baseline':'RGB检测基线','thermal_baseline':'热成像检测基线','legacy_fixed_01':'原C++固定0.1两次空间融合（主要）','classic_ds':'经典D-S三证据','candidate_geometry_repeated':'候选几何s重复空间融合（未验证）','recommended_spatial_then_ds':'候选空间RGB/热+UV经典DS（探索变体）'}
sha=existing.sha

def save(name,obj):
    (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False,default=str),encoding='utf-8')

def csv_save(name,rows):
    with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)

def binary_metrics(y,p):
    if len(y)!=len(p) or not y or any(x not in (0,1) for x in list(y)+list(p)):
        raise ValueError('Equal nonempty binary label/prediction arrays required')
    tp=sum(a==1 and b==1 for a,b in zip(y,p)); fp=sum(a==0 and b==1 for a,b in zip(y,p))
    fn=sum(a==1 and b==0 for a,b in zip(y,p)); tn=sum(a==0 and b==0 for a,b in zip(y,p))
    return dict(N=len(y),TP=tp,FP=fp,FN=fn,TN=tn,accuracy=(tp+tn)/len(y),precision=tp/(tp+fp) if tp+fp else None,recall=tp/(tp+fn) if tp+fn else None,F1=2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else None)

def check_alignment(samples,rows,mode):
    if len(samples)!=len(rows) or len({r['pair_index'] for r in rows})!=len(samples):
        raise ValueError('Incomplete or duplicated pair IDs')
    for i,(sample,row) in enumerate(zip(samples,rows)):
        if row['pair_index']!=i or row['image_name']!=getattr(sample,mode):
            raise ValueError('Pair ID/name/order mismatch')

def predict_images(model,paths,device):
    # Ultralytics list-of-filenames loses names during EXIF transpose and assigns imageN.jpg.
    # Explicit OpenCV decoding preserves original BGR pixels; bind names only after pixel equality.
    images=[cv2.imread(p) for p in paths]
    if any(im is None for im in images): raise ValueError('Image decode failure')
    results=list(model.predict(source=images,conf=.5,imgsz=640,device=device,classes=[0],batch=16,verbose=False,save=False,half=False,rect=False,iou=.7,max_det=300))
    if len(results)!=len(paths): raise ValueError('Prediction count mismatch')
    for result,image,path in zip(results,images,paths):
        np.testing.assert_array_equal(result.orig_img,image)
        result.path=path
    return results


def run():
    start=time.perf_counter()
    OUT.mkdir(exist_ok=True); (OUT/'examples').mkdir(exist_ok=True)
    samples=[c.parse_sample_line(l) for l in (SOURCE/'paper_sample.txt').read_text(encoding='utf-8-sig').splitlines()]
    assert len(samples)==1528
    assert len({x.rgb for x in samples})==len({x.thermal for x in samples})==1528
    paths={mode:[str(SOURCE/('bgr' if mode=='rgb' else 'thermal')/getattr(s,mode)) for s in samples] for mode in ['rgb','thermal']}
    missing={mode:[p for p in ps if not Path(p).is_file()] for mode,ps in paths.items()}
    assert all(not x for x in missing.values()),missing
    sources=list((existing.CPP/'src').glob('*.cpp'))+list((existing.CPP/'main').glob('*.cpp'))+list((existing.CPP/'include').glob('*.h'))
    source_hashes={str(p):sha(p) for p in sources}
    models={}; metadata={}
    expected={'rgb':'1068cfddcca133ac2f6c66522d7a43f4988daa91023c1c8898bb6fbf4a894e32','thermal':'c068422549d5de50d532431892a940a44ce2300b1cb0fae2e12887749adc12e6'}
    for mode,name in [('rgb','best.pt'),('thermal','best-2.pt')]:
        p=ATT/name; assert sha(p)==expected[mode]
        copied=OUT/'weights'/name; copied.parent.mkdir(exist_ok=True); shutil.copy2(p,copied); assert sha(copied)==sha(p)
        model=YOLO(str(copied)); assert len(model.names)==1 and model.names[0].lower()=='fire',model.names
        models[mode]=model
        metadata[mode]=dict(original_checkpoint=str(p),checkpoint=str(copied),bytes=p.stat().st_size,checkpoint_sha256=sha(p),names=model.names,class_mapping='class0 fire/Fire -> fire; no smoke class',model_class=type(model.model).__name__,yaml=model.model.yaml,layer_types=[type(x).__name__ for x in model.model.model],architecture='Ultralytics YOLO26n (yolo26n.yaml, end2end=True, reg_max=1); not PPYOLOv2 or YOLO11',train_args=model.ckpt.get('train_args'),checkpoint_keys=list(model.ckpt),checkpoint_epoch=model.ckpt.get('epoch'),checkpoint_version=model.ckpt.get('version'),checkpoint_date=model.ckpt.get('date'),parameters=sum(p.numel() for p in model.model.parameters()),end2end=model.model.end2end)
    device='0' if torch.cuda.is_available() else 'cpu'; failures=[]
    try:
        if device=='0':
            a=torch.ones((16,16),device='cuda'); assert torch.isfinite(a@a).all(); torch.cuda.synchronize()
        smoke={mode:predict_images(model,paths[mode][:3],device) for mode,model in models.items()}
    except Exception as exc:
        failures.append(dict(stage='GPU smoke',error=repr(exc))); device='cpu'
        smoke={mode:predict_images(model,paths[mode][:3],device) for mode,model in models.items()}
    assert all(len(r)==3 for r in smoke.values())
    smoke_record={mode:[dict(pair_index=i,image_path=r.path,boxes=r.boxes.data.cpu().tolist(),speed=r.speed) for i,r in enumerate(rs)] for mode,rs in smoke.items()}
    save('smoke_3pairs.json',dict(device=device,results=smoke_record,failures=failures))
    print('REAL SMOKE PASS: 3 pairs each modality; device='+device,flush=True)
    from ultralytics.utils import nms
    nms_text=inspect.getsource(nms.non_max_suppression)
    (OUT/'ultralytics_postprocess_source.txt').write_text(nms_text,encoding='utf-8')
    assert 'pred[:, 4] > conf_thres' in nms_text, 'Inspect confidence boundary before execution'
    predictions={}; inference_seconds={}
    for mode,model in models.items():
        t=time.perf_counter(); records=[]
        # Batches are real calls; preserve the exact input order independently of loader behavior.
        for first in range(0,1528,16):
            batchpaths=paths[mode][first:first+16]
            try: results=predict_images(model,batchpaths,device)
            except Exception as exc:
                if device=='cpu': raise
                failures.append(dict(stage=mode+' full batch '+str(first),error=repr(exc),fallback='cpu')); device='cpu'
                results=predict_images(model,batchpaths,device)
            assert len(results)==len(batchpaths)
            for off,r in enumerate(results):
                i=first+off; p=Path(paths[mode][i]); assert Path(r.path).resolve()==p.resolve()
                ds=[]
                for b in r.boxes.data.cpu().tolist():
                    x1,y1,x2,y2,score,cls=b
                    assert int(cls)==0 and score>.5 and x2>x1 and y2>y1
                    ds.append(dict(bbox_xyxy=[x1,y1,x2,y2],score=score,class_id=int(cls),class_name=model.names[int(cls)]))
                records.append(dict(pair_index=i,source_line=i+1,image_name=p.name,image_path=str(p),image_sha256=sha(p),height=r.orig_shape[0],width=r.orig_shape[1],detections=ds,speed_ms=r.speed,device=device))
            if first%160==0: print(mode+' processed '+str(len(records))+'/1528',flush=True)
        check_alignment(samples,records,mode); inference_seconds[mode]=time.perf_counter()-t
        predictions[mode]=records; metadata[mode].update(processed_images=len(records),unique_ids=len({r['pair_index'] for r in records}),missing_images=[],retained_fire_boxes=sum(len(r['detections']) for r in records),actual_devices=sorted({r['device'] for r in records}),inference_seconds=inference_seconds[mode])
        save(mode+'_predictions.json',dict(metadata=metadata[mode],images=records))
    save('model_metadata.json',metadata)
    existing.OUT=OUT
    h,cal=existing.calibration()
    state=c.SensorState(); previous=None; sequence=0; resets=[]; rows=[]; pairing=[]; inputs=[]; expected_cpp=[]
    for i,sample in enumerate(samples):
        ts=datetime.strptime(sample.rgb.removesuffix('_bgr.jpg'),'%Y_%m_%d_%H_%M_%S_%f')
        gap=None if previous is None else (ts-previous).total_seconds()
        reset=previous is None or gap<=0 or gap>120
        if reset:
            state=c.SensorState(); sequence+=1; resets.append(dict(pair_index=i,source_line=i+1,gap_seconds=gap))
        previous=ts
        rgb=predictions['rgb'][i]; thermal=predictions['thermal'][i]
        rc=max((d['score'] for d in rgb['detections']),default=0.); tc=max((d['score'] for d in thermal['detections']),default=0.)
        sr=c.sensor_step(state,sample,rc,tc); state=sr.state
        spatial,pair=existing.geometry(rgb,thermal,h)
        fused={'legacy_fixed_01':c.fuse3(sr.rgb,sr.thermal,sr.uv,'spatial',.1),'classic_ds':c.fuse3(sr.rgb,sr.thermal,sr.uv,'ds'),'candidate_geometry_repeated':c.fuse3(sr.rgb,sr.thermal,sr.uv,'spatial',spatial),'recommended_spatial_then_ds':c.fuse_rgb_thermal_uv(sr.rgb,sr.thermal,sr.uv,'spatial',spatial)}
        row=dict(pair_index=i,source_line=i+1,rgb_name=sample.rgb,thermal_name=sample.thermal,Tmax=sample.temperature,UV_input=sample.uv,UV_reference_label=sample.uv,sequence_id=sequence,state_reset=int(reset),gap_seconds=gap,rgb_confidence=rc,thermal_confidence=tc,rgb_box_count=len(rgb['detections']),thermal_box_count=len(thermal['detections']),rgb_baseline_binary05=int(rc>=.5),thermal_baseline_binary05=int(tc>=.5),candidate_spatial_iou=spatial,top_score_pair_iou=pair['top_score_iou'],geometry_status=pair['status'])
        for name,m in [('rgb_bpa',sr.rgb),('thermal_bpa',sr.thermal),('uv_bpa',sr.uv)]:
            for key,val in zip(['fire','nofire','uncertain'],m.values()): row[name+'_'+key]=val
        for name,res in fused.items():
            for key,val in zip(['fire','nofire','uncertain'],res.mass.values()): row[name+'_'+key]=val
            row[name+'_binary05']=int(res.mass.fire>=.5)
            row[name+'_decision095']=c.decide(res.mass,.95)
            row[name+'_conflict_ab']=res.conflict; row[name+'_degenerate']=int(res.degenerate)
        rows.append(row); pairing.append(dict(pair_index=i,spatial_consistency=spatial,**pair))
        inputs.append(f'{int(reset)} {rc:.17g} {tc:.17g} {sample.temperature:.17g} {sample.uv}\n')
        expected_cpp.append([*sr.rgb.values(),*sr.thermal.values(),*sr.uv.values(),*fused['legacy_fixed_01'].mass.values(),*fused['classic_ds'].mass.values()])
    exe=ROOT/'fusion/cpp_reference_harness.exe'; shutil.copy2(exe,OUT/exe.name)
    proc=subprocess.run([str(OUT/exe.name)],input=''.join(inputs),capture_output=True,text=True,check=True)
    ref=np.array([list(map(float,l.split())) for l in proc.stdout.splitlines()]); exp=np.array(expected_cpp)
    assert ref.shape==exp.shape==(1528,15) and np.isfinite(ref).all()
    maxerr=float(np.abs(ref-exp).max()); assert maxerr<1e-10
    (OUT/'cpp_reference_inputs.txt').write_text(''.join(inputs),encoding='utf-8'); (OUT/'cpp_reference_results.txt').write_text(proc.stdout,encoding='utf-8')
    save('cpp_parity_verification.json',dict(rows=1528,values_per_row=15,maximum_absolute_error=maxerr,reference_exit_code=proc.returncode,reference_exe_sha256=sha(exe),reference_source=str(ROOT/'fusion/cpp_reference_harness.cpp'),reference_source_sha256=sha(ROOT/'fusion/cpp_reference_harness.cpp'),source_unchanged=all(sha(p)==v for p,v in source_hashes.items())))
    y=[r['UV_reference_label'] for r in rows]; metrics={}; coverage={}
    for name in METHODS:
        metrics[name]=binary_metrics(y,[r[name+'_binary05'] for r in rows])
        if name not in ['rgb_baseline','thermal_baseline']:
            ctr=Counter(r[name+'_decision095'] for r in rows)
            selected=[r for r in rows if r[name+'_decision095']!='uncertain']
            coverage[name]=dict(fire=ctr['fire'],nofire=ctr['nofire'],uncertain=ctr['uncertain'],decision_coverage=len(selected)/1528,uncertainty_rate=ctr['uncertain']/1528,degenerate_count=sum(r[name+'_degenerate'] for r in rows),covered_only_metrics=binary_metrics([r['UV_reference_label'] for r in selected],[int(r[name+'_decision095']=='fire') for r in selected]) if selected else None)
    csv_save('fusion_results.csv',rows); save('fusion_results.json',dict(primary='legacy_fixed_01',label_semantics='UV0 nofire, UV1 fire, user-supplied reference standard; not independent scene truth',rows=rows)); save('geometry_pairing.json',pairing)
    save('metrics.json',dict(reference='text column4 UV 0=nofire 1=fire',circularity='UV is simultaneously fusion input and evaluation reference; metrics are agreement with UV, not independent detection validity',threshold_protocol='Retain class0 boxes with YOLO.predict(conf=0.5) end2end postprocess strict score>0.5. No boxes -> frame confidence0. All binary methods use confidence/mF>=0.5.',UV_label_counts=dict(Counter(y)),methods=metrics,fusion_abstention_threshold=.95,abstention=coverage))
    # Representative real originals with actual retained NN boxes, scores and projected thermal polygon.
    indices=sorted({0,2,1527,int(np.argmax([r['candidate_spatial_iou'] for r in rows])),next((i for i,r in enumerate(rows) if r['UV_reference_label']==0),0),next((i for i,r in enumerate(rows) if r['UV_reference_label']==1),0),next((i for i,r in enumerate(rows) if r['rgb_baseline_binary05']!=r['UV_reference_label']),0)})
    examples=[]
    for i in indices:
        canvases=[]
        for mode in ['rgb','thermal']:
            rec=predictions[mode][i]; image=cv2.imread(rec['image_path']); assert image is not None
            for detection in rec['detections']:
                x1,y1,x2,y2=map(round,detection['bbox_xyxy']); cv2.rectangle(image,(x1,y1),(x2,y2),(0,255,0),2)
                cv2.putText(image,'fire '+format(detection['score'],'.3f'),(x1,max(15,y1-4)),cv2.FONT_HERSHEY_SIMPLEX,.42,(0,255,0),1,cv2.LINE_AA)
            if mode=='rgb' and pairing[i].get('thermal_detection'):
                poly=c.project_polygon(pairing[i]['thermal_detection']['bbox_xyxy'],h)
                cv2.polylines(image,[np.round(poly).astype(np.int32)],True,(255,0,255),2)
            target_height=360; width=round(image.shape[1]*target_height/image.shape[0]); image=cv2.resize(image,(width,target_height))
            band=np.zeros((70,width,3),np.uint8); cv2.putText(band,mode+' pair '+str(i)+' UV='+str(y[i]),(8,23),cv2.FONT_HERSHEY_SIMPLEX,.5,(255,255,255),1)
            cv2.putText(band,'s='+format(rows[i]['candidate_spatial_iou'],'.4f')+' boxes='+str(len(rec['detections'])),(8,49),cv2.FONT_HERSHEY_SIMPLEX,.5,(255,255,255),1)
            canvases.append(np.vstack((band,image)))
        dest=OUT/'examples'/('pair_'+str(i)+'.jpg'); assert cv2.imwrite(str(dest),np.hstack(canvases))
        examples.append(dict(pair_index=i,path=str(dest),sha256=sha(dest),note='Green real detector boxes; magenta candidate thermal polygon on RGB. UV is reference, not independent scene ground truth.'))
    elapsed=time.perf_counter()-start
    execution=dict(created_utc=datetime.now(timezone.utc).isoformat(),total_pairs=1528,models=metadata,device=device,GPU_name=torch.cuda.get_device_name(0) if torch.cuda.is_available() else None,torch_version=torch.__version__,ultralytics_version=ultralytics.__version__,fallback_failures=failures,inference_seconds=inference_seconds,elapsed_before_report_seconds=elapsed,paper_sample_sha256=sha(SOURCE/'paper_sample.txt'),source_algorithm_sha256=source_hashes,parameters=dict(window=10,time_lambda=.7,reliability=.3,T0=50,uv_alpha=.01,uv_beta=.1,imgsz=640,confidence=.5),temporal_policy='Exact file order, reset at initial/nonincreasing timestamp/gap>120s; heuristic acquisition boundaries, not independently verified scenes; adaptation from uninterrupted original loop',resets=resets,sequence_count=sequence,calibration=cal,spatial_nonzero_rows=sum(r['candidate_spatial_iou']>0 for r in rows),geometry_status_counts=dict(Counter(r['geometry_status'] for r in rows)),metrics=metrics,abstention=coverage,UV_label_counts=dict(Counter(y)),cpp_reference_max_abs_error=maxerr,examples=examples)
    save('execution_summary.json',execution)
    make_report(execution)
    with zipfile.ZipFile(REPORT) as z:
        assert z.testzip() is None and 'word/document.xml' in z.namelist()
    doc=Document(REPORT); readback='\n'.join([p.text for p in doc.paragraphs]+[cell.text for tab in doc.tables for rr in tab.rows for cell in rr.cells])
    assert all(s in readback for s in ['1528','YOLO26n','UV','0.1','未验证','循环'])
    (OUT/'report_readback.txt').write_text(readback,encoding='utf-8')
    loaded=list(csv.DictReader((OUT/'fusion_results.csv').open(encoding='utf-8-sig')))
    assert len(loaded)==1528 and len({r['pair_index'] for r in loaded})==1528 and set(loaded[0])==set(rows[0])
    assert all(sha(p)==v for p,v in source_hashes.items())
    for mode in ['rgb','thermal']:
        check_alignment(samples,json.loads((OUT/(mode+'_predictions.json')).read_text(encoding='utf-8'))['images'],mode)
    verification=dict(csv_rows=len(loaded),unique_ids=1528,csv_columns=len(loaded[0]),CSV_schema=list(loaded[0]),docx_crc_ok=True,docx_readback_ok=True,docx_embedded_examples=len(examples),docx_sha256=sha(REPORT),csv_sha256=sha(OUT/'fusion_results.csv'),json_sha256=sha(OUT/'fusion_results.json'),metrics_sha256=sha(OUT/'metrics.json'),source_cpp_unchanged=True,per_modality_unique_ids={'rgb':1528,'thermal':1528},missing_images=missing,elapsed_seconds=time.perf_counter()-start,report_path=str(REPORT))
    save('deliverable_verification.json',verification)
    print(json.dumps(dict(metrics=metrics,abstention=coverage,verification=verification,models={k:{a:v[a] for a in ['architecture','names','checkpoint_sha256','retained_fire_boxes','actual_devices']} for k,v in metadata.items()},spatial_nonzero_rows=execution['spatial_nonzero_rows'],geometry_status_counts=execution['geometry_status_counts'],cpp_max_abs_error=maxerr,calibration_loo_rmse=cal['leave_one_out_rmse_px']),ensure_ascii=False,indent=2),flush=True)


def make_report(s):
    REPORT.parent.mkdir(exist_ok=True); d=Document(); st=d.styles['Normal']; st.font.name='Microsoft YaHei'; st.font.size=Pt(10); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    def p(text): d.add_paragraph(str(text))
    def h(text): d.add_heading(text,1)
    def table(headers,records):
        t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
        for cell,val in zip(t.rows[0].cells,headers): cell.text=str(val)
        for rec in records:
            for cell,val in zip(t.add_row().cells,rec): cell.text=str(val)
    def fmt(v): return '未定义（分母为0）' if v is None else format(v,'.6f')
    d.add_heading('上传模型多模态实际测试：UV参考标准报告',0)
    p('真实测试1528对RGB与热成像。用户指定第四列UV：0无火、1有火。本报告依该标准计算全部帧级指标，同时明确UV既是融合输入又是评价标签，存在循环评价/信息泄漏，不是独立场景真值，也不是目标检测AP或IoU定位性能。不能将本结果解释为独立火灾检测准确率、泛化能力或融合必然提升。')
    p('生成UTC：'+s['created_utc'])
    h('1. 实际模型、输入与推理协议')
    p('来源：C:/Users/Administrator/AppData/Local/hermes/attachments；best.pt严格用于RGB，best-2.pt用于thermal。原始文件只读，复制权重校验SHA256一致，未替换或重训练。训练配置bgr.yaml/thermal.yaml与该映射吻合。当前检查实际为Ultralytics YOLO26n，而不是PPYOLOv2，也不应根据C3k2/C2PSA就称YOLO11。yaml_file=yolo26n.yaml、end2end=True、reg_max=1，DetectionModel/Detect，完整结构、类别名与训练参数见model_metadata.json。')
    for mode,v in s['models'].items():
        p(mode+'：'+v['checkpoint']); p('SHA256：'+v['checkpoint_sha256']); p('实际类别：'+json.dumps(v['names'],ensure_ascii=False)+'；类0 fire/Fire均映射fire，无smoke类。训练imgsz='+str(v['train_args']['imgsz'])+'，训练epochs='+str(v['train_args']['epochs'])+'；本次统一推理imgsz640，不修改模型。')
    table(['模态','图像数','唯一ID','缺失','保留fire框','执行设备','秒'],[[mode,v['processed_images'],v['unique_ids'],0,v['retained_fire_boxes'],','.join(v['actual_devices']),fmt(v['inference_seconds'])] for mode,v in s['models'].items()])
    p('源paper_sample.txt：'+str(SOURCE/'paper_sample.txt')+'；SHA256 '+s['paper_sample_sha256']+'。逐行RGB/thermal文件名与ID对齐，各1528唯一图像，原图SHA256写入预测JSON。第三列Tmax输入热BPA，第四列UV同时作为用户评价标准。温度单位及标定未独立验证。')
    p('Ultralytics '+s['ultralytics_version']+'，torch '+s['torch_version']+'；GPU '+str(s['GPU_name'])+'；device='+s['device']+'。真实GPU矩阵运算与两模态各3图smoke通过；失败/回退记录='+json.dumps(s['fallback_failures'],ensure_ascii=False)+'。')
    p('使用cv2.imread逐张解码原始BGR像素，并将真实图像数组输入YOLO.predict；Ultralytics对文件名列表会在EXIF转置后生成imageN.jpg，本次逐结果验证orig_img与对应解码数组像素完全一致后绑定原文件路径，不依赖自动生成文件名。全量显式每16张一批，末批8张。predict(conf=0.5,imgsz=640,classes=[0],half=False,rect=False,batch=16,iou=0.7,max_det=300)。实际end2end路径保留score>0.5（严格大于，而非≥），源码摘录已存。端到端头采用自身后处理，不假称标准NMS。帧置信度为保留fire框最大分数，无框为0；所有二分类方法统一conf或mF≥0.5判fire，其他判nofire；因此恰好0.5的原始框会被后处理剔除。无低阈值旧框参与几何，无smoke映射。')
    h('2. UV参考帧级结果（完整1528帧）')
    p('UV标签计数：'+json.dumps(s['UV_label_counts'],ensure_ascii=False)+'。Accuracy=(TP+TN)/N，Precision=TP/(TP+FP)，Recall=TP/(TP+FN)，F1=2TP/(2TP+FP+FN)；零分母标未定义，不伪造0或1。')
    table(['方法','TP','FP','FN','TN','Accuracy','Precision','Recall','F1'],[[NAMES[name],v['TP'],v['FP'],v['FN'],v['TN'],fmt(v['accuracy']),fmt(v['precision']),fmt(v['recall']),fmt(v['F1'])] for name,v in s['metrics'].items()])
    p('这些是真实输出与UV标准的符合程度；UV参与融合意味着融合方法评价存在循环性。RGB/thermal不直接使用UV标签，但训练样本独立性与用户模型训练/测试重叠未审计，亦不能假称无泄漏。未预设或承诺任何性能改善，未按测试结果调阈值。')
    h('3. 原C++主结果、BPA公式与安全边界')
    p('权威来源C:/lyk/paper/bpa_test/main/main.cpp及src四类BPA/融合源，均只读。main使用窗口10、λ0.7、RGB可靠性0.3、T0=50、UV α0.01/β0.1；只将随机置信度/温度替换为真实模型与文本输入。')
    p('稳定性S=exp(-std(最近10个置信度))。RGB当前(F,N,U)=(c*S*0.3,(1-c)*S,1-F-N)。热PC=1-exp(-10c)，PT=clamp(1-exp(-0.1*(Tmax-50)),0,1)，热(F,N,U)=(PC*PT*S,(1-c)*(1-PT)*S,1-F-N)。各当前BPA与previous按λ*previous+(1-λ)*current平滑后归一化。cold温度先clamp可避免指数溢出，结果与源码温度clamp等价。')
    p('UV初始q=0.001；UV1时q←q+0.01*q*(1-q)，UV0时q←q-0.1*q*(1-q)，限制[0,1]；UV当前(q,0.8*(1-q),0.2*(1-q))。RGB/热previous=(0.5,0.3,0.2)，UV previous=(0.1,0.7,0.05)，保留原不归一化startup，第一次平滑后才归一化。UV0/1不是直接注入(0,1,0)/(1,0,0)，而是按原状态更新。')
    p('二证据：K=aF*bN+aN*bF；F=aF*bF+aF*bU+aU*bF；N=aN*bN+aN*bU+aU*bN；U=aU*bU。经典DS先除1-K再normalize。空间分支Ksp=(1-s)K，Fsp=sF/(1-Ksp)，Nsp=N/(1-Ksp)，Usp=U/(1-Ksp)，再normalize；共同分母最终抵消，非退化输出实为(sF,N,U)/(sF+N+U)，不能声称动态K单独改善最终比例。')
    p('原fuse3先RGB/热再UV，两次s都固定0.1，内部虽计算PCR/经典DS但return abc_sp。主要结果忠实保留固定0.1两次规则；它不是测量的空间一致性。候选几何替换s并两次重复是明确扩展，仍把无几何位置的UV应用同一s，不能当成合理UV空间模型；可选探索变体只RGB/热用空间，UV用经典DS。未启用源码未调用UV gate。')
    p('实际同1528行真实输入运行原C++参考exe，比较每行RGB/热/UV BPA、原固定融合与经典DS共15值；最大绝对误差='+format(s['cpp_reference_max_abs_error'],'.3e')+' <1e-10，exit0。原源SHA附execution_summary.json，运行后全部不变。总/近总冲突或零支持Python回退vacuous(0,0,1)并degenerate标记，避免原normalize零除；实际发生计数见下表。这是显式安全差异，不能假称所有边界源码行为完全相同。')
    h('4. 时序边界')
    p(s['temporal_policy']+'。文件名小数按%f解析，真实时间尺度未确认；按样本平滑而非物理时长。真实重置记录='+json.dumps(s['resets'],ensure_ascii=False)+'；sequence数='+str(s['sequence_count'])+'。120秒仅继承已有采集间断启发式，非独立验证的场景边界；不据此消除序列/训练泄漏风险。')
    h('5. 候选配准与实际空间一致性（未验证）')
    cal=s['calibration']; p('从当前homo_main.cpp未注释七对点重新解析，thermal pts2→RGB pts1；RANSAC3像素、seed7，H='+json.dumps(cal['H'])+'。没有任何手工像素偏移、矩形挪动或伪造s。')
    p('拟合RMSE='+fmt(cal['fit_rmse_px'])+'px；留一RMSE='+fmt(cal['leave_one_out_rmse_px'])+'px，留一最大='+fmt(cal['leave_one_out_max_px'])+'px。点集中狭窄y带、留一不稳定；仅候选/探索，未验证跨场景、全视场、镜头位姿、深度视差。H来源图与测试场景独立性未确认，不能宣称已完成可靠配准。')
    table(['thermal点','RGB点','拟合误差px','留一误差px'],[[a,b,fmt(f),fmt(l)] for a,b,f,l in zip(cal['thermal_points'],cal['rgb_points'],cal['fit_errors_px'],cal['leave_one_out_errors_px'])])
    p('每帧最高分RGB框为锚，所有实际保留score>0.5 thermal框四角经H投影成多边形；polygonIoU选最大thermal框，相同IoU优先分数高。BPA仍用各模态独立最高分，几何最佳thermal可非最高分，rank/框/分数记录geometry_pairing.json。缺一模态框s=0；不重叠s=0；非法投影标记且s=0。不做人工offset，不用投影包围矩形IoU，不做多目标一对一匹配。s=0并不证明无火。')
    p('真实非零s帧数='+str(s['spatial_nonzero_rows'])+'；几何状态计数='+json.dumps(s['geometry_status_counts'],ensure_ascii=False)+'。')
    h('6. 0.95融合弃权协议（不同于统一0.5二分类）')
    p('mF≥0.95 fire，mN≥0.95 nofire，否则uncertain。fire阈值沿原main0.95；对称nofire及弃权为显式评价扩展。覆盖率=(fire+nofire)/1528；仅覆盖子集指标另列metrics.json，不能把弃权当正确，不与全量0.5指标混比。')
    table(['方法','Fire','NoFire','弃权','覆盖率','退化数'],[[NAMES[name],v['fire'],v['nofire'],v['uncertain'],fmt(v['decision_coverage']),v['degenerate_count']] for name,v in s['abstention'].items()])
    h('7. 原图真实检测样例')
    p('绿色是模型保留框与分数；紫色是候选H投影thermal polygon，仅可视化探索几何，不是标注真值。样例选择包括首尾、有/无UV、模型分歧与最大IoU，不是随机代表性性能抽样。')
    for ex in s['examples']:
        p('pair_index '+str(ex['pair_index'])); d.add_picture(ex['path'],width=Inches(6))
    h('8. 交付与解释边界')
    p('新独立目录fusion/supplied_models_uv_reference：两模态完整预测JSON、1528行CSV/JSON、metrics.json、执行/模型元数据、真实3对smoke、H与误差、真实几何配对、C++参考输入输出与误差、带框样例和校验文件。旧fusion数据及报告保持不变。')
    p('本次按用户UV标准给出完整指标，不再因缺独立GT拒绝计算；但其含义仅是UV符合度。需要独立盲评场景标签、用户模型训练数据重叠审计、全视场多场景外部配准验证及真正时序边界，才能评估独立真实火灾检测与融合收益。不能据这份循环评价论文式宣称泛化改善。')
    d.save(REPORT)


if __name__=='__main__': run()
