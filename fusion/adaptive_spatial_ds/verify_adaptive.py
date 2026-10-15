"""Independent result verification, no adaptive_core fusion or geometry calls."""
from pathlib import Path
import sys,json,csv,zipfile,hashlib,collections
import numpy as np
import cv2
from scipy.optimize import linear_sum_assignment
from scipy.stats import beta
from docx import Document
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(OUT));import run_adaptive as run
load=run.load

def polygon_iou(rb,tb,B):
 x,y,X,Y=rb;p=np.array([[x,y],[X,y],[X,Y],[x,Y]],np.float32);x,y,X,Y=tb;q=np.array([[x,y],[X,y],[X,Y],[x,Y]],np.float64);q=np.float32(np.c_[q,np.ones(4)]@np.asarray(B).T)
 inter=float(cv2.intersectConvexConvex(p,q)[0]);union=abs(cv2.contourArea(p))+abs(cv2.contourArea(q))-inter
 return max(0.,min(1.,inter/union)) if union else 0.
def metric(y,p):
 tp=int(sum((y==1)&(p==1)));fp=int(sum((y==0)&(p==1)));fn=int(sum((y==1)&(p==0)));tn=int(sum((y==0)&(p==0)))
 return dict(TP=tp,FP=fp,FN=fn,TN=tn)
def verify():
 y=np.array([v['UV_reference'] for v in run.prior.old.parse()]);groups=np.array(run.P['all_groups']);m=load(OUT/'metrics.json');am=load(OUT/'anchor_metrics.json');a=load(OUT/'anchor_evidence.json');ev=load(OUT/'frame_evidence.json');R=load(OUT/'rgb_all_full320.json');T=load(OUT/'thermal_all_predictions.json');B=run.A
 for p,h in load(OUT/'protected_previous.json').items():assert run.sha(p)==h,('prior changed',p)
 for name,v in run.P['models'].items():assert run.sha(v['source'])==v['sha256']
 receipt=load(OUT/'forward_all_receipt.json');views=load(OUT/'views_all.json');assert len(views)==receipt['actual_detector_views']==18336
 assert all(s[-2:] in [[320,320],[640,640]] for s in receipt['input_shapes'])
 counted=collections.Counter((v['modality'],v['pair_index'],v['view']) for v in views);assert all(v==1 for v in counted.values());assert len(counted)==18336
 hash_count=0
 for mod in ['rgb','thermal']:
  data=load(OUT/(mod+'_all_predictions.json'));old=load(run.PREV/(mod+'_full_predictions.json'));assert len(data)==1528
  for i,(r,prev) in enumerate(zip(data,old)):
   assert r['pair_index']==i and r['image_path']==prev['image_path'];assert run.sha(r['image_path'])==r['image_sha256']==prev['image_sha256'];im=cv2.imread(r['image_path']);assert list(im.shape[:2])==r['orig_shape'];hash_count+=1
   for box in r['detections']:
    assert box['actual_detector'] and box['score']>.005 and box['view'] in ['full320','full640','tile0_320','tile1_320','tile2_320','tile3_320'];x,y0,X,Y=box['crop'];xx,yy,XX,YY=box['bbox'];assert x-1e-3<=xx<XX<=X+1e-3 and y0-1e-3<=yy<YY<=Y+1e-3
  for pool,key in [('all_predictions','multiview'),('all_full320','original320')]:
   p=np.array([int(any(b['score']>.5 for b in r['detections'])) for r in load(OUT/(mod+'_'+pool+'.json'))]);assert metric(y,p)=={k:m['baselines'][mod+'_'+key][k] for k in ['TP','FP','FN','TN']}
 # Independent all32 anchor predictions from separately recomputed polygon matching.
 grid=load(OUT/'anchor_protocol.json')['grid'];assert len(grid)==32;pred=np.zeros((32,len(y)),int);all_matches=[]
 for i,(r,t) in enumerate(zip(R,T)):
  n=len(r['detections']);nt=len(t['detections']);s=np.array([[polygon_iou(rb['bbox'],tb['bbox'],B) for tb in t['detections']] for rb in r['detections']]).reshape(n,nt);w=np.zeros((n+nt,n+nt));w[:n,:nt]=np.where(s>=.001,s,-1);rr,tt=linear_sum_assignment(-w);pairs=[(int(j),int(k),float(s[j,k])) for j,k in zip(rr,tt) if j<n and k<nt and s[j,k]>=.001];all_matches.append(pairs)
  for k,cfg in enumerate(grid):
   values=[1-(1-r['detections'][j]['score'])*(1-cfg['rT']*(min(1.,ss/.3) if cfg['soft'] else 1.)*t['detections'][l]['score']) for j,l,ss in pairs];pred[k,i]=bool(values and max(values)>cfg['threshold'])
  for scope in ['fixed','full','oof']:
   o=a[i][scope];q=o['selected'];mass=o['masses'];assert abs(sum(mass)-1)<1e-10 and min(mass)>=0
   if q:
    j=q['rgb_index'];k=q['thermal_index'];assert q['rgb']==r['detections'][j] and q['thermal']==t['detections'][k];ss=polygon_iou(q['rgb']['bbox'],q['thermal']['bbox'],B);assert ss>=.001 and abs(ss-o['iou'])<1e-6
    f=1-(1-q['bpa_rgb'][0])*(1-q['bpa_thermal'][0]);assert abs(f-mass[0])<1e-10 and mass[1]==0 and o['K']==0
   else:assert o['fire']==0 and mass==[0.,0.,1.]
 trials=list(csv.DictReader((OUT/'anchor_all_32_trials.csv').open(encoding='utf-8-sig')));assert len(trials)==32
 for k,row in enumerate(trials):assert metric(y,pred[k])=={s:int(row[s]) for s in ['TP','FP','FN','TN']}
 for g in np.unique(groups):
  f=load(OUT/('anchor_fold_'+str(g)+'_freeze.json'));train=np.flatnonzero(groups!=g);test=np.flatnonzero(groups==g);assert train.tolist()==f['train_indices'] and test.tolist()==f['test_indices'];counts=[(metric(y[train],p[train])['FP'],-metric(y[train],p[train])['TP'],k) for k,p in enumerate(pred)];kk=min(counts)[2];assert kk==f['index']
  for i in test:assert a[i]['oof']['fire']==pred[kk,i]
 for scope,key in [('fixed','fixed_exposed'),('full','full_descriptive'),('oof','group_oof_exploratory')]:assert metric(y,np.array([o[scope]['fire'] for o in a]))=={k:am['metrics'][key][k] for k in ['TP','FP','FN','TN']}
 # Failed symmetric/calibrated branches: full-row independent BPA arithmetic, actual selected geometry and no missing-evidence overrides.
 for i,e in enumerate(ev):
  for scope in ['full','oof']:
   for branch,o in e[scope].items():
    mm=o['masses'];assert abs(sum(mm)-1)<1e-9;assert min(mm)>=-1e-12;q=o['selected']
    if q:
     af,an,au=q['bpa_rgb'];bf,bn,bu=q['bpa_thermal'];kk=af*bn+an*bf
     expected=[af*bf+af*bu+au*bf,an*bn+an*bu+au*bn,au*bu+kk] if kk>=.95 else [(af*bf+af*bu+au*bf)/(1-kk),(an*bn+an*bu+au*bn)/(1-kk),au*bu/(1-kk)]
     np.testing.assert_allclose(mm,expected,atol=1e-10);assert o['fire']==int(mm[0]>o['threshold'])
     if branch.startswith('nominal'):assert abs(polygon_iou(q['rgb']['bbox'],q['thermal']['bbox'],B)-o['iou'])<1e-5
    else:assert o['fire']==0 and mm==[0.,0.,1.]
 for branch in m['results']:
  for scope in ['full','oof']:assert metric(y,np.array([e[scope][branch]['fire'] for e in ev]))=={k:m['results'][branch][scope][k] for k in ['TP','FP','FN','TN']}
 # Explicit native320 thermal boundary drift, never silently claim exact old baseline reproduction.
 old=load(run.PREV/'thermal_full_predictions.json');new=load(OUT/'thermal_all_full320.json');drift=[]
 for i,(u,v) in enumerate(zip(old,new)):
  before=max([b['score'] for b in u['detections']],default=0);after=max([b['score'] for b in v['detections']],default=0)
  if (before>.5)!=(after>.5):drift.append(dict(pair_index=i,before_max=before,after_max=after,reference=int(y[i])))
 result=dict(utc=run.now(),verified=True,images_hash_verified=hash_count,neural_views_verified=len(views),anchor_trials_verified=32,anchor_folds_verified=6,full_rows_verified=len(y),all_previous_hashes_unchanged=True,models_unchanged=True,UV_runtime=False,thermal_native_boundary_drift=drift,zero_of_240_two_sided95_upper=float(beta.ppf(.975,1,240)),source_sha256=run.sha(run.prior.old.SOURCE))
 path=run.ROOT/'reports/自适应空间DS_双基线复验报告.docx'
 if path.exists():
  with zipfile.ZipFile(path) as z:assert z.testzip() is None
  doc=Document(path);txt='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join('|'.join(c.text for c in rr.cells) for t in doc.tables for rr in t.rows);(OUT/'report_readback.txt').write_text(txt,encoding='utf-8');assert 'OOF' in txt and '1265' in txt and '23' in txt and '181' in txt;result['docx']=dict(path=str(path),sha256=run.sha(path),crc_ok=True,tables=len(doc.tables))
 run.save('verification.json',result);print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':verify()
