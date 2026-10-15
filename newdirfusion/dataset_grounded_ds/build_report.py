from pathlib import Path
import json,csv,hashlib,zipfile,collections
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
OUT=Path(__file__).resolve().parent;ROOT=Path('C:/lyk/fire/qualityDetection');sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
m=json.loads((OUT/'safe_metrics.json').read_text());old=json.loads((OUT/'metrics.json').read_text());audit=json.loads((OUT/'training_dataset_audit.json').read_text());exp=json.loads((OUT/'safe_export_receipt.json').read_text());training=json.loads((OUT/'safe_training_complete.json').read_text());exe=json.loads((OUT/'safe_execution.json').read_text());metadata=json.loads((OUT/'safe_model_metadata.json').read_text());cfg=m['config_unchanged'];doc=Document()
normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10);normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
for style in ['Title','Heading 1','Heading 2']:
 doc.styles[style].font.name='Microsoft YaHei';doc.styles[style].element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.add_heading('真实数据审计与空间IoU + D-S火灾融合复核',0)
doc.add_paragraph('新增工作目录：'+str(OUT));doc.add_paragraph('架构：Ultralytics YOLO26n，分别训练可见光与热成像单类fire检测器；不是PPYOLOv2。原TorchDetection工程、历史模型、图表和输出均保持不变。')
doc.add_heading('1. 本次结论与口径',level=1)
full=m['results']['full'];held=m['results']['prior202_diagnostic'];passes=m['primary_full_target_met']
doc.add_paragraph('完整1528对、同一新模型基线的主目标：'+('达到' if passes else '未达到')+'。本结论由完整真实前向和逐帧混淆矩阵计算，不能把表中未达标项省略。')
doc.add_paragraph('主实验UV不进入网络、不进入融合决策；第4列UV0/1仅按用户约定作为结果参考。因此这些是UV参考一致性指标，不是经独立人工标注的场景真值准确率。整个1528对已经被历次实验查看，融合阈值来自此前探索，本次不属于独立验证。')
doc.add_paragraph('先前202帧留出子集现在仅称“诊断子集”：'+('达到严格目标' if held['target_met_same_models'] else '未达到严格目标')+'。如果该子集的热基线FNR=0，则数学上不存在更小的FNR；融合FNR=0也只能等于0，不能写成严格改善。')
doc.add_paragraph('对保留历史热基线（640输入、TP732/FP0/FN666/TN130）的完整集比较：'+('达到' if full['target_met_vs_historical_thermal'] else '未达到')+'。同一新模型基线是主要公平比较，历史更低FPR控制额外保留，未用弱热模型替换基线。')
doc.add_heading('2. 只读源工程、数据泄漏与重建数据',level=1)
doc.add_paragraph('读取train.py、cfgs/bgr.yaml和thermal.yaml、create_dataset.py、现有训练args/results与checkpoint字段。源train.py使用imgsz=320，所查附件best.pt与源RGB run5完全同SHA256；best-2.pt与热run5完全同SHA256。源RGB run6按其存储验证mAP50:95略优而进入诊断控制，不是根据1528标签挑选替代模型。')
t=doc.add_table(rows=1,cols=5);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['源模态/划分','图像数','与测试精确哈希重合','空标注数','格式无效']):c.text=v
for key,stats in audit['stats'].items():
 for cell,v in zip(t.add_row().cells,[key,stats['images'],stats['test_overlap'],stats['empty_annotation'],stats['invalid']]):cell.text=str(v)
doc.add_paragraph('重要新发现：当前RGB源train/val合计372张与测试图像精确SHA256重合；热源合计89张。历史checkpoint指向这些配置，但未保存其原始训练清单，因此不能反推每个旧权重实际训练成员；也不能声称旧权重的保留模型评测没有训练泄漏。先前仅对另一个RGB来源做审计不能代表本次TorchDetection来源。')
doc.add_paragraph('为避免继承火灾checkpoint潜在测试记忆，本次两个新网络均从源工程generic yolo26n.pt（80类COCO预训练）初始化，不使用best.pt/best-2.pt微调。将所有3056张保留测试图像精确哈希同时从新训练和新验证中排除。重复图像去重；同图像哈希却不同标注文件的组全部排除；真正越界框排除，不裁剪、不补造标注。空标签仅来自原有显式标签文件，未把无标签图当阴性。')
t=doc.add_table(rows=1,cols=5);t.style='Table Grid'
for cell,v in zip(t.rows[0].cells,['新模态','训练','验证','训练空标签','验证空标签']):cell.text=v
for mode,ss in exp['summary'].items():
 for cell,v in zip(t.add_row().cells,[mode,ss['train'],ss['val'],ss['empty_train'],ss['empty_val']]):cell.text=str(v)
reasons=collections.Counter(r['reason'] for r in exp['excluded_bounds']);doc.add_paragraph('额外排除记录：'+json.dumps(dict(reasons),ensure_ascii=False)+'。精确哈希训练—测试和训练—验证重合均为0。')
doc.add_paragraph('额外源域核查：RGB372张、热89张实验传感器文件名图像全部与保留集精确重合，剔除后未找到可核实的非重合同相机域训练图。非空fire框与UV0参考同时出现的重合图像，两模态各70张，说明框任务和UV参考语义不能直接等同；未擅自重标为负样本。源工程未找到README或独立sample/paired清单，不编造独立配对校准集。')
doc.add_paragraph('新RGB真实conf.01前向有1519框/917帧，但最高score仅0.11068566，冻结RGB门槛.2没有任何支持，所以双真实框主实验全漏检；这不是未运行推理，也不是用映射框补检测。新RGB>.5基线FPR也为0，严格“进一步降低RGB FPR”不可能低于0。保留真实失败结果，没有换弱基线或重调诊断标签阈值。')
doc.add_paragraph('划分按.rf及_aug增广家族、实验文件名的分钟块代理分组；数值文件名的真实视频/场景来源未知，近重复和同源视频泄漏不能排除。训练集哈希隔离不是新的独立融合测试。')
doc.add_heading('3. 实际训练、设备及模型选择',level=1)
doc.add_paragraph('实际GPU：'+exe['device']+'。两模态各完成30轮，输入320，batch32，workers0，seed42，AdamW lr0=.001，amp=False；热成像hsv_h/hsv_s设0以避免无意义色相/饱和度增强。所有缓存、复制数据、checkpoint、CSV和训练日志均写入新增目录。best.pt由验证fitness选择，没有用1528参考标签选择checkpoint。')
t=doc.add_table(rows=1,cols=5);t.style='Table Grid'
for cell,v in zip(t.rows[0].cells,['模态','已完成','耗时秒','最佳保存AP50','最佳保存AP50:95']):cell.text=v
for mode,ss in training['status'].items():
 bm=ss['best_saved_metrics']
 for cell,v in zip(t.add_row().cells,[mode,30,round(ss['seconds'],2),bm['metrics/mAP50(B)'],bm['metrics/mAP50-95(B)']]):cell.text=str(v)
doc.add_paragraph('验证AP来自各自有框数据集，不能拿两模态不同验证集的AP直接推导融合优劣；以下帧级FPR/FNR全部来自同一对齐保留配对集。')
for mode,md in metadata.items():doc.add_paragraph(mode+' best checkpoint：'+md['checkpoint']+'\nSHA256：'+md['checkpoint_sha256'])
doc.add_heading('4. 固定空间融合算法与局限',level=1)
doc.add_paragraph('在新前向之前冻结此前参数，不做新阈值搜索：'+json.dumps(cfg,ensure_ascii=False)+'。真实候选重新前向conf=.01，RGB支持筛选>=.2，热支持筛选>=.1；单模态基线仍采用真实score>.5。低置信度新前向不能用旧.1缓存假装生成。')
doc.add_paragraph('将每个真实热框四个角以thermal→RGB全仿射映射，计算凸多边形交并比，再在筛选后的候选间用带dummy的Hungarian最大和一对一匹配。没有用映射框冒充RGB检测。保留IoU>=.01的配对，选单对最大火质量，不把相关候选累乘。无真实双模态匹配时质量为0。')
doc.add_paragraph('名义火证据BPA=(r*c,0,1-r*c)，r=rmin+(1-rmin)*IoU。经典D-S给mF=1-(1-r*cR)*(1-r*cT)，严格mF>.3判火。已冻结rmin=1，动态空间可靠性折扣实际上关闭；空间只起低门槛门控作用。该BPA的冲突K恒为0，不能声称冲突解决或校准后验概率带来了增益。')
doc.add_paragraph('沿用七点全仿射候选：留一误差约5.77像素，独立全视场/新场景配准尚未验证；IoU=.01是宽松一致性门槛，不能描述为强空间一致性。UV乘法门控未用于主实验，避免输入—参考循环评价。')
doc.add_heading('5. 新模型完整集与诊断子集真实结果',level=1)
for subset,title in [('full','完整1528对：1398阳性/130阴性'),('prior202_diagnostic','先前202帧诊断：180阳性/22阴性')]:
 doc.add_heading(title,level=2);t=doc.add_table(rows=1,cols=9);t.style='Table Grid'
 for cell,v in zip(t.rows[0].cells,['方法','TP','FP','FN','TN','FPR%','FNR%','Precision','F1']):cell.text=v
 for method,label in [('rgb_baseline','RGB单模态>.5'),('thermal_baseline','热单模态>.5'),('fusion','固定空间IoU+D-S')]:
  mm=m['results'][subset][method];vals=[label,*[mm[x] for x in ['TP','FP','FN','TN']],f"{100*mm['FPR']:.4f}",f"{100*mm['FNR']:.4f}",'未定义' if mm['precision'] is None else f"{mm['precision']:.6f}",f"{mm['F1']:.6f}"]
  for cell,v in zip(t.add_row().cells,vals):cell.text=str(v)
 doc.add_paragraph('同模型严格目标判定：'+str(m['results'][subset]['target_met_same_models']))
drop=full['rgb_baseline']['FPR']-full['fusion']['FPR'];doc.add_paragraph(f'相对于同一RGB模型，完整集经验FPR绝对下降{100*drop:.4f}个百分点。'+('该观测下降不等同独立样本显著性检验。' if drop>0 else '本次没有降低，不满足显著降低RGB误报率的要求。')+'序列帧强相关，130阴性样本不能证明总体FPR为0。')
doc.add_heading('6. 为什么先前180帧RGB完全缺候选',level=1)
doc.add_paragraph('首先直接检查三个失败原始RGB/热配对与训练例图。RGB画面昏暗，主要是金属容器/器件和暗背景，三个例图没有清楚可辨的橙色明火；热图表现为亮热源。公共训练RGB例图主要是明显橙火/烟，热例图有高对比白色火形，存在可见性和域差异。不能从截图单独推断真实无火，也不能因UV1就补造RGB框。')
doc.add_picture(str(OUT/'failed_original_pairs.png'),width=Inches(6.3));doc.add_picture(str(OUT/'training_examples.png'),width=Inches(6.3))
doc.add_paragraph('新的真实对照表明，输入尺度是主要可复核原因之一：附件RGB run5在320下完整集TP1195而640下TP452（同>.5）；此前180帧在640 conf.1无候选，320新前向conf.01可恢复多数支持，固定双模态融合在该子集TP153。按源验证mAP95预选RGB run6在320时固定融合恢复180/180，但其源数据存在精确测试重合，不能据此宣称独立泛化。尺度改变同时改变热基线误报，所以本次完整报告同时保留同尺度和历史640热基线控制。')
t=doc.add_table(rows=1,cols=7);t.style='Table Grid'
for cell,v in zip(t.rows[0].cells,['旧模型诊断控制','融合TP','FP','FN','TN','FNR%','202诊断TP']):cell.text=v
for variant,rr in old['results'].items():
 fm=rr['full']['fusion'];vals=[variant,*[fm[k] for k in ['TP','FP','FN','TN']],f"{100*fm['FNR']:.4f}",rr['prior202_diagnostic']['fusion']['TP']]
 for cell,v in zip(t.add_row().cells,vals):cell.text=str(v)
doc.add_paragraph('“stronger_thermal_control”只表示该源模型储存的验证AP较高，保留实验域上的表现反而很差；没有将其选成主基线。新输出640热TP730而历史732的小差异被保留，可能是批量32对16等浮点阈值敏感，不把新旧计数混用。')
doc.add_heading('7. 新模型真实框示例、复现与交付验证',level=1)
doc.add_picture(str(OUT/'safe_failed_pair_detections.png'),width=Inches(6.3));doc.add_paragraph('示例框仅来自真实新网络检测；没有把仿射投影当作目标检测结果。匹配坐标、得分、IoU和质量保存在safe_matched_evidence.json。')
doc.add_paragraph('复现顺序：export_safe.py → train_safe.py → evaluate_safe.py → build_report.py；使用qualityDetection/.venv/Scripts/python.exe，在本新增目录运行。train_safe.py禁止resume、exist_ok=False，重复运行会生成新命名输出，不应覆盖本次结论；复现应使用另一个新目录以保留所有历史证据。')
doc.add_paragraph('关键文件：safe_export_receipt.json（逐样本哈希与分组）；safe_training.log/results.csv/best.pt/last.pt；safe_evaluation_plan.json（前向前冻结）；safe_rgb_predictions.json与safe_thermal_predictions.json（3056真实前向）；safe_frame_predictions.csv；safe_metrics.json；source_protected_sha256.json；safe_execution.json。')
doc.add_paragraph(f"实际只读源代码/配置/标签/权重/日志SHA复核文件数：{exe['protected_source_files_unchanged']}。训练/测试图片本次清单亦逐张SHA复核。未变更任何源标签、模型、训练代码或原报告。")
for p in ['safe_metrics.json','safe_export_receipt.json','safe_evaluation_plan.json','safe_training_complete.json']:doc.add_paragraph(p+' SHA256: '+sha(OUT/p))
report=OUT/'dataset_grounded_spatial_DS_Chinese.docx';doc.save(report)
with zipfile.ZipFile(report) as z:assert z.testzip() is None
read=Document(report);assert len(read.tables)>=6 and len(read.inline_shapes)==3;assert '不属于独立验证' in '\n'.join(p.text for p in read.paragraphs)
(OUT/'report_readback.json').write_text(json.dumps({'report':str(report),'sha256':sha(report),'paragraph_count':len(read.paragraphs),'table_count':len(read.tables),'image_count':len(read.inline_shapes),'zip_crc_ok':True,'headings':[p.text for p in read.paragraphs if p.style.name.startswith('Heading')],'tables':[[[cell.text for cell in row.cells] for row in table.rows] for table in read.tables]},ensure_ascii=False,indent=2),encoding='utf-8');print(str(report))
