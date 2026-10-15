from pathlib import Path
import json,csv,hashlib,zipfile,collections
from docx import Document
OUT=Path(__file__).resolve().parent;TEST=Path('C:/lyk/paper_data/paper_data_/paper_data');sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
metrics=json.loads((OUT/'safe_metrics.json').read_text());rows=list(csv.DictReader((OUT/'safe_frame_predictions.csv').open(encoding='utf-8-sig')));lines=[x.split() for x in (TEST/'paper_sample.txt').read_text().splitlines() if x.strip()]
assert len(rows)==1528 and [int(r['pair_index']) for r in rows]==list(range(1528));assert len(lines)==1528
cfg=metrics['config_unchanged'];predictions={m:json.loads((OUT/('safe_'+m+'_predictions.json')).read_text()) for m in ['rgb','thermal']};matched=json.loads((OUT/'safe_matched_evidence.json').read_text())
for i,(r,line) in enumerate(zip(rows,lines)):
 assert r['rgb_name']==line[0] and r['thermal_name']==line[1] and int(r['UV_reference'])==int(line[3])
 assert int(r['fusion_fire'])==int(float(r['mass_fire'])>cfg['decision_threshold'])
 for mode,col,folder,field in [('rgb',0,'bgr','rgb_baseline_fire'),('thermal',1,'thermal','thermal_baseline_fire')]:
  p=predictions[mode][i];assert p['pair_index']==i and p['image_name']==line[col];assert p['sha256']==sha(TEST/folder/line[col])
  assert int(r[field])==int(any(b['score']>.5 for b in p['detections']))
 evidence=matched[i];assert evidence['pair_index']==i
 mf=max([b['mass_fire'] for b in evidence['matched_evidence']],default=0.);assert abs(mf-float(r['mass_fire']))<1e-12
 for e in evidence['matched_evidence']:
  assert e['rgb_box'] in predictions['rgb'][i]['detections'] and e['thermal_box'] in predictions['thermal'][i]['detections']
  assert e['rgb_box']['score']>=cfg['tR'] and e['thermal_box']['score']>=cfg['tT'] and e['polygon_iou']>=cfg['iou_gate']
  assert abs(e['mass_fire']-(1-(1-e['rgb_box']['score'])*(1-e['thermal_box']['score'])))<1e-12
recomputed={}
for subset,rr in [('full',rows),('prior202_diagnostic',[r for r in rows if r['split']=='heldout'])]:
 methods={}
 for method,field in [('fusion','fusion_fire'),('rgb_baseline','rgb_baseline_fire'),('thermal_baseline','thermal_baseline_fire')]:
  cnt=collections.Counter((int(r['UV_reference']),int(r[field])) for r in rr);tp=cnt[(1,1)];fp=cnt[(0,1)];fn=cnt[(1,0)];tn=cnt[(0,0)]
  actual=metrics['results'][subset][method];assert [actual[k] for k in ['TP','FP','FN','TN']]==[tp,fp,fn,tn] and actual['N']==len(rr)
  assert actual['FPR']==fp/(fp+tn) and actual['FNR']==fn/(fn+tp);methods[method]={'TP':tp,'FP':fp,'FN':fn,'TN':tn,'FPR':fp/(fp+tn),'FNR':fn/(fn+tp)}
 fm=methods['fusion'];tm=methods['thermal_baseline'];rm=methods['rgb_baseline'];target=fm['FPR']<=tm['FPR'] and fm['FNR']<tm['FNR'] and fm['FPR']<rm['FPR'];assert target==metrics['results'][subset]['target_met_same_models'];methods['target_met_same_models']=target;recomputed[subset]=methods
export=json.loads((OUT/'safe_export_receipt.json').read_text());sourceaudit=json.loads((OUT/'training_dataset_audit.json').read_text());hashes={p['sha256'] for ps in predictions.values() for p in ps};assert len(hashes)==3056
uv_by_hash={p['sha256']:int(lines[p['pair_index']][3]) for ps in predictions.values() for p in ps};domain=json.loads((OUT/'domain_failure_evidence.json').read_text())
for mode,label in [('fireOrigin','rgb'),('thermal','thermal')]:
 sensor=[r for r in sourceaudit['records'] if r['mode']==mode and Path(r['image']).name.startswith('2026_')]
 assert len(sensor)==domain['sensor_named_source_images'][label] and all(r['test_overlap'] for r in sensor)
 negatives=sum(r['mode']==mode and r['test_overlap'] and uv_by_hash[r['sha256']]==0 and r['box_count']>0 for r in sourceaudit['records'])
 assert negatives==domain['negative_UV_reference_nonempty_fire_annotation_overlap'][label]
scores=[b['score'] for r in predictions['rgb'] for b in r['detections']];assert len(scores)==domain['clean_RGB_conf01_real_boxes'] and max(scores)==domain['clean_RGB_max_score']
assert sum(any(b['score']>=cfg['tR'] for b in r['detections']) for r in predictions['rgb'])==domain['clean_RGB_frames_at_prior_conf02']
assert not hashes&{r['sha256'] for r in export['manifest']}
for r in sourceaudit['records']:assert sha(r['image'])==r['sha256'] and sha(r['label'])==r['label_sha256']
for r in export['manifest']:assert sha(r['export_image'])==r['sha256'] and sha(r['export_label'])==r['label_sha256']
for mode in ['fireOrigin','thermal']:
 tr={r['sha256'] for r in export['manifest'] if r['mode']==mode and r['new_split']=='train'};vl={r['sha256'] for r in export['manifest'] if r['mode']==mode and r['new_split']=='val'};assert not tr&vl
 data=list(csv.DictReader((OUT/'safe_training'/mode/'results.csv').open()));assert len(data)==30 and [int(r['epoch']) for r in data]==list(range(1,31))
protected=json.loads((OUT/'source_protected_sha256.json').read_text());assert all(sha(p)==h for p,h in protected.items())
status=json.loads((OUT/'safe_training_complete.json').read_text());assert status['success'];assert all(v['successful'] for v in status['status'].values())
for v in status['status'].values():assert sha(v['best'])==v['best_sha256']
exe=json.loads((OUT/'safe_execution.json').read_text());assert sha(OUT/'safe_evaluation_plan.json')==exe['freeze_sha256']
report=OUT/'dataset_grounded_spatial_DS_Chinese.docx';readback=json.loads((OUT/'report_readback.json').read_text());assert sha(report)==readback['sha256']
with zipfile.ZipFile(report) as z:assert z.testzip() is None
read=Document(report);assert len(read.tables)==readback['table_count'] and len(read.inline_shapes)==readback['image_count']
assert any('不属于独立验证' in p.text for p in read.paragraphs)
receipt={'success':True,'independently_recomputed_confusion':recomputed,'train_val_test_exact_hash_disjoint':True,'verified_exported_images':len(export['manifest']),'verified_original_dataset_image_records':len(sourceaudit['records']),'verified_unique_reserved_images':len(hashes),'verified_source_protected_files':len(protected),'epochs_completed':{'rgb':30,'thermal':30},'report_sha256':sha(report),'report_crc_ok':True,'report_tables':len(read.tables),'report_images':len(read.inline_shapes),'source_unchanged':True,'training_completed':True,'no_UV_in_primary_decisions':True,'artifact_sha256':{p.name:sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in ['.json','.csv','.py','.docx'] and p.name!='final_verification.json'}}
(OUT/'final_verification.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8');print(json.dumps({k:v for k,v in receipt.items() if k!='artifact_sha256'},indent=2))
